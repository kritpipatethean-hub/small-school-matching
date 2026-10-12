"""OBEC small-school matching pipeline for QGIS 3.40 LTR.

This controller intentionally uses only the Python runtime bundled with QGIS
and invokes QNEAT3 through qgis_process.  External GIS sources are read from
the exact paths in machine.json; this module never recursively discovers GIS
files outside the workspace.
"""

from __future__ import annotations

import argparse
import csv
import ctypes
import datetime as dt
import hashlib
import json
import math
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import textwrap
import time
import traceback
import uuid
import xml.etree.ElementTree as ET
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Iterable

# โครงสร้างโฟลเดอร์ย่อยรายพื้นที่ — แหล่งความจริงเดียวอยู่ที่ area_layout.py ในโฟลเดอร์เดียวกัน
import area_layout as layout


VERSION = "1.8.1"
SCRIPT_DIR = Path(__file__).resolve().parent
WORKSPACE = SCRIPT_DIR.parent
DATA_DIR = WORKSPACE / "01 ข้อมูลกลาง"
AREA_ROOT = WORKSPACE / "02 งานรายพื้นที่"
PORTAL_ROOT = WORKSPACE / "05 พอร์ทัลส่งมอบ สพฐ"
MACHINE_PATH = SCRIPT_DIR / "machine.json"
AREAS_PATH = SCRIPT_DIR / "areas.json"

MASTER_CSV = DATA_DIR / "DMC682_สพป.csv"
SMALL_CSV = DATA_DIR / "DMC682_สพป_เล็ก.csv"
MAGNET_CSV = DATA_DIR / "Magnet_schools_with_coordinates.csv"
BOUNDARY_AUDIT_CSV = DATA_DIR / "cleaned_school_districts_สพป.csv"

GIB = 1024 ** 3
ROAD_EXPRESSION_FIELD = "fclass"
BOUNDARY_NAME_FIELD = "cleaned_sc"
ORIGIN_ID_FIELD = "โรงเรียน"
DESTINATION_ID_FIELD = "school_code"

# ช่วงขนาดโรงเรียนตามจำนวนนักเรียนรวม (ตามแผนที่ตัวอย่างของโครงการ) และสีประจำช่วง
# ใช้ร่วมกันทั้งตัววาดแผนที่ QC และพอร์ทัล (สีเดียวกันประกาศไว้ใน portal.css คลาส .size-bin-0 ถึง .size-bin-5
# สำหรับช่วงที่วาดบนแผนที่ และ .size-bin-off สำหรับแถวรวมของช่วงที่ไม่วาด)
SIZE_BINS: tuple[tuple[int, int, str], ...] = (
    (0, 10, "#D62728"),
    (11, 20, "#F26B21"),
    (21, 30, "#F7941D"),
    (31, 40, "#FFC000"),
    (41, 50, "#FFE84D"),
    (51, 60, "#B5D334"),
    (61, 100, "#3AB5E5"),
    (101, 120, "#1F4FD8"),
)
SMALL_SCHOOL_MAX_STUDENTS = 120
# แผนที่ขนาดโรงเรียนวาดเฉพาะโรงเรียนที่มีนักเรียนไม่เกินค่านี้ (ตัดสินร่วมกับผู้ใช้ 8 ก.ย. 2569)
# ช่วงที่เกินค่านี้ยังนับใน QC school_size และตารางบนเว็บ แต่ไม่วาดจุดบนแผนที่
SIZE_MAP_MAX_STUDENTS = 60
# พื้นที่จุด (points^2) ต่อช่วงที่วาด: นักเรียนมาก = จุดใหญ่; คำอธิบายสัญลักษณ์ใช้ค่าเดียวกัน
SIZE_MAP_MARKER_AREA: dict[str, float] = {
    "0-10": 34.0,
    "11-20": 56.0,
    "21-30": 84.0,
    "31-40": 118.0,
    "41-50": 158.0,
    "51-60": 204.0,
}
# ป้ายชื่อบนแผนที่ขนาดโรงเรียน (ตัดสินร่วมกับผู้ใช้ 8 ก.ย. 2569 รอบห้า): ทุกโรงเรียนที่วาด (ไม่เกิน SIZE_MAP_MAX_STUDENTS คน)
# มีป้ายชื่อตามทะเบียน DMC682 ส่วนแผนที่ผลการจับคู่ยังมีป้ายเฉพาะโรงเรียนคุณภาพชุมชน; รุ่นของกติกาป้ายบันทึกใน
# presentation.size_map_labels_version และ verify แบบเข้มบังคับให้ตรงค่านี้ (เพิ่มเลขเมื่อกติกาป้ายเปลี่ยน)
SIZE_MAP_LABELS_VERSION = 1
SIZE_MAP_LABEL_FONTSIZE = 9.0
# หน้าที่ 4 ของ PDF: คำอธิบายขนาดโรงเรียนขนาดเล็ก (ตารางช่วงขนาด ร้อยละ ประโยคสรุป วิธีอ่านแผนที่) ตัดสินร่วมกับผู้ใช้ 14 ก.ย. 2569
# รุ่นของหน้านี้บันทึกใน presentation.size_summary_page_version และ verify แบบเข้มบังคับให้ตรง (เพิ่มเลขเมื่อโครงหน้าเปลี่ยน)
SIZE_SUMMARY_PAGE_VERSION = 1
# แนวเขตตำบลจาก tha_admin3 ภายนอก (อ่านอย่างเดียวด้วย exact path ตอนวาด ไม่คัดลอกเข้า workspace)
TAMBON_LINE_COLOR = "#8A97A6"
# ภาพจำลองหลังจับคู่ (consolidation): รุ่นของบล็อกใน QC/catalog; verify แบบเข้มบังคับให้ตรงค่านี้
CONSOLIDATION_VERSION = 1
# เกณฑ์ระยะทางแผน 1.2 (เมตร) สำหรับตรวจคู่ OD ที่ตรวจรับแล้วในภาพจำลอง
MAX_DISTANCE_METRES = 10_000.0
# เกณฑ์ 7 ขนาดของ สพฐ. ตามคอลัมน์ "เกณฑ์7ขนาด" ใน DMC682 (รหัสตรงตามไฟล์ เพื่อให้ verify เทียบซ้ำได้)
SIZE7_BOUNDS: tuple[tuple[int, int | None, str], ...] = (
    (0, 0, "0_(0)"), (1, 20, "1.1_(1-20)"), (21, 40, "1.2_(21-40)"), (41, 60, "1.3_(41-60)"),
    (61, 80, "1.4_(61-80)"), (81, 100, "1.5_(81-100)"), (101, 120, "1.6_(101-120)"),
    (121, 200, "2_(121-200)"), (201, 300, "3_(201-300)"), (301, 499, "4_(301-499)"),
    (500, 1499, "5_(500-1499)"), (1500, 2499, "6_(1500-2499)"), (2500, None, "7_(>=2500)"),
)

_ACTIVE_SUBPROCESS: subprocess.Popen[str] | None = None
_QGIS_APP: Any | None = None


class PipelineError(RuntimeError):
    """Base error presented to the operator without a Python traceback."""


class SafetyError(PipelineError):
    """A machine, locking, path, or overwrite safety check failed."""


class QualityError(PipelineError):
    """A data or result quality-control check failed."""


class AreaTimeout(PipelineError):
    """The 90-minute total area deadline was reached."""


def configure_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure:
            try:
                reconfigure(encoding="utf-8", errors="replace")
            except (ValueError, OSError):
                pass


def utc_now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def local_now() -> str:
    return dt.datetime.now().astimezone().isoformat(timespec="seconds")


def human_gib(value: int | float) -> str:
    return f"{float(value) / GIB:,.2f} GB"


def ensure_file(path: Path, label: str) -> Path:
    try:
        resolved = path.resolve(strict=True)
    except FileNotFoundError as exc:
        raise SafetyError(f"ไม่พบ {label}: {path}") from exc
    if not resolved.is_file():
        raise SafetyError(f"{label} ไม่ใช่ไฟล์: {resolved}")
    return resolved


def ensure_within(path: Path, parent: Path, label: str) -> Path:
    resolved = path.resolve(strict=False)
    root = parent.resolve(strict=True)
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise SafetyError(f"{label} อยู่นอกขอบเขตที่อนุญาต: {resolved}") from exc
    if resolved == root:
        raise SafetyError(f"{label} ต้องไม่ชี้ที่ root โดยตรง: {resolved}")
    return resolved


def ensure_workspace_or_local_state(path: Path, label: str) -> Path:
    """Allow generated data only in the workspace or the dedicated temp root."""
    try:
        return ensure_within(path, WORKSPACE, label)
    except SafetyError:
        return ensure_within(path, local_state_root(), label)


def load_json(path: Path, label: str) -> dict[str, Any]:
    exact = ensure_file(path, label)
    try:
        with exact.open("r", encoding="utf-8-sig") as handle:
            value = json.load(handle)
    except (OSError, json.JSONDecodeError) as exc:
        raise PipelineError(f"อ่าน {label} ไม่สำเร็จ: {exc}") from exc
    if not isinstance(value, dict):
        raise PipelineError(f"{label} ต้องเป็น JSON object")
    return value


def load_settings() -> tuple[dict[str, Any], dict[str, Any]]:
    machine = load_json(MACHINE_PATH, "machine.json")
    catalog = load_json(AREAS_PATH, "areas.json")
    if machine.get("schema_version") != 1:
        raise PipelineError("machine.json schema_version ต้องเป็น 1")
    if catalog.get("schema_version") != 1:
        raise PipelineError("areas.json schema_version ต้องเป็น 1")
    for key in ("qgis", "external_gis", "limits", "analysis"):
        if not isinstance(machine.get(key), dict):
            raise PipelineError(f"machine.json ไม่มี object: {key}")
    if not isinstance(catalog.get("areas"), list):
        raise PipelineError("areas.json ไม่มีรายการ areas")
    return machine, catalog


def write_json(path: Path, value: Any) -> None:
    ensure_workspace_or_local_state(path, "ไฟล์ JSON ผลลัพธ์")
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def sha256_file(path: Path) -> str:
    checksum = hashlib.sha256()
    with ensure_file(path, "ไฟล์สำหรับ fingerprint").open("rb") as handle:
        for chunk in iter(lambda: handle.read(4 * 1024 * 1024), b""):
            checksum.update(chunk)
    return checksum.hexdigest().upper()


def sampled_file_fingerprint(path: Path) -> dict[str, Any]:
    """Fingerprint one exact file without traversing its parent directory."""
    exact = ensure_file(path, "ไฟล์ GIS ภายนอก")
    stat = exact.stat()
    checksum = hashlib.sha256()
    checksum.update(str(stat.st_size).encode("ascii"))
    checksum.update(b"\0")
    with exact.open("rb") as handle:
        head = handle.read(1024 * 1024)
        checksum.update(head)
        if stat.st_size > len(head):
            handle.seek(max(0, stat.st_size - 1024 * 1024))
            checksum.update(handle.read(1024 * 1024))
    return {
        "path": str(exact),
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "algorithm": "sha256(size+first1MiB+last1MiB)",
        "digest": checksum.hexdigest().upper(),
    }


def shapefile_fingerprint(path: Path) -> dict[str, Any]:
    """Read a fixed set of shapefile components; never glob or recurse."""
    shp = ensure_file(path, "Shapefile")
    components: list[dict[str, Any]] = []
    for suffix in (".shp", ".shx", ".dbf", ".prj"):
        component = shp.with_suffix(suffix)
        components.append(sampled_file_fingerprint(component))
    optional_cpg = shp.with_suffix(".cpg")
    if optional_cpg.is_file():
        components.append(sampled_file_fingerprint(optional_cpg))
    combined = hashlib.sha256()
    for item in components:
        combined.update(Path(item["path"]).suffix.lower().encode("ascii"))
        combined.update(str(item["size"]).encode("ascii"))
        combined.update(item["digest"].encode("ascii"))
    return {
        "path": str(shp),
        "algorithm": "combined-exact-sidecar-sampled-v1",
        "digest": combined.hexdigest().upper(),
        "components": components,
    }


def full_file_fingerprint(path: Path, label: str = "ไฟล์ข้อมูลกลาง") -> dict[str, Any]:
    """Hash one small, exact file in full without inspecting its directory."""
    exact = ensure_file(path, label)
    return {
        "path": str(exact),
        "size": exact.stat().st_size,
        "algorithm": "sha256-full-v1",
        "digest": sha256_file(exact),
    }


def resolve_catalog_data_path(raw: Any, label: str) -> Path:
    """Resolve an areas.json path and confine it to 01 ข้อมูลกลาง."""
    if not isinstance(raw, str) or not raw.strip():
        raise QualityError(f"{label} ไม่มี relative path")
    if "|" in raw:
        raise SafetyError(f"{label} ต้องเป็นพาธไฟล์ล้วน ห้ามต่อ provider URI")
    relative = Path(raw)
    if relative.is_absolute():
        raise SafetyError(f"{label} ต้องเป็น relative path: {raw}")
    resolved = ensure_within(SCRIPT_DIR / relative, DATA_DIR, label)
    return ensure_file(resolved, label)


MACHINE_SPECIFIC_FINGERPRINT_KEYS = frozenset({"path", "mtime_ns"})


def portable_fingerprint(value: Any) -> Any:
    """ตัดส่วนที่ผูกกับเครื่อง (พาธเต็มและเวลาแก้ไขไฟล์) ออกก่อนเทียบ fingerprint ให้เทียบตามเนื้อไฟล์
    (ขนาด + digest + วิธีคำนวณ + ชนิด) เพื่อให้ workspace ที่ clone ไปเครื่องอื่นตรวจรับผ่าน (1.8.1)"""
    if isinstance(value, dict):
        return {key: portable_fingerprint(item) for key, item in value.items() if key not in MACHINE_SPECIFIC_FINGERPRINT_KEYS}
    if isinstance(value, list):
        return [portable_fingerprint(item) for item in value]
    return value


def vector_source_fingerprint(path: Path) -> dict[str, Any]:
    exact = ensure_file(path, "แหล่งขอบเขต")
    suffix = exact.suffix.lower()
    if suffix == ".shp":
        value = shapefile_fingerprint(exact)
        value["kind"] = "external_shapefile"
        return value
    if suffix == ".gpkg":
        value = full_file_fingerprint(exact, "boundary override")
        value["kind"] = "workspace_override"
        return value
    raise SafetyError(f"ชนิดไฟล์ขอบเขตไม่รองรับ: {exact.suffix}")


def resolve_boundary_source(
    machine: dict[str, Any], area: dict[str, Any]
) -> tuple[Path, str]:
    raw_override = area.get("boundary_override")
    if raw_override is not None:
        return resolve_catalog_data_path(
            raw_override, f"boundary_override ของ {area.get('name', area.get('areacode'))}"
        ), "workspace_override"
    external = exact_external_paths(machine)
    return external["education_boundaries"], "external_shapefile"


def csv_rows(path: Path) -> Iterable[dict[str, str]]:
    exact = ensure_file(path, "CSV ต้นทาง")
    with exact.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise QualityError(f"CSV ไม่มี header: {exact}")
        for row in reader:
            yield {str(key): (value or "") for key, value in row.items()}


def validate_catalog(catalog: dict[str, Any], verify_hashes: bool = True) -> dict[str, Any]:
    areas = catalog["areas"]
    if len(areas) != 183:
        raise QualityError(f"areas.json มี {len(areas)} พื้นที่; ต้องเป็น 183")

    by_code: dict[str, dict[str, Any]] = {}
    ids: set[str] = set()
    names: set[str] = set()
    provinces: set[str] = set()
    allowed_status = {"ready", "pending", "blocked", "empty"}
    override_areas: list[dict[str, Any]] = []
    for item in areas:
        if not isinstance(item, dict):
            raise QualityError("รายการใน areas.json ต้องเป็น object")
        code = str(item.get("areacode", ""))
        if not re.fullmatch(r"\d{8}", code):
            raise QualityError(f"areacode ไม่ใช่เลข 8 หลัก: {code!r}")
        if code in by_code:
            raise QualityError(f"areacode ซ้ำ: {code}")
        expected_id = f"spap-{code}"
        if item.get("id") != expected_id or expected_id in ids:
            raise QualityError(f"id พื้นที่ไม่ถูกต้องหรือซ้ำ: {item.get('id')!r}")
        name = str(item.get("name", "")).strip()
        province = str(item.get("province", "")).strip()
        if not name.startswith("สพป.") or not province:
            raise QualityError(f"ชื่อพื้นที่/จังหวัดไม่สมบูรณ์สำหรับ {code}")
        if name in names:
            raise QualityError(f"ชื่อพื้นที่ซ้ำ: {name}")
        if item.get("status") not in allowed_status:
            raise QualityError(f"status ไม่รองรับสำหรับ {code}: {item.get('status')}")
        folder_name = str(item.get("folder_name", ""))
        if not folder_name or any(char in folder_name for char in '<>:"/\\|?*'):
            raise QualityError(f"folder_name ไม่ปลอดภัยสำหรับ {code}: {folder_name!r}")
        if folder_name in {".", ".."}:
            raise QualityError(f"folder_name ไม่ปลอดภัยสำหรับ {code}")
        if "boundary_override" in item:
            resolve_catalog_data_path(
                item["boundary_override"], f"boundary_override ของ {code}"
            )
            override_areas.append(item)
        by_code[code] = item
        ids.add(expected_id)
        names.add(name)
        provinces.add(province)

    if len(provinces) != 77:
        raise QualityError(f"areas.json มี {len(provinces)} จังหวัด; ต้องเป็น 77")

    source_meta = catalog.get("source", {})
    if resolve_catalog_data_path(source_meta.get("path"), "source DMC master") != MASTER_CSV.resolve():
        raise QualityError("areas.json source.path ไม่ชี้ DMC master ที่กำหนด")
    expected_master_hash = str(source_meta.get("sha256", "")).upper()
    actual_hash = sha256_file(MASTER_CSV) if verify_hashes else None
    if verify_hashes and actual_hash != expected_master_hash:
        raise QualityError(
            "fingerprint DMC master เปลี่ยนจากตอนสร้าง areas.json; "
            "ต้องสร้าง catalog ใหม่และตรวจรับก่อนรัน"
        )

    master_mapping: dict[str, tuple[str, str]] = {}
    master_rows = 0
    for row in csv_rows(MASTER_CSV):
        master_rows += 1
        code = row.get("areacode", "").strip()
        value = (
            row.get("ชื่อจังหวัด", "").strip(),
            row.get("ชื่อเขต", "").strip(),
        )
        if not re.fullmatch(r"\d{8}", code) or not all(value):
            raise QualityError(f"DMC master มีรหัส/ชื่อว่างที่แถว {master_rows + 1}")
        previous = master_mapping.setdefault(code, value)
        if previous != value:
            raise QualityError(f"DMC master map รหัส {code} ไปหลายชื่อ")
    expected_rows = int(source_meta.get("row_count", -1))
    if master_rows != expected_rows or len(master_mapping) != 183:
        raise QualityError(
            f"DMC master count เปลี่ยน: rows={master_rows}, areas={len(master_mapping)}"
        )
    for code, (province, name) in master_mapping.items():
        catalog_item = by_code.get(code)
        if not catalog_item or catalog_item["province"] != province or catalog_item["name"] != name:
            raise QualityError(f"areas.json ไม่ตรง DMC master สำหรับ areacode {code}")

    if verify_hashes:
        checks = catalog.get("cross_checks", {})
        pairs = (
            (SMALL_CSV, checks.get("small_source", {}), "small_source"),
            (MAGNET_CSV, checks.get("magnet_source", {}), "magnet_source"),
        )
        for path, metadata, label in pairs:
            if resolve_catalog_data_path(metadata.get("path"), label) != path.resolve():
                raise QualityError(f"areas.json {label}.path ชี้ไฟล์ผิด")
            expected = str(metadata.get("sha256", "")).upper()
            if not expected or sha256_file(path) != expected:
                raise QualityError(
                    f"fingerprint เปลี่ยน: {path.name}; ต้องตรวจข้อมูลกลางก่อนรัน"
                )

        mapping_meta = checks.get("boundary_mapping_source", {})
        mapping_path = resolve_catalog_data_path(
            mapping_meta.get("path"), "boundary_mapping_source"
        )
        if mapping_path != BOUNDARY_AUDIT_CSV.resolve():
            raise QualityError("boundary_mapping_source.path ชี้ไฟล์ผิด")
        mapping_rows = sum(1 for _ in csv_rows(mapping_path))
        if mapping_rows != int(mapping_meta.get("row_count", -1)):
            raise QualityError(
                f"จำนวนแถว boundary mapping เปลี่ยน: ได้ {mapping_rows}, "
                f"catalog={mapping_meta.get('row_count')}"
            )
        if sha256_file(mapping_path) != str(mapping_meta.get("sha256", "")).upper():
            raise QualityError("fingerprint boundary mapping เปลี่ยน; ต้องตรวจข้อมูลกลางก่อนรัน")

        override_meta = checks.get("boundary_override_source", {})
        if override_areas:
            override_path = resolve_catalog_data_path(
                override_meta.get("path"), "boundary_override_source"
            )
            expected_override_hash = str(override_meta.get("sha256", "")).upper()
            if sha256_file(override_path) != expected_override_hash:
                raise QualityError("fingerprint boundary override เปลี่ยน; ต้องตรวจรับก่อนรัน")
            for item in override_areas:
                if resolve_catalog_data_path(
                    item["boundary_override"], f"boundary_override ของ {item['areacode']}"
                ) != override_path:
                    raise QualityError(
                        f"boundary_override ของ {item['areacode']} ไม่ตรง cross_checks"
                    )
            validate_boundary_override_file(override_path, override_meta, override_areas)
        elif override_meta:
            raise QualityError("มี boundary_override_source แต่ไม่มีพื้นที่ใดใช้งาน")

    status_counts: dict[str, int] = {key: 0 for key in sorted(allowed_status)}
    for item in areas:
        status_counts[item["status"]] += 1
    expected_status_counts = catalog.get("expected_status_counts")
    if not isinstance(expected_status_counts, dict):
        raise QualityError("areas.json ไม่มี expected_status_counts")
    normalized_expected = {
        key: int(expected_status_counts.get(key, -1)) for key in sorted(allowed_status)
    }
    if set(expected_status_counts) != allowed_status or status_counts != normalized_expected:
        raise QualityError(
            f"จำนวนสถานะ areas.json ไม่ตรงค่าคาด: ได้ {status_counts}, "
            f"คาด {normalized_expected}"
        )
    return {
        "area_count": len(areas),
        "province_count": len(provinces),
        "row_count": master_rows,
        "source_sha256": actual_hash or expected_master_hash,
        "status_counts": status_counts,
        "by_code": by_code,
    }


class MEMORYSTATUSEX(ctypes.Structure):
    _fields_ = [
        ("dwLength", ctypes.c_ulong),
        ("dwMemoryLoad", ctypes.c_ulong),
        ("ullTotalPhys", ctypes.c_ulonglong),
        ("ullAvailPhys", ctypes.c_ulonglong),
        ("ullTotalPageFile", ctypes.c_ulonglong),
        ("ullAvailPageFile", ctypes.c_ulonglong),
        ("ullTotalVirtual", ctypes.c_ulonglong),
        ("ullAvailVirtual", ctypes.c_ulonglong),
        ("ullAvailExtendedVirtual", ctypes.c_ulonglong),
    ]


def available_memory_bytes() -> int:
    if os.name != "nt":
        try:
            page_size = os.sysconf("SC_PAGE_SIZE")
            pages = os.sysconf("SC_AVPHYS_PAGES")
            return int(page_size * pages)
        except (AttributeError, OSError, ValueError) as exc:
            raise SafetyError("ไม่สามารถตรวจ RAM ว่างบนระบบนี้") from exc
    status = MEMORYSTATUSEX()
    status.dwLength = ctypes.sizeof(status)
    if not ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(status)):
        raise SafetyError("GlobalMemoryStatusEx ตรวจ RAM ไม่สำเร็จ")
    return int(status.ullAvailPhys)


def workspace_size_bytes() -> int:
    """Recursively count only this workspace and never follow links/junctions."""
    total = 0
    stack = [WORKSPACE.resolve(strict=True)]
    while stack:
        current = stack.pop()
        try:
            with os.scandir(current) as iterator:
                for entry in iterator:
                    try:
                        if entry.is_symlink():
                            continue
                        if entry.is_dir(follow_symlinks=False):
                            stack.append(Path(entry.path))
                        elif entry.is_file(follow_symlinks=False):
                            total += entry.stat(follow_symlinks=False).st_size
                    except OSError:
                        continue
        except OSError:
            continue
    return total


if os.name == "nt":
    from ctypes import wintypes

    TH32CS_SNAPPROCESS = 0x00000002
    INVALID_HANDLE_VALUE = ctypes.c_void_p(-1).value

    class PROCESSENTRY32W(ctypes.Structure):
        _fields_ = [
            ("dwSize", wintypes.DWORD),
            ("cntUsage", wintypes.DWORD),
            ("th32ProcessID", wintypes.DWORD),
            ("th32DefaultHeapID", ctypes.POINTER(ctypes.c_ulong)),
            ("th32ModuleID", wintypes.DWORD),
            ("cntThreads", wintypes.DWORD),
            ("th32ParentProcessID", wintypes.DWORD),
            ("pcPriClassBase", wintypes.LONG),
            ("dwFlags", wintypes.DWORD),
            ("szExeFile", wintypes.WCHAR * 260),
        ]


def process_snapshot() -> list[tuple[int, str]]:
    if os.name != "nt":
        return []
    kernel32 = ctypes.windll.kernel32
    snapshot = kernel32.CreateToolhelp32Snapshot(TH32CS_SNAPPROCESS, 0)
    if snapshot == INVALID_HANDLE_VALUE:
        raise SafetyError("ตรวจรายการ process ไม่สำเร็จ")
    entry = PROCESSENTRY32W()
    entry.dwSize = ctypes.sizeof(entry)
    result: list[tuple[int, str]] = []
    try:
        success = kernel32.Process32FirstW(snapshot, ctypes.byref(entry))
        while success:
            result.append((int(entry.th32ProcessID), str(entry.szExeFile)))
            success = kernel32.Process32NextW(snapshot, ctypes.byref(entry))
    finally:
        kernel32.CloseHandle(snapshot)
    return result


def qgis_processes() -> dict[str, list[tuple[int, str]]]:
    gui_names = {
        "qgis.exe",
        "qgis-bin.exe",
        "qgis-ltr.exe",
        "qgis-ltr-bin.exe",
    }
    worker_names = {"qgis_process.exe", "qgis_process-qgis-ltr.exe"}
    gui: list[tuple[int, str]] = []
    workers: list[tuple[int, str]] = []
    for pid, name in process_snapshot():
        lowered = name.lower()
        if lowered in gui_names:
            gui.append((pid, name))
        elif lowered in worker_names:
            workers.append((pid, name))
    return {"gui": gui, "workers": workers}


def pid_is_running(pid: int) -> bool:
    if pid <= 0:
        return False
    if os.name != "nt":
        try:
            os.kill(pid, 0)
            return True
        except OSError:
            return False
    PROCESS_QUERY_LIMITED_INFORMATION = 0x1000
    handle = ctypes.windll.kernel32.OpenProcess(
        PROCESS_QUERY_LIMITED_INFORMATION, False, pid
    )
    if not handle:
        return False
    ctypes.windll.kernel32.CloseHandle(handle)
    return True


def local_state_root(create: bool = True) -> Path:
    raw = os.environ.get("LOCALAPPDATA")
    if not raw:
        raise SafetyError("ไม่พบตัวแปร LOCALAPPDATA")
    root = Path(raw).resolve(strict=True) / "OBEC_SMS_Pipeline"
    if create:
        root.mkdir(parents=True, exist_ok=True)
    return root


def current_lock_info() -> dict[str, Any] | None:
    lock = local_state_root(create=False) / "sms_pipeline.lock"
    if not lock.exists():
        return None
    try:
        with lock.open("r", encoding="utf-8-sig") as handle:
            value = json.load(handle)
        return value if isinstance(value, dict) else {"invalid": True}
    except (OSError, json.JSONDecodeError):
        return {"invalid": True, "path": str(lock)}


@contextmanager
def pipeline_lock(action: str, area_code: str | None = None) -> Iterable[Path]:
    root = local_state_root()
    lock = root / "sms_pipeline.lock"
    nonce = uuid.uuid4().hex
    payload = {
        "pid": os.getpid(),
        "nonce": nonce,
        "action": action,
        "areacode": area_code,
        "started_at": local_now(),
        "workspace": str(WORKSPACE),
    }
    try:
        descriptor = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL)
    except FileExistsError:
        info = current_lock_info() or {}
        stale_pid = int(info.get("pid", -1)) if str(info.get("pid", "")).lstrip("-").isdigit() else -1
        if stale_pid > 0 and pid_is_running(stale_pid):
            raise SafetyError(
                f"มีงานระบบกำลังทำงานอยู่ (PID {stale_pid}, action={info.get('action')})"
            )
        stale = root / f"sms_pipeline.stale.{dt.datetime.now():%Y%m%d-%H%M%S}.{uuid.uuid4().hex[:8]}.lock"
        try:
            os.replace(lock, stale)
        except OSError as exc:
            raise SafetyError(f"พบ lock เก่าและย้ายออกไม่ได้: {lock}") from exc
        descriptor = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.flush()
            os.fsync(handle.fileno())
        yield lock
    finally:
        try:
            with lock.open("r", encoding="utf-8-sig") as handle:
                current = json.load(handle)
            if current.get("nonce") == nonce:
                lock.unlink()
        except (FileNotFoundError, OSError, json.JSONDecodeError):
            pass


def remaining_seconds(deadline: float | None) -> float | None:
    if deadline is None:
        return None
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise AreaTimeout("เกินเวลารวมต่อพื้นที่ 90 นาที")
    return remaining


def require_deadline(deadline: float, stage: str) -> None:
    if time.monotonic() >= deadline:
        raise AreaTimeout(f"ครบเวลารวมต่อพื้นที่ก่อนขั้น {stage}")


def stop_process(proc: subprocess.Popen[str]) -> None:
    if proc.poll() is not None:
        return
    try:
        if os.name == "nt":
            proc.send_signal(signal.CTRL_BREAK_EVENT)
        else:
            proc.terminate()
        proc.wait(timeout=10)
    except (OSError, subprocess.TimeoutExpired):
        try:
            proc.kill()
            proc.wait(timeout=10)
        except (OSError, subprocess.TimeoutExpired):
            pass


def run_external(
    arguments: list[str],
    timeout_seconds: float,
    stdin_path: Path | None = None,
    log_path: Path | None = None,
    minimum_available_memory_bytes: int | None = None,
) -> str:
    """Run exactly one child process, never through shell=True."""
    global _ACTIVE_SUBPROCESS
    if _ACTIVE_SUBPROCESS is not None:
        raise SafetyError("ตรวจพบความพยายามเปิด subprocess พร้อมกันมากกว่าหนึ่งงาน")
    creationflags = 0
    if os.name == "nt":
        creationflags = subprocess.CREATE_NEW_PROCESS_GROUP | subprocess.CREATE_NO_WINDOW
    input_handle = None
    try:
        if stdin_path is not None:
            input_handle = ensure_file(stdin_path, "JSON parameter file").open("r", encoding="utf-8")
        proc = subprocess.Popen(
            arguments,
            stdin=input_handle,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
            shell=False,
            creationflags=creationflags,
        )
        _ACTIVE_SUBPROCESS = proc
        finish_by = time.monotonic() + max(1.0, timeout_seconds)
        while True:
            remaining = finish_by - time.monotonic()
            if remaining <= 0:
                stop_process(proc)
                raise AreaTimeout(
                    f"subprocess เกินเวลา {timeout_seconds / 60:.1f} นาที"
                )
            try:
                # Poll in short intervals so the memory safety floor and
                # Ctrl+C remain effective during a long QNEAT3 calculation.
                output, _ = proc.communicate(timeout=min(5.0, remaining))
                break
            except KeyboardInterrupt:
                stop_process(proc)
                raise
            except subprocess.TimeoutExpired:
                if (
                    minimum_available_memory_bytes is not None
                    and available_memory_bytes() < minimum_available_memory_bytes
                ):
                    stop_process(proc)
                    raise SafetyError(
                        "RAM ว่างต่ำกว่า safety floor ระหว่าง subprocess; "
                        "หยุดงานก่อนเผยแพร่ผล"
                    )
        if log_path is not None:
            log_path.parent.mkdir(parents=True, exist_ok=True)
            log_path.write_text(output, encoding="utf-8", errors="replace")
        if proc.returncode != 0:
            excerpt = "\n".join(output.splitlines()[-25:])
            raise PipelineError(
                f"คำสั่งภายนอกจบด้วยรหัส {proc.returncode}\n{excerpt}"
            )
        return output
    finally:
        _ACTIVE_SUBPROCESS = None
        if input_handle is not None:
            input_handle.close()


def qgis_paths(machine: dict[str, Any]) -> tuple[Path, Path]:
    qgis = machine["qgis"]
    python_bat = ensure_file(Path(qgis["python_bat"]), "QGIS Python batch")
    process_bat = ensure_file(Path(qgis["process_bat"]), "qgis_process batch")
    return python_bat, process_bat


def check_qgis_installation(machine: dict[str, Any]) -> dict[str, Any]:
    _, process_bat = qgis_paths(machine)
    timeout = float(machine["limits"]["status_timeout_seconds"])
    version_text = run_external([str(process_bat), "--version"], timeout)
    expected_version = str(machine["qgis"]["version"])
    if expected_version not in version_text:
        raise SafetyError(
            f"QGIS version ไม่ตรง: ต้องการ {expected_version}; ได้ {version_text.strip()}"
        )
    plugin_text = run_external([str(process_bat), "plugins"], timeout)
    plugin_name = str(machine["qgis"]["required_plugin"])
    loaded_pattern = re.compile(rf"^\*\s+{re.escape(plugin_name)}\s*$", re.MULTILINE)
    if not loaded_pattern.search(plugin_text):
        raise SafetyError(f"QGIS plugin {plugin_name} ไม่ได้ติดตั้ง/เปิดใช้งาน")
    algorithm = str(machine["qgis"]["required_algorithm"])
    help_text = run_external([str(process_bat), "help", algorithm], timeout)
    if algorithm not in help_text:
        raise SafetyError(f"ไม่พบ QNEAT3 algorithm: {algorithm}")
    return {
        "version": expected_version,
        "plugin": plugin_name,
        "algorithm": algorithm,
    }


def exact_external_paths(machine: dict[str, Any]) -> dict[str, Path]:
    result: dict[str, Path] = {}
    for key in (
        "education_boundaries",
        "roads",
        "admin1",
        "admin2",
        "admin3",
        "admin_lines",
    ):
        raw = machine["external_gis"].get(key)
        if not isinstance(raw, str) or not raw:
            raise PipelineError(f"machine.json ไม่มี external_gis.{key}")
        result[key] = ensure_file(Path(raw), f"external_gis.{key}")
    return result


def status_report(
    machine: dict[str, Any],
    catalog: dict[str, Any],
    deep_fingerprints: bool = True,
) -> dict[str, Any]:
    limits = machine["limits"]
    report: dict[str, Any] = {
        "checked_at": local_now(),
        "workspace": str(WORKSPACE),
        "errors": [],
        "warnings": [],
        "checks": {},
    }

    memory = available_memory_bytes()
    disk = shutil.disk_usage(WORKSPACE).free
    workspace_bytes = workspace_size_bytes()
    report["checks"]["resources"] = {
        "available_ram_bytes": memory,
        "free_disk_bytes": disk,
        "workspace_bytes": workspace_bytes,
    }
    if memory < float(limits["minimum_available_ram_gb"]) * GIB:
        report["errors"].append(
            f"RAM ว่าง {human_gib(memory)} ต่ำกว่า {limits['minimum_available_ram_gb']} GB"
        )
    if disk < float(limits["minimum_free_disk_gb"]) * GIB:
        report["errors"].append(
            f"ดิสก์ว่าง {human_gib(disk)} ต่ำกว่า {limits['minimum_free_disk_gb']} GB"
        )
    if workspace_bytes > float(limits["workspace_warning_gb"]) * GIB:
        report["warnings"].append(
            f"workspace มีขนาด {human_gib(workspace_bytes)} เกิน {limits['workspace_warning_gb']} GB; ระบบจะไม่ลบข้อมูลให้"
        )

    processes = qgis_processes()
    report["checks"]["processes"] = processes
    if processes["gui"]:
        report["errors"].append(
            "QGIS GUI เปิดอยู่: "
            + ", ".join(f"{name} PID {pid}" for pid, name in processes["gui"])
        )
    if processes["workers"]:
        report["errors"].append(
            "พบ qgis_process งานเดิม: "
            + ", ".join(f"{name} PID {pid}" for pid, name in processes["workers"])
        )

    lock_info = current_lock_info()
    report["checks"]["lock"] = lock_info
    if lock_info and int(lock_info.get("pid", -1) or -1) != os.getpid():
        pid = int(lock_info.get("pid", -1) or -1)
        if pid_is_running(pid):
            report["errors"].append(f"มี pipeline lock โดย PID {pid}")
        else:
            report["warnings"].append("พบ lock เก่า; งานถัดไปจะเก็บเป็น stale lock")

    try:
        catalog_check = validate_catalog(catalog, verify_hashes=deep_fingerprints)
        report["checks"]["catalog"] = {
            key: value for key, value in catalog_check.items() if key != "by_code"
        }
    except PipelineError as exc:
        report["errors"].append(str(exc))

    try:
        external = exact_external_paths(machine)
        external_check: dict[str, Any] = {}
        for key, path in external.items():
            if deep_fingerprints and key in {"education_boundaries", "roads"}:
                external_check[key] = shapefile_fingerprint(path)
            else:
                external_check[key] = {"path": str(path), "size": path.stat().st_size}
        report["checks"]["external_gis"] = external_check
    except PipelineError as exc:
        report["errors"].append(str(exc))

    try:
        qgis_paths(machine)
        if not processes["gui"] and not processes["workers"]:
            report["checks"]["qgis"] = check_qgis_installation(machine)
        else:
            report["warnings"].append(
                "ข้ามการเรียก qgis_process เพื่อตรวจรุ่น เพราะมี QGIS/งาน QGIS เปิดอยู่"
            )
    except PipelineError as exc:
        report["errors"].append(str(exc))

    report["ready"] = not report["errors"]
    return report


def print_status(report: dict[str, Any]) -> None:
    resources = report["checks"].get("resources", {})
    catalog = report["checks"].get("catalog", {})
    qgis = report["checks"].get("qgis", {})
    print("\n" + "=" * 72)
    print("ตรวจสถานะข้อมูลและเครื่อง")
    print("=" * 72)
    print(f"Workspace: {report['workspace']}")
    if resources:
        print(f"RAM ว่าง:      {human_gib(resources['available_ram_bytes'])}")
        print(f"ดิสก์ว่าง:    {human_gib(resources['free_disk_bytes'])}")
        print(f"ขนาด workspace: {human_gib(resources['workspace_bytes'])}")
    if catalog:
        print(
            f"บัญชีพื้นที่:  {catalog.get('area_count')} พื้นที่ / "
            f"{catalog.get('province_count')} จังหวัด"
        )
        print(f"สถานะพื้นที่:  {catalog.get('status_counts')}")
        print(f"DMC fingerprint: {catalog.get('source_sha256')}")
    if qgis:
        print(f"QGIS:          {qgis.get('version')} / {qgis.get('plugin')} พร้อมใช้")
    for warning in report.get("warnings", []):
        print(f"[คำเตือน] {warning}")
    for error in report.get("errors", []):
        print(f"[ไม่ผ่าน] {error}")
    print("พร้อมคำนวณ" if report.get("ready") else "ยังไม่พร้อมคำนวณ")


def require_ready_status(machine: dict[str, Any], catalog: dict[str, Any]) -> dict[str, Any]:
    report = status_report(machine, catalog, deep_fingerprints=True)
    print_status(report)
    if not report["ready"]:
        raise SafetyError("preflight ไม่ผ่าน; ยังไม่เริ่มคำนวณและไม่ได้แก้ผลเดิม")
    return report


def qgis_parameter_file(run_dir: Path, sequence: int, algorithm: str, inputs: dict[str, Any]) -> Path:
    safe_algorithm = re.sub(r"[^A-Za-z0-9_-]+", "_", algorithm)
    path = run_dir / f"{sequence:02d}_{safe_algorithm}.parameters.json"
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump({"inputs": inputs}, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        handle.flush()
        os.fsync(handle.fileno())
    return path


def run_qgis_algorithm(
    machine: dict[str, Any],
    run_dir: Path,
    sequence: int,
    algorithm: str,
    inputs: dict[str, Any],
    deadline: float,
    qneat: bool = False,
) -> dict[str, Any]:
    processes = qgis_processes()
    if processes["gui"]:
        raise SafetyError("QGIS ถูกเปิดระหว่างงาน; หยุดก่อนเขียนผล")
    if processes["workers"]:
        raise SafetyError("พบ qgis_process อื่นระหว่างงาน; หยุดเพื่อคง concurrency = 1")
    _, process_bat = qgis_paths(machine)
    parameter_path = qgis_parameter_file(run_dir, sequence, algorithm, inputs)
    log_path = run_dir / f"{sequence:02d}_{re.sub(r'[^A-Za-z0-9_-]+', '_', algorithm)}.log"
    configured = float(machine["limits"]["qgis_process_timeout_seconds"])
    remaining = remaining_seconds(deadline)
    timeout = min(configured if qneat else min(configured, 1800.0), float(remaining))
    print(f"  [{sequence:02d}] {algorithm}")
    output = run_external(
        [str(process_bat), "--json", "run", algorithm, "-"],
        timeout,
        stdin_path=parameter_path,
        log_path=log_path,
        minimum_available_memory_bytes=int(
            float(machine["limits"]["minimum_available_ram_gb"]) * GIB
        ),
    )
    try:
        parsed = json.loads(output)
        return parsed if isinstance(parsed, dict) else {"raw": output}
    except json.JSONDecodeError:
        # Some providers write a warning before the JSON object.  The complete
        # text remains in the exact run log for audit.
        start = output.find("{")
        if start >= 0:
            try:
                parsed = json.loads(output[start:])
                return parsed if isinstance(parsed, dict) else {"raw": output}
            except json.JSONDecodeError:
                pass
        return {"raw": output}


def normalize_school_id(value: Any, label: str) -> str:
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if not re.fullmatch(r"\d+", text):
        raise QualityError(f"{label} ไม่ใช่รหัสตัวเลข: {value!r}")
    return text


def number(value: str, label: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise QualityError(f"{label} ไม่ใช่ตัวเลข: {value!r}") from exc
    if not (-1.0e12 < result < 1.0e12):
        raise QualityError(f"{label} อยู่นอกช่วงที่รองรับ: {result}")
    return result


def load_area_records(area: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], set[str]]:
    code = area["areacode"]
    name = area["name"]
    small: list[dict[str, Any]] = []
    magnets: list[dict[str, Any]] = []

    for row in csv_rows(SMALL_CSV):
        if row.get("ชื่อเขต", "").strip() != name:
            continue
        if row.get("areacode", "").strip() != code:
            raise QualityError(f"areacode ใน DMC small ไม่ตรงสำหรับ {name}")
        school_id = normalize_school_id(row.get("โรงเรียน"), "รหัสโรงเรียนต้นทาง")
        latitude = number(row.get("lat", ""), f"lat โรงเรียน {school_id}")
        longitude = number(row.get("long", ""), f"long โรงเรียน {school_id}")
        if not (5 <= latitude <= 21 and 97 <= longitude <= 106):
            raise QualityError(f"พิกัดโรงเรียน {school_id} อยู่นอกประเทศไทยโดยประมาณ")
        students_text = row.get("รวมนักเรียน", "").strip()
        students = int(float(students_text)) if students_text else 0
        small.append(
            {
                "school_id": school_id,
                "school_name": row.get("ชื่อโรงเรียน", "").strip(),
                "area_name": name,
                "areacode": code,
                "province": row.get("ชื่อจังหวัด", "").strip(),
                "district": row.get("ชื่ออำเภอ", "").strip(),
                "subdistrict": row.get("ชื่อตำบล", "").strip(),
                "students": students,
                "lat": latitude,
                "lon": longitude,
            }
        )

    for row in csv_rows(MAGNET_CSV):
        if row.get("spt", "").strip() != name:
            continue
        school_id = normalize_school_id(row.get("school_code"), "รหัสโรงเรียนคุณภาพชุมชน")
        latitude = number(row.get("lat", ""), f"lat โรงเรียนคุณภาพชุมชน {school_id}")
        longitude = number(row.get("lon", ""), f"lon โรงเรียนคุณภาพชุมชน {school_id}")
        if not (5 <= latitude <= 21 and 97 <= longitude <= 106):
            raise QualityError(f"พิกัดโรงเรียนคุณภาพชุมชน {school_id} อยู่นอกประเทศไทยโดยประมาณ")
        magnets.append(
            {
                "school_id": school_id,
                "school_name": row.get("school_name", "").strip(),
                "area_name": name,
                "province": row.get("province", "").strip(),
                "district": row.get("district", "").strip(),
                "subdistrict": row.get("subdistrict", "").strip(),
                "lat": latitude,
                "lon": longitude,
            }
        )

    small_ids = [row["school_id"] for row in small]
    magnet_ids = [row["school_id"] for row in magnets]
    if len(small_ids) != len(set(small_ids)):
        raise QualityError(f"รหัสโรงเรียนต้นทางซ้ำใน {name}")
    if len(magnet_ids) != len(set(magnet_ids)):
        raise QualityError(f"รหัสโรงเรียนคุณภาพชุมชนซ้ำใน {name}")
    if not small:
        raise QualityError(f"ไม่พบประชากรโรงเรียนขนาดเล็กสำหรับ {name}")
    if not magnets:
        raise QualityError(f"ไม่พบโรงเรียนคุณภาพชุมชนสำหรับ {name}")

    overlap = set(small_ids) & set(magnet_ids)
    expected = area.get("school_counts", {})
    comparisons = {
        "registered": len(small),
        "potential_origins": len(small) - len(overlap),
        "magnets": len(magnets),
        "excluded_is_magnet": len(overlap),
    }
    for key, actual in comparisons.items():
        if key in expected and int(expected[key]) != actual:
            raise QualityError(
                f"จำนวน {key} ของ {name} เปลี่ยน: ได้ {actual}, catalog={expected[key]}"
            )
    return small, magnets, overlap


def ogr_modules() -> tuple[Any, Any]:
    try:
        from osgeo import ogr, osr
    except ImportError as exc:
        raise PipelineError(
            "ไม่พบ GDAL/OGR ใน Python ปัจจุบัน; ต้องเรียกผ่าน python-qgis-ltr.bat"
        ) from exc
    ogr.UseExceptions()
    return ogr, osr


def spatial_transform() -> tuple[Any, Any]:
    _, osr = ogr_modules()
    source = osr.SpatialReference()
    source.ImportFromEPSG(4326)
    target = osr.SpatialReference()
    target.ImportFromEPSG(32647)
    if hasattr(osr, "OAMS_TRADITIONAL_GIS_ORDER"):
        source.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
        target.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    return target, osr.CoordinateTransformation(source, target)


def point_utm(lon: float, lat: float, transform: Any) -> Any:
    ogr, _ = ogr_modules()
    point = ogr.Geometry(ogr.wkbPoint)
    point.AddPoint_2D(float(lon), float(lat))
    point.Transform(transform)
    return point


def create_school_points(
    path: Path,
    layer_name: str,
    records: list[dict[str, Any]],
    magnet: bool,
) -> None:
    ensure_workspace_or_local_state(path, "ผลจุดโรงเรียน")
    ogr, _ = ogr_modules()
    target_srs, transform = spatial_transform()
    driver = ogr.GetDriverByName("GPKG")
    dataset = driver.CreateDataSource(str(path))
    if dataset is None:
        raise PipelineError(f"สร้าง GeoPackage ไม่สำเร็จ: {path}")
    layer = dataset.CreateLayer(layer_name, target_srs, ogr.wkbPoint)
    if magnet:
        fields = [
            ("school_code", ogr.OFTInteger64, 0),
            ("school_name", ogr.OFTString, 254),
            ("spt", ogr.OFTString, 120),
            ("province", ogr.OFTString, 80),
            ("district", ogr.OFTString, 120),
            ("subdistrict", ogr.OFTString, 120),
            ("lat", ogr.OFTReal, 0),
            ("lon", ogr.OFTReal, 0),
        ]
    else:
        fields = [
            ("โรงเรียน", ogr.OFTInteger64, 0),
            ("ชื่อโรงเรียน", ogr.OFTString, 254),
            ("ชื่อเขต", ogr.OFTString, 120),
            ("ชื่อจังหวัด", ogr.OFTString, 80),
            ("ชื่ออำเภอ", ogr.OFTString, 120),
            ("ชื่อตำบล", ogr.OFTString, 120),
            ("รวมนักเรียน", ogr.OFTInteger, 0),
            ("areacode", ogr.OFTString, 8),
            ("lat", ogr.OFTReal, 0),
            ("long", ogr.OFTReal, 0),
        ]
    for name, field_type, width in fields:
        definition = ogr.FieldDefn(name, field_type)
        if width:
            definition.SetWidth(width)
        if layer.CreateField(definition) != 0:
            raise PipelineError(f"สร้าง field {name} ไม่สำเร็จใน {path.name}")

    definition = layer.GetLayerDefn()
    for row in records:
        feature = ogr.Feature(definition)
        if magnet:
            feature.SetField("school_code", int(row["school_id"]))
            feature.SetField("school_name", row["school_name"])
            feature.SetField("spt", row["area_name"])
            feature.SetField("province", row["province"])
            feature.SetField("district", row["district"])
            feature.SetField("subdistrict", row["subdistrict"])
            feature.SetField("lat", row["lat"])
            feature.SetField("lon", row["lon"])
        else:
            feature.SetField("โรงเรียน", int(row["school_id"]))
            feature.SetField("ชื่อโรงเรียน", row["school_name"])
            feature.SetField("ชื่อเขต", row["area_name"])
            feature.SetField("ชื่อจังหวัด", row["province"])
            feature.SetField("ชื่ออำเภอ", row["district"])
            feature.SetField("ชื่อตำบล", row["subdistrict"])
            feature.SetField("รวมนักเรียน", row["students"])
            feature.SetField("areacode", row["areacode"])
            feature.SetField("lat", row["lat"])
            feature.SetField("long", row["lon"])
        feature.SetGeometry(point_utm(row["lon"], row["lat"], transform))
        if layer.CreateFeature(feature) != 0:
            raise PipelineError(f"เขียนจุดโรงเรียน {row['school_id']} ไม่สำเร็จ")
        feature = None
    layer.SyncToDisk()
    dataset = None


def ogr_layer_metadata(path: Path) -> dict[str, Any]:
    ogr, _ = ogr_modules()
    exact = ensure_file(path, "GeoPackage")
    dataset = ogr.Open(str(exact), 0)
    if dataset is None or dataset.GetLayerCount() < 1:
        raise QualityError(f"เปิด layer ไม่ได้: {exact}")
    layer = dataset.GetLayer(0)
    fields = [
        layer.GetLayerDefn().GetFieldDefn(index).GetName()
        for index in range(layer.GetLayerDefn().GetFieldCount())
    ]
    spatial_ref = layer.GetSpatialRef()
    authority = None
    if spatial_ref:
        authority_name = spatial_ref.GetAuthorityName(None)
        authority_code = spatial_ref.GetAuthorityCode(None)
        if authority_name and authority_code:
            authority = f"{authority_name}:{authority_code}"
    result = {
        "path": str(exact),
        "layer_name": layer.GetName(),
        "feature_count": int(layer.GetFeatureCount()),
        "geometry_type": int(layer.GetGeomType()),
        "crs": authority,
        "fields": fields,
    }
    dataset = None
    return result


def validate_boundary_override_file(
    path: Path,
    metadata: dict[str, Any],
    configured_areas: list[dict[str, Any]],
) -> dict[str, Any]:
    """Validate the small, versioned override without reading external GIS folders."""
    ogr, _ = ogr_modules()
    exact = ensure_file(path, "boundary override")
    dataset = ogr.Open(str(exact), 0)
    if dataset is None or dataset.GetLayerCount() != 1:
        raise QualityError("boundary override ต้องมี layer เดียว")
    layer = dataset.GetLayer(0)
    layer_name = layer.GetName()
    if layer_name != str(metadata.get("layer_name", "")):
        raise QualityError(
            f"ชื่อ layer boundary override เปลี่ยน: ได้ {layer_name!r}, "
            f"catalog={metadata.get('layer_name')!r}"
        )
    feature_count = int(layer.GetFeatureCount())
    if feature_count != int(metadata.get("feature_count", -1)):
        raise QualityError(
            f"จำนวน feature boundary override เปลี่ยน: ได้ {feature_count}, "
            f"catalog={metadata.get('feature_count')}"
        )
    spatial_ref = layer.GetSpatialRef()
    authority = None
    if spatial_ref:
        authority_name = spatial_ref.GetAuthorityName(None)
        authority_code = spatial_ref.GetAuthorityCode(None)
        if authority_name and authority_code:
            authority = f"{authority_name}:{authority_code}"
    if authority != str(metadata.get("crs", "")):
        raise QualityError(
            f"CRS boundary override เปลี่ยน: ได้ {authority}, catalog={metadata.get('crs')}"
        )
    if ogr.GT_Flatten(layer.GetGeomType()) != ogr.wkbMultiPolygon:
        raise QualityError("boundary override ต้องเป็น MultiPolygon")

    definition = layer.GetLayerDefn()
    fields = {
        definition.GetFieldDefn(index).GetName()
        for index in range(definition.GetFieldCount())
    }
    required_fields = {
        "areacode",
        BOUNDARY_NAME_FIELD,
        "province",
        "district_count",
        "district_codes",
        "source_ref",
    }
    if not required_fields.issubset(fields):
        raise QualityError(
            "boundary override ขาด field: "
            + ", ".join(sorted(required_fields - fields))
        )

    by_code = {str(item["areacode"]): item for item in configured_areas}
    expected_details = metadata.get("expected_areas")
    if not isinstance(expected_details, dict) or set(expected_details) != set(by_code):
        raise QualityError("boundary_override_source.expected_areas ไม่ตรงพื้นที่ที่ใช้งาน")
    seen: set[str] = set()
    geometries: list[Any] = []
    for feature in layer:
        code = str(feature.GetField("areacode")).strip()
        if code not in by_code or code in seen:
            raise QualityError(f"boundary override มี areacode ผิดหรือซ้ำ: {code!r}")
        area = by_code[code]
        if str(feature.GetField(BOUNDARY_NAME_FIELD)).strip() != area["boundary_name"]:
            raise QualityError(f"boundary override มีชื่อเขตไม่ตรงสำหรับ {code}")
        if str(feature.GetField("province")).strip() != area["province"]:
            raise QualityError(f"boundary override มีจังหวัดไม่ตรงสำหรับ {code}")
        expected = expected_details[code]
        try:
            district_count = int(feature.GetField("district_count"))
        except (TypeError, ValueError) as exc:
            raise QualityError(f"district_count ไม่ถูกต้องสำหรับ {code}") from exc
        actual_codes = sorted(
            value.strip()
            for value in str(feature.GetField("district_codes")).split(",")
            if value.strip()
        )
        expected_codes = sorted(str(value) for value in expected.get("district_codes", []))
        if district_count != int(expected.get("district_count", -1)):
            raise QualityError(f"district_count เปลี่ยนสำหรับ {code}")
        if actual_codes != expected_codes or len(actual_codes) != district_count:
            raise QualityError(f"district_codes เปลี่ยนสำหรับ {code}")
        geometry = feature.GetGeometryRef()
        if geometry is None or geometry.IsEmpty() or not geometry.IsValid():
            raise QualityError(f"geometry boundary override ไม่สมบูรณ์สำหรับ {code}")
        geometries.append(geometry.Clone())
        seen.add(code)
    if seen != set(by_code):
        raise QualityError("boundary override มีพื้นที่ไม่ครบตาม areas.json")

    for left_index, left in enumerate(geometries):
        for right in geometries[left_index + 1 :]:
            intersection = left.Intersection(right)
            if intersection is not None and not intersection.IsEmpty() and intersection.GetArea() > 1e-12:
                raise QualityError("boundary override มี polygon ซ้อนทับกัน")
    dataset = None
    return {
        "path": str(exact),
        "layer_name": layer_name,
        "feature_count": feature_count,
        "crs": authority,
        "area_codes": sorted(seen),
    }


def ogr_value_counts(path: Path, field: str) -> dict[str, int]:
    ogr, _ = ogr_modules()
    dataset = ogr.Open(str(ensure_file(path, "GeoPackage สำหรับนับค่า")), 0)
    if dataset is None or dataset.GetLayerCount() < 1:
        raise QualityError(f"เปิด layer เพื่อนับค่าไม่ได้: {path}")
    layer = dataset.GetLayer(0)
    if layer.GetLayerDefn().GetFieldIndex(field) < 0:
        raise QualityError(f"ชั้น {path.name} ไม่มี field {field}")
    counts: dict[str, int] = {}
    for feature in layer:
        value = feature.GetField(field)
        key = "" if value is None else str(value).strip()
        counts[key] = counts.get(key, 0) + 1
    dataset = None
    return counts


def load_boundary_geometry(path: Path) -> Any:
    ogr, _ = ogr_modules()
    dataset = ogr.Open(str(ensure_file(path, "ขอบเขตพื้นที่")), 0)
    if dataset is None:
        raise QualityError(f"เปิดขอบเขตไม่ได้: {path}")
    layer = dataset.GetLayer(0)
    combined = None
    for feature in layer:
        geometry = feature.GetGeometryRef()
        if geometry is None or geometry.IsEmpty():
            continue
        combined = geometry.Clone() if combined is None else combined.Union(geometry)
    dataset = None
    if combined is None or combined.IsEmpty():
        raise QualityError(f"ขอบเขตไม่มี geometry: {path}")
    return combined


def output_names(area: dict[str, Any]) -> dict[str, str]:
    folder = area["folder_name"]
    generic = {
        "boundary": f"ขอบเขต_{folder}_UTM47N.gpkg",
        "small": f"โรงเรียนขนาดเล็ก_{folder}.gpkg",
        "magnets": f"โรงเรียนแม่เหล็ก_{folder}.gpkg",
        "roads": f"ถนน_{folder}_สำหรับคำนวณ.gpkg",
        "od": f"OD_{folder}_ไม่เกิน10กม.gpkg",
        "status": f"สถานะการจับคู่_โรงเรียนขนาดเล็ก_{folder}.gpkg",
        "image": f"แผนที่ผลการจับคู่_{folder}.png",
        "pdf": f"สรุปผลการจับคู่_{folder}.pdf",
        "size_image": f"แผนที่ขนาดโรงเรียน_{folder}.png",
        "qc": f"ผลตรวจ_{area['areacode']}.json",
    }
    pilot_files = area.get("pilot", {}).get("files", {})
    for key in ("boundary", "small", "magnets", "roads", "od", "status"):
        if key in pilot_files:
            generic[key] = pilot_files[key]
    for key, name in generic.items():
        if Path(name).name != name or any(char in name for char in '<>:"/\\|?*'):
            raise SafetyError(f"ชื่อผลลัพธ์ไม่ปลอดภัย ({key}): {name}")
    return generic


def make_partial(final: Path, token: str) -> Path:
    return final.with_name(f"{final.stem}.partial.{token}{final.suffix}")


def output_paths(area: dict[str, Any], token: str) -> tuple[Path, dict[str, Path], dict[str, Path]]:
    area_dir = ensure_within(AREA_ROOT / area["folder_name"], AREA_ROOT, "โฟลเดอร์รายพื้นที่")
    names = output_names(area)
    # A pending area does not have a directory until its first permanent
    # result is ready.  Validate against AREA_ROOT here; validating against a
    # non-existent area_dir would make Path.resolve(strict=True) fail before
    # run_area gets a chance to create that directory.
    final = {
        key: ensure_within(area_dir / layout.area_relpath(key, name), AREA_ROOT, f"ผล {key}")
        for key, name in names.items()
    }
    partial = {key: make_partial(path, token) for key, path in final.items()}
    return area_dir, final, partial


def ensure_area_subdirs(area_dir: Path) -> None:
    """สร้างโฟลเดอร์ย่อยมาตรฐานทั้งสามชุดเสมอ เพื่อให้ทุกพื้นที่มี pattern เดียวกัน"""
    for sub in layout.SUBDIRS:
        ensure_within(area_dir / sub, AREA_ROOT, "โฟลเดอร์ย่อยรายพื้นที่").mkdir(
            parents=True, exist_ok=True
        )


def remove_empty_area_subdirs(area_dir: Path) -> None:
    """ลบเฉพาะโฟลเดอร์ย่อยที่ว่างจริง ใช้ตอนงานล้มก่อนมีผลถาวร (Project Tor.md หัวข้อ 12; บันทึกการเปลี่ยนแปลง.md §14.0)"""
    for sub in layout.SUBDIRS:
        try:
            (area_dir / sub).rmdir()
        except OSError:
            pass


def file_is_locked(path: Path) -> bool:
    if not path.exists() or os.name != "nt":
        return False
    GENERIC_READ = 0x80000000
    GENERIC_WRITE = 0x40000000
    OPEN_EXISTING = 3
    FILE_ATTRIBUTE_NORMAL = 0x80
    handle = ctypes.windll.kernel32.CreateFileW(
        str(path),
        GENERIC_READ | GENERIC_WRITE,
        0,
        None,
        OPEN_EXISTING,
        FILE_ATTRIBUTE_NORMAL,
        None,
    )
    if handle == INVALID_HANDLE_VALUE:
        return True
    ctypes.windll.kernel32.CloseHandle(handle)
    return False


def preflight_outputs(final: dict[str, Path], overwrite: bool) -> None:
    existing = [path for path in final.values() if path.exists()]
    if existing and not overwrite:
        listed = "\n".join(f"  - {path.name}" for path in existing)
        raise SafetyError(
            "พบผลเดิมและระบบไม่เขียนทับโดยค่าเริ่มต้น:\n"
            f"{listed}\nใช้การยืนยัน overwrite แยกต่างหากหากต้องการสร้างรุ่นใหม่"
        )
    if overwrite:
        locked = [path for path in existing if file_is_locked(path)]
        if locked:
            raise SafetyError(
                "ไฟล์เป้าหมายถูกล็อก:\n" + "\n".join(f"  - {path}" for path in locked)
            )


def transactional_promote(
    partial: dict[str, Path],
    final: dict[str, Path],
    overwrite: bool,
    token: str,
) -> None:
    preflight_outputs(final, overwrite)
    missing = [path for path in partial.values() if not path.is_file()]
    if missing:
        raise QualityError(
            "ไฟล์ .partial ไม่ครบก่อน publish:\n"
            + "\n".join(f"  - {path}" for path in missing)
        )
    backups: dict[str, Path] = {}
    promoted: list[str] = []
    try:
        if overwrite:
            for key, target in final.items():
                if target.exists():
                    backup = target.with_name(f"{target.stem}.backup.{token}{target.suffix}")
                    os.replace(target, backup)
                    backups[key] = backup
        for key, staging in partial.items():
            target = final[key]
            if target.exists():
                raise SafetyError(f"เป้าหมายปรากฏขึ้นระหว่าง publish: {target}")
            os.replace(staging, target)
            promoted.append(key)
    except BaseException:
        for key in reversed(promoted):
            target = final[key]
            try:
                if target.exists():
                    target.unlink()
            except OSError:
                pass
        for key, backup in backups.items():
            try:
                if backup.exists() and not final[key].exists():
                    os.replace(backup, final[key])
            except OSError:
                pass
        raise
    else:
        for backup in backups.values():
            try:
                backup.unlink()
            except OSError:
                print(f"[คำเตือน] ลบ backup หลัง publish ไม่ได้: {backup}")


def qgis_sink(path: Path, layer_name: str) -> str:
    # qgis_process JSON output parameters expect a filesystem destination.
    # A provider URI such as file.gpkg|layername=... is treated as a literal
    # database filename and QGIS appends another .gpkg, causing creation to
    # fail.  Every generated GPKG here contains one layer, so the provider's
    # default layer name is sufficient.
    del layer_name
    return str(path)


def prepare_boundary_and_roads(
    machine: dict[str, Any],
    area: dict[str, Any],
    partial: dict[str, Path],
    run_dir: Path,
    deadline: float,
) -> dict[str, Any]:
    external = exact_external_paths(machine)
    boundary_source, boundary_source_kind = resolve_boundary_source(machine, area)
    boundary_source_fingerprint = vector_source_fingerprint(boundary_source)
    buffer_path = run_dir / "boundary_buffer_15km.gpkg"
    extracted_path = run_dir / "boundary_selected_wgs84.gpkg"
    clipped_path = run_dir / "roads_clipped_first.gpkg"
    accepted_path = run_dir / "roads_vehicle_classes.gpkg"
    fixed_path = run_dir / "roads_fixed.gpkg"

    run_qgis_algorithm(
        machine,
        run_dir,
        1,
        "native:extractbyattribute",
        {
            "INPUT": str(boundary_source),
            "FIELD": BOUNDARY_NAME_FIELD,
            "OPERATOR": 0,
            "VALUE": area["boundary_name"],
            "OUTPUT": qgis_sink(extracted_path, "boundary_selected"),
        },
        deadline,
    )
    extracted_meta = ogr_layer_metadata(extracted_path)
    if extracted_meta["feature_count"] != 1:
        raise QualityError(
            f"ขอบเขต {area['name']} ได้ {extracted_meta['feature_count']} feature; ต้องเป็น 1"
        )

    run_qgis_algorithm(
        machine,
        run_dir,
        2,
        "native:reprojectlayer",
        {
            "INPUT": str(extracted_path),
            "TARGET_CRS": machine["analysis"]["crs"],
            "CONVERT_CURVED_GEOMETRIES": False,
            "OUTPUT": qgis_sink(partial["boundary"], "boundary_area"),
        },
        deadline,
    )
    boundary_meta = ogr_layer_metadata(partial["boundary"])
    if boundary_meta["feature_count"] != 1 or boundary_meta["crs"] != "EPSG:32647":
        raise QualityError(f"ขอบเขตหลัง reproject ไม่ผ่าน QC: {boundary_meta}")

    run_qgis_algorithm(
        machine,
        run_dir,
        3,
        "native:buffer",
        {
            "INPUT": str(partial["boundary"]),
            "DISTANCE": float(machine["analysis"]["buffer_metres"]),
            "SEGMENTS": 8,
            "END_CAP_STYLE": 0,
            "JOIN_STYLE": 0,
            "MITER_LIMIT": 2,
            "DISSOLVE": True,
            "SEPARATE_DISJOINT": False,
            "OUTPUT": qgis_sink(buffer_path, "boundary_buffer"),
        },
        deadline,
    )

    # GDAL clips the exact nationwide roads file before any in-memory class
    # filtering.  No parent directory enumeration is performed.
    run_qgis_algorithm(
        machine,
        run_dir,
        4,
        "gdal:clipvectorbypolygon",
        {
            "INPUT": str(external["roads"]),
            "MASK": str(buffer_path),
            "OPTIONS": "",
            "OUTPUT": str(clipped_path),
        },
        deadline,
    )
    classes = machine["analysis"]["accepted_road_classes"]
    if not isinstance(classes, list) or not classes:
        raise PipelineError("accepted_road_classes ต้องเป็น list ที่ไม่ว่าง")
    if any(not re.fullmatch(r"[a-z_]+", str(value)) for value in classes):
        raise SafetyError("accepted_road_classes มีค่าที่ไม่ปลอดภัย")
    quoted = ", ".join("'" + str(value) + "'" for value in classes)
    expression = f'"{ROAD_EXPRESSION_FIELD}" IN ({quoted})'
    run_qgis_algorithm(
        machine,
        run_dir,
        5,
        "native:extractbyexpression",
        {
            "INPUT": str(clipped_path),
            "EXPRESSION": expression,
            "OUTPUT": qgis_sink(accepted_path, "roads_vehicle"),
        },
        deadline,
    )
    run_qgis_algorithm(
        machine,
        run_dir,
        6,
        "native:fixgeometries",
        {
            "INPUT": str(accepted_path),
            "METHOD": 1,
            "OUTPUT": qgis_sink(fixed_path, "roads_fixed"),
        },
        deadline,
    )
    run_qgis_algorithm(
        machine,
        run_dir,
        7,
        "native:multiparttosingleparts",
        {
            "INPUT": str(fixed_path),
            "OUTPUT": qgis_sink(partial["roads"], "roads_for_routing"),
        },
        deadline,
    )
    road_meta = ogr_layer_metadata(partial["roads"])
    required_fields = {"fclass", "oneway"}
    if road_meta["feature_count"] <= 0 or not required_fields.issubset(road_meta["fields"]):
        raise QualityError(f"ถนนหลังเตรียมข้อมูลไม่ผ่าน QC: {road_meta}")
    if road_meta["crs"] != "EPSG:32647":
        raise QualityError(f"ถนนต้องเป็น EPSG:32647; ได้ {road_meta['crs']}")
    oneway_counts = ogr_value_counts(partial["roads"], "oneway")
    unexpected_oneway = sorted(set(oneway_counts) - {"", "F", "T", "B"})
    if unexpected_oneway:
        raise QualityError(
            "ถนนมีค่า oneway นอกโดเมน F/T/B/ว่าง: " + ", ".join(unexpected_oneway)
        )
    return {
        "boundary": boundary_meta,
        "roads": road_meta,
        "oneway_counts": oneway_counts,
        "boundary_source": {
            "path": str(boundary_source),
            "kind": boundary_source_kind,
            "fingerprint": boundary_source_fingerprint,
        },
    }


def run_qneat3(
    machine: dict[str, Any],
    partial: dict[str, Path],
    run_dir: Path,
    deadline: float,
) -> Path:
    full_od = run_dir / "od_full_matrix.gpkg"
    run_qgis_algorithm(
        machine,
        run_dir,
        8,
        str(machine["qgis"]["required_algorithm"]),
        {
            "INPUT": str(partial["roads"]),
            "FROM_POINT_LAYER": str(partial["small"]),
            "FROM_ID_FIELD": ORIGIN_ID_FIELD,
            "TO_POINT_LAYER": str(partial["magnets"]),
            "TO_ID_FIELD": DESTINATION_ID_FIELD,
            "STRATEGY": 0,
            "ENTRY_COST_CALCULATION_METHOD": 1,
            "MATRIX_GEOMETRY_TYPE": 0,
            "DIRECTION_FIELD": "oneway",
            "VALUE_FORWARD": "F",
            "VALUE_BACKWARD": "T",
            "VALUE_BOTH": "B",
            "DEFAULT_DIRECTION": 2,
            "SPEED_FIELD": None,
            "DEFAULT_SPEED": 5,
            "TOLERANCE": 0,
            "OUTPUT": qgis_sink(full_od, "od_full_matrix"),
        },
        deadline,
        qneat=True,
    )
    metadata = ogr_layer_metadata(full_od)
    required = {"origin_id", "destination_id", "total_cost"}
    if not required.issubset(metadata["fields"]):
        raise QualityError(f"Full OD ไม่มี field ที่กำหนด: {required - set(metadata['fields'])}")
    return full_od


def _create_ogr_field(layer: Any, name: str, field_type: int, width: int = 0) -> None:
    ogr, _ = ogr_modules()
    definition = ogr.FieldDefn(name, field_type)
    if width:
        definition.SetWidth(width)
    if layer.CreateField(definition) != 0:
        raise PipelineError(f"สร้าง field {name} ไม่สำเร็จใน {layer.GetName()}")


def analyse_od_and_write_outputs(
    machine: dict[str, Any],
    full_od_path: Path,
    boundary_path: Path,
    small_all: list[dict[str, Any]],
    analysis_origins: list[dict[str, Any]],
    magnets: list[dict[str, Any]],
    overlap: set[str],
    od_output: Path,
    status_output: Path,
) -> dict[str, Any]:
    """Choose one deterministic nearest magnet and create the two final GIS layers."""
    ogr, _ = ogr_modules()
    dataset = ogr.Open(str(ensure_file(full_od_path, "Full OD")), 0)
    if dataset is None or dataset.GetLayerCount() != 1:
        raise QualityError(f"เปิด Full OD ไม่ได้: {full_od_path}")
    source_layer = dataset.GetLayer(0)
    source_srs = source_layer.GetSpatialRef()

    origin_ids = {row["school_id"] for row in analysis_origins}
    destination_ids = {row["school_id"] for row in magnets}
    rows_by_origin: dict[str, list[dict[str, Any]]] = {value: [] for value in origin_ids}
    seen_pairs: set[tuple[str, str]] = set()
    full_rows = 0
    valid_rows = 0

    for feature in source_layer:
        full_rows += 1
        origin = normalize_school_id(feature.GetField("origin_id"), "origin_id ใน Full OD")
        destination = normalize_school_id(
            feature.GetField("destination_id"), "destination_id ใน Full OD"
        )
        if origin not in origin_ids:
            raise QualityError(f"Full OD มี origin_id นอกประชากรพื้นที่: {origin}")
        if destination not in destination_ids:
            raise QualityError(f"Full OD มี destination_id นอกโรงเรียนคุณภาพชุมชน: {destination}")
        pair = (origin, destination)
        if pair in seen_pairs:
            raise QualityError(f"Full OD มีคู่ซ้ำ: {origin} → {destination}")
        seen_pairs.add(pair)
        raw_cost = feature.GetField("total_cost")
        try:
            cost = float(raw_cost)
        except (TypeError, ValueError):
            continue
        if not math.isfinite(cost) or cost < 0:
            continue
        geometry = feature.GetGeometryRef()
        valid_rows += 1
        rows_by_origin[origin].append(
            {
                "origin_id": origin,
                "destination_id": destination,
                "entry_cost": feature.GetField("entry_cost"),
                "network_cost": feature.GetField("network_cost"),
                "exit_cost": feature.GetField("exit_cost"),
                "total_cost": cost,
                "geometry_wkb": geometry.ExportToWkb() if geometry is not None else None,
            }
        )
    dataset = None

    tie_tolerance = float(machine["analysis"]["tie_tolerance_metres"])
    maximum = float(machine["analysis"]["maximum_distance_metres"])
    sensitivity_min = float(machine["analysis"]["sensitivity_min_metres"])
    sensitivity_max = float(machine["analysis"]["sensitivity_max_metres"])
    nearest: dict[str, dict[str, Any]] = {}
    for origin in sorted(origin_ids):
        candidates = sorted(
            rows_by_origin.get(origin, []),
            key=lambda row: (row["total_cost"], row["destination_id"]),
        )
        if not candidates:
            continue
        minimum = candidates[0]["total_cost"]
        tied_destinations = {
            row["destination_id"]
            for row in candidates
            if abs(row["total_cost"] - minimum) <= tie_tolerance
        }
        if len(tied_destinations) > 1:
            raise QualityError(
                f"พบปลายทางระยะเท่ากันภายใน {tie_tolerance} เมตรสำหรับ origin {origin}: "
                + ", ".join(sorted(tied_destinations))
            )
        nearest[origin] = candidates[0]

    driver = ogr.GetDriverByName("GPKG")
    od_dataset = driver.CreateDataSource(str(od_output))
    if od_dataset is None:
        raise PipelineError(f"สร้าง OD output ไม่สำเร็จ: {od_output}")
    od_layer = od_dataset.CreateLayer("od_matched_10km", source_srs, ogr.wkbLineString)
    _create_ogr_field(od_layer, "origin_id", ogr.OFTInteger64)
    _create_ogr_field(od_layer, "destination_id", ogr.OFTInteger64)
    for field_name in ("entry_cost", "network_cost", "exit_cost", "total_cost", "distance_km"):
        _create_ogr_field(od_layer, field_name, ogr.OFTReal)
    _create_ogr_field(od_layer, "sensitivity", ogr.OFTInteger)

    boundary = load_boundary_geometry(boundary_path)
    _, transform = spatial_transform()
    status_rows: list[dict[str, Any]] = []
    counts = {
        "registered": len(small_all),
        "origins": len(analysis_origins),
        "magnets": len(magnets),
        "matched": 0,
        "over10km": 0,
        "unreachable": 0,
        "excluded": 0,
        "coordExceptions": 0,
    }
    matched_pair_tokens: list[str] = []
    sensitivity_count = 0

    for row in small_all:
        school_id = row["school_id"]
        point = point_utm(row["lon"], row["lat"], transform)
        outside = not bool(boundary.Intersects(point))
        selected = nearest.get(school_id)
        sensitivity = bool(
            selected
            and sensitivity_min <= selected["total_cost"] <= sensitivity_max
        )
        if sensitivity:
            sensitivity_count += 1

        if school_id in overlap:
            status = "EXCLUDED_IS_MAGNET"
            counts["excluded"] += 1
        elif outside:
            status = "COORD_BOUNDARY_EXCEPTION"
            counts["coordExceptions"] += 1
        elif selected is None:
            status = "UNREACHABLE"
            counts["unreachable"] += 1
        elif selected["total_cost"] <= maximum:
            status = "MATCHED_10KM"
            counts["matched"] += 1
        else:
            status = "OVER_10KM"
            counts["over10km"] += 1

        status_rows.append(
            {
                "source": row,
                "point": point,
                "status": status,
                "selected": selected,
                "outside": outside,
                "sensitivity": sensitivity,
            }
        )

        if status == "MATCHED_10KM" and selected is not None:
            if selected["geometry_wkb"] is None:
                raise QualityError(f"คู่ผ่านเกณฑ์ไม่มี geometry: {school_id}")
            feature = ogr.Feature(od_layer.GetLayerDefn())
            feature.SetField("origin_id", int(school_id))
            feature.SetField("destination_id", int(selected["destination_id"]))
            for field_name in ("entry_cost", "network_cost", "exit_cost"):
                value = selected.get(field_name)
                if value is not None:
                    feature.SetField(field_name, float(value))
            feature.SetField("total_cost", float(selected["total_cost"]))
            feature.SetField("distance_km", float(selected["total_cost"]) / 1000.0)
            feature.SetField("sensitivity", int(sensitivity))
            feature.SetGeometry(ogr.CreateGeometryFromWkb(selected["geometry_wkb"]))
            if od_layer.CreateFeature(feature) != 0:
                raise PipelineError(f"เขียน OD คู่ {school_id} ไม่สำเร็จ")
            feature = None
            matched_pair_tokens.append(
                f"{school_id}:{selected['destination_id']}:{selected['total_cost']:.6f}"
            )
    od_layer.SyncToDisk()
    od_dataset = None

    status_dataset = driver.CreateDataSource(str(status_output))
    if status_dataset is None:
        raise PipelineError(f"สร้าง status output ไม่สำเร็จ: {status_output}")
    target_srs, _ = spatial_transform()
    status_layer = status_dataset.CreateLayer(
        "school_matching_status", target_srs, ogr.wkbPoint
    )
    _create_ogr_field(status_layer, "โรงเรียน", ogr.OFTInteger64)
    _create_ogr_field(status_layer, "ชื่อโรงเรียน", ogr.OFTString, 254)
    _create_ogr_field(status_layer, "ชื่อเขต", ogr.OFTString, 120)
    _create_ogr_field(status_layer, "ชื่อจังหวัด", ogr.OFTString, 80)
    _create_ogr_field(status_layer, "ชื่ออำเภอ", ogr.OFTString, 120)
    _create_ogr_field(status_layer, "รวมนักเรียน", ogr.OFTInteger)
    _create_ogr_field(status_layer, "สถานะ", ogr.OFTString, 40)
    _create_ogr_field(status_layer, "destination_id", ogr.OFTInteger64)
    _create_ogr_field(status_layer, "total_cost", ogr.OFTReal)
    _create_ogr_field(status_layer, "distance_km", ogr.OFTReal)
    _create_ogr_field(status_layer, "sensitivity", ogr.OFTInteger)
    _create_ogr_field(status_layer, "outside_boundary", ogr.OFTInteger)

    for item in status_rows:
        row = item["source"]
        selected = item["selected"]
        feature = ogr.Feature(status_layer.GetLayerDefn())
        feature.SetField("โรงเรียน", int(row["school_id"]))
        feature.SetField("ชื่อโรงเรียน", row["school_name"])
        feature.SetField("ชื่อเขต", row["area_name"])
        feature.SetField("ชื่อจังหวัด", row["province"])
        feature.SetField("ชื่ออำเภอ", row["district"])
        feature.SetField("รวมนักเรียน", int(row["students"]))
        feature.SetField("สถานะ", item["status"])
        if selected is not None:
            feature.SetField("destination_id", int(selected["destination_id"]))
            feature.SetField("total_cost", float(selected["total_cost"]))
            feature.SetField("distance_km", float(selected["total_cost"]) / 1000.0)
        feature.SetField("sensitivity", int(item["sensitivity"]))
        feature.SetField("outside_boundary", int(item["outside"]))
        feature.SetGeometry(item["point"])
        if status_layer.CreateFeature(feature) != 0:
            raise PipelineError(f"เขียนสถานะโรงเรียน {row['school_id']} ไม่สำเร็จ")
        feature = None
    status_layer.SyncToDisk()
    status_dataset = None

    if sum(counts[key] for key in ("matched", "over10km", "unreachable", "excluded", "coordExceptions")) != counts["registered"]:
        raise QualityError(f"ผลรวมสถานะไม่เท่าประชากร: {counts}")
    if ogr_layer_metadata(od_output)["feature_count"] != counts["matched"]:
        raise QualityError("จำนวนเส้น OD ไม่เท่าจำนวน MATCHED_10KM")
    if ogr_layer_metadata(status_output)["feature_count"] != counts["registered"]:
        raise QualityError("จำนวนแถวสถานะไม่เท่าประชากรตามทะเบียน")

    pair_digest = hashlib.sha256(
        "\n".join(sorted(matched_pair_tokens)).encode("utf-8")
    ).hexdigest().upper()
    expected_pairs = len(analysis_origins) * len(magnets)
    warnings: list[str] = []
    if full_rows != expected_pairs:
        warnings.append(
            f"Full OD มี {full_rows} แถว จากค่าคาดหมาย {expected_pairs}; "
            "ระบบจัด origin ที่ไม่มีผลที่ใช้ได้เป็น UNREACHABLE"
        )
    return {
        "summary": counts,
        "full_od_rows": full_rows,
        "full_od_valid_rows": valid_rows,
        "expected_full_od_rows": expected_pairs,
        "nearest_rows": len(nearest),
        "sensitivity_9_5_to_10_5_km": sensitivity_count,
        "tie_tolerance_metres": tie_tolerance,
        "tie_count": 0,
        "matched_pair_digest": pair_digest,
        "warnings": warnings,
    }


def cleanup_partial_outputs(partial: dict[str, Path], token: str) -> None:
    for path in partial.values():
        try:
            ensure_within(path, AREA_ROOT, "ไฟล์ .partial สำหรับ cleanup")
            if f".partial.{token}" not in path.name:
                raise SafetyError(f"ชื่อ partial ไม่มี token ที่คาดหมาย: {path}")
            if path.is_file():
                path.unlink()
        except (OSError, PipelineError) as exc:
            print(f"[คำเตือน] เก็บ partial ไม่สำเร็จ: {path} ({exc})")


def cleanup_successful_run_dir(run_dir: Path) -> None:
    root = local_state_root()
    resolved = run_dir.resolve(strict=True)
    try:
        relative = resolved.relative_to(root.resolve(strict=True))
    except ValueError as exc:
        raise SafetyError(f"run directory อยู่นอก LOCALAPPDATA ที่กำหนด: {resolved}") from exc
    if len(relative.parts) < 3 or relative.parts[0] != "runs":
        raise SafetyError(f"ปฏิเสธการลบ run directory ที่กว้างเกินไป: {resolved}")
    shutil.rmtree(resolved)


def validate_expected_summary(area: dict[str, Any], summary: dict[str, Any]) -> None:
    """Block publication when a protected pilot no longer reconciles.

    New/pending areas do not have an accepted baseline yet.  Ready pilots do,
    so every published aggregate must equal areas.json exactly; a warning is
    not sufficient for those areas.
    """
    expected = area.get("expected")
    if not isinstance(expected, dict):
        expected = area.get("pilot", {}).get("expected")
    if not isinstance(expected, dict):
        return
    differences: list[str] = []
    for key, expected_value in expected.items():
        if key not in summary:
            differences.append(f"{key}: ไม่มีค่า (คาด {expected_value})")
            continue
        try:
            actual_value = int(summary[key])
            target_value = int(expected_value)
        except (TypeError, ValueError) as exc:
            raise QualityError(f"ค่าตรวจรับ {key} ไม่ใช่จำนวนเต็ม") from exc
        if actual_value != target_value:
            differences.append(f"{key}: ได้ {actual_value}, คาด {target_value}")
    if differences:
        raise QualityError(
            f"ผล {area['name']} ไม่ตรงค่าตรวจรับ จึงไม่เผยแพร่:\n  - "
            + "\n  - ".join(differences)
        )


def validate_expected_result(area: dict[str, Any], result: dict[str, Any]) -> None:
    """Reconcile detailed accepted metrics for pipeline-generated areas."""
    expected = area.get("accepted_result")
    if not isinstance(expected, dict):
        return
    differences: list[str] = []
    for key, expected_value in expected.items():
        actual_value = result.get(key)
        if isinstance(expected_value, int):
            try:
                matches = int(actual_value) == expected_value
            except (TypeError, ValueError):
                matches = False
        else:
            matches = str(actual_value) == str(expected_value)
        if not matches:
            differences.append(f"{key}: ได้ {actual_value!r}, คาด {expected_value!r}")
    if differences:
        raise QualityError(
            f"รายละเอียดผล {area['name']} ไม่ตรงค่าตรวจรับ:\n  - "
            + "\n  - ".join(differences)
        )


def build_qc_payload(
    machine: dict[str, Any],
    catalog: dict[str, Any],
    area: dict[str, Any],
    result: dict[str, Any],
    partial: dict[str, Path],
    preflight: dict[str, Any],
    size_profile: dict[str, Any] | None = None,
    presentation: dict[str, Any] | None = None,
    consolidation: dict[str, Any] | None = None,
) -> dict[str, Any]:
    external = exact_external_paths(machine)
    preparation = result.get("preparation", {})
    boundary_source = preparation.get("boundary_source", {})
    boundary_path_raw = boundary_source.get("path")
    if not isinstance(boundary_path_raw, str) or not boundary_path_raw:
        raise QualityError("ผลเตรียมข้อมูลไม่มี boundary source")
    current_boundary_fingerprint = vector_source_fingerprint(Path(boundary_path_raw))
    if current_boundary_fingerprint != boundary_source.get("fingerprint"):
        raise QualityError("boundary source เปลี่ยนระหว่างคำนวณ; หยุดก่อนเผยแพร่")
    output_meta: dict[str, Any] = {}
    for key in ("boundary", "small", "magnets", "roads", "od", "status", "image", "pdf", "size_image"):
        path = ensure_file(partial[key], f"ผล {key} ก่อน publish")
        output_meta[key] = {
            "filename": output_names(area)[key],
            "bytes": path.stat().st_size,
            "sha256": sha256_file(path),
        }
    return {
        "schema_version": 1,
        "status": "PASSED",
        "publication_label": "ฉบับนำร่องเพื่อการตรวจสอบ",
        "dataset": "ชุดข้อมูล DMC682",
        "method": "แผน 1.2 — ระยะทางตามโครงข่ายถนนไม่เกิน 10 กิโลเมตร",
        "area": {
            "id": area["id"],
            "areacode": area["areacode"],
            "name": area["name"],
            "province": area["province"],
        },
        "generated_at": local_now(),
        "pipeline_version": VERSION,
        "analysis": result,
        "source_fingerprints": {
            "dmc_master_sha256": catalog["source"]["sha256"],
            "dmc_small_sha256": catalog["cross_checks"]["small_source"]["sha256"],
            "magnets_sha256": catalog["cross_checks"]["magnet_source"]["sha256"],
            "boundary_mapping_sha256": catalog["cross_checks"]["boundary_mapping_source"]["sha256"],
            "boundary_source_kind": boundary_source.get("kind"),
            "education_boundaries": current_boundary_fingerprint,
            "roads": shapefile_fingerprint(external["roads"]),
        },
        "machine_checks": {
            "qgis": preflight.get("checks", {}).get("qgis"),
            "resources": preflight.get("checks", {}).get("resources"),
        },
        "outputs": output_meta,
        "school_size": size_profile,
        "presentation": presentation,
        "consolidation": consolidation,
    }


def run_area(
    machine: dict[str, Any],
    catalog: dict[str, Any],
    area: dict[str, Any],
    confirmation: str,
    overwrite: bool = False,
) -> dict[str, Any]:
    code = area["areacode"]
    if confirmation.strip() != code:
        raise SafetyError("รหัสยืนยันไม่ตรง; ยังไม่เริ่มคำนวณ")
    if area.get("status") in {"blocked", "empty"}:
        raise SafetyError(f"พื้นที่ {area['name']} มีสถานะ {area['status']} และยังรันไม่ได้")
    if area.get("pilot", {}).get("never_rerun"):
        raise SafetyError(f"{area['name']} เป็นผล legacy ที่กำหนดห้ามรันเขียนทับ")

    preflight = require_ready_status(machine, catalog)
    token = uuid.uuid4().hex[:12]
    area_dir, final, partial = output_paths(area, token)
    preflight_outputs(final, overwrite)
    area_dir_existed = area_dir.exists()

    with pipeline_lock("run-area", code):
        # Repeat the process check after acquiring the lock so a GUI cannot be
        # opened in the gap between preflight and the first output.
        processes = qgis_processes()
        if processes["gui"] or processes["workers"]:
            raise SafetyError("พบ QGIS/qgis_process หลัง preflight; ยังไม่เริ่มเขียนผล")
        area_dir.mkdir(parents=True, exist_ok=True)
        ensure_area_subdirs(area_dir)
        run_root = local_state_root() / "runs" / code
        run_root.mkdir(parents=True, exist_ok=True)
        run_dir = run_root / f"{dt.datetime.now():%Y%m%d-%H%M%S}-{token}"
        run_dir.mkdir(parents=False, exist_ok=False)
        deadline = time.monotonic() + float(machine["limits"]["area_timeout_seconds"])
        success = False
        try:
            require_deadline(deadline, "อ่านประชากรโรงเรียน")
            small_all, magnets, overlap = load_area_records(area)
            size_profile = school_size_profile(small_all)
            analysis_origins = [
                row for row in small_all if row["school_id"] not in overlap
            ]
            create_school_points(
                partial["small"], "small_school_origins", analysis_origins, magnet=False
            )
            create_school_points(
                partial["magnets"], "magnet_schools", magnets, magnet=True
            )
            preparation = prepare_boundary_and_roads(
                machine, area, partial, run_dir, deadline
            )
            require_deadline(deadline, "อ่านแนวเขตตำบล")
            tambon_lines, tambon_meta = load_tambon_lines(machine, read_vector_frame(partial["boundary"]))
            require_deadline(deadline, "QNEAT3")
            full_od = run_qneat3(machine, partial, run_dir, deadline)
            require_deadline(deadline, "วิเคราะห์ OD")
            result = analyse_od_and_write_outputs(
                machine,
                full_od,
                partial["boundary"],
                small_all,
                analysis_origins,
                magnets,
                overlap,
                partial["od"],
                partial["status"],
            )
            result["preparation"] = preparation
            validate_expected_summary(area, result["summary"])
            validate_expected_result(area, result)
            require_deadline(deadline, "ภาพจำลองหลังจับคู่")
            consolidation = build_consolidation(area, result["summary"], partial["od"])
            require_deadline(deadline, "วาดแผนที่")
            render_standard_map(
                area,
                partial["boundary"],
                partial["roads"],
                partial["od"],
                partial["status"],
                partial["magnets"],
                partial["image"],
                tambon_lines=tambon_lines,
            )
            require_deadline(deadline, "วาดแผนที่ขนาดโรงเรียน")
            render_standard_size_map(
                area, partial["boundary"], partial["roads"], small_all, partial["size_image"], size_profile,
                tambon_lines=tambon_lines,
            )
            require_deadline(deadline, "สร้าง PDF")
            render_summary_pdf(
                area, result["summary"], partial["image"], partial["pdf"],
                size_image_path=partial["size_image"], consolidation=consolidation, size_profile=size_profile,
            )
            require_deadline(deadline, "สร้าง QC")
            qc_payload = build_qc_payload(
                machine, catalog, area, result, partial, preflight, size_profile=size_profile,
                presentation=fresh_presentation(tambon_meta), consolidation=consolidation,
            )
            write_json(partial["qc"], qc_payload)
            transactional_promote(partial, final, overwrite, token)
            success = True
            print(f"เผยแพร่ผลในโฟลเดอร์พื้นที่แล้ว: {area_dir}")
            return qc_payload
        finally:
            if success:
                try:
                    cleanup_successful_run_dir(run_dir)
                except (OSError, PipelineError) as exc:
                    # Publication already succeeded; a leftover diagnostic
                    # directory is a warning, not a failed published result.
                    print(f"[คำเตือน] ลบ run directory หลังสำเร็จไม่ได้: {exc}")
            else:
                cleanup_partial_outputs(partial, token)
                if not area_dir_existed and area_dir.is_dir():
                    remove_empty_area_subdirs(area_dir)
                    try:
                        area_dir.rmdir()  # succeeds only when truly empty
                    except OSError:
                        pass
                print(f"เก็บ log งานที่ไม่สำเร็จไว้ตรวจสอบ: {run_dir}")


def _plot_font_name() -> str:
    try:
        from matplotlib import font_manager
    except ImportError as exc:
        raise PipelineError("ไม่พบ matplotlib ใน QGIS Python") from exc
    candidates = (
        Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "tahoma.ttf",
        Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "leelawui.ttf",
    )
    for candidate in candidates:
        if candidate.is_file():
            font_manager.fontManager.addfont(str(candidate))
            return font_manager.FontProperties(fname=str(candidate)).get_name()
    return "DejaVu Sans"


def _nice_scale_length(width_metres: float) -> float:
    target = max(1000.0, width_metres * 0.18)
    choices = [1000, 2000, 5000, 10000, 20000, 50000, 100000, 200000]
    valid = [value for value in choices if value <= target]
    return float(valid[-1] if valid else choices[0])


def render_map_frames(
    area: dict[str, Any],
    boundary: Any,
    roads: Any,
    od_lines: Any,
    status_points: Any,
    magnets: Any,
    output_path: Path,
    tambon_lines: Any = None,
) -> None:
    ensure_within(output_path, WORKSPACE, "PNG แผนที่")
    try:
        import matplotlib

        matplotlib.use("Agg", force=True)
        import matplotlib.pyplot as plt
        from matplotlib.lines import Line2D
        from matplotlib.transforms import Bbox
    except ImportError as exc:
        raise PipelineError("ไม่พบ matplotlib สำหรับสร้างแผนที่") from exc

    font_name = _plot_font_name()
    plt.rcParams.update({"font.family": font_name, "axes.unicode_minus": False})
    for frame, label in (
        (boundary, "ขอบเขต"),
        (roads, "ถนน"),
        (od_lines, "OD"),
        (status_points, "สถานะโรงเรียน"),
        (magnets, "โรงเรียนคุณภาพชุมชน"),
    ):
        if frame is None or getattr(frame, "crs", None) is None:
            raise QualityError(f"ชั้น {label} ไม่มี CRS")
        if str(frame.crs).upper() != "EPSG:32647":
            frame.to_crs("EPSG:32647", inplace=True)

    # Frame the actual study features, never the larger road buffer. Include
    # coordinate exceptions so a closer crop cannot hide a school under review.
    extents = [frame.total_bounds for frame in (boundary, status_points, magnets, od_lines) if len(frame)]
    xmin = min(float(bounds[0]) for bounds in extents)
    ymin = min(float(bounds[1]) for bounds in extents)
    xmax = max(float(bounds[2]) for bounds in extents)
    ymax = max(float(bounds[3]) for bounds in extents)
    width = max(xmax - xmin, 1000.0)
    height = max(ymax - ymin, 1000.0)
    pad_x = width * 0.03
    pad_y = height * 0.03
    # An explicit, aspect-matched axes avoids the wide empty side margins from
    # equal-aspect data inside a fixed-height subplot. Coordinates stay metric.
    figure_width, side_margin = 12.0, 0.24
    header_height, footer_height = 0.86, 1.35
    map_width = figure_width - 2 * side_margin
    map_height = map_width * (height + 2 * pad_y) / (width + 2 * pad_x)
    figure_height = map_height + header_height + footer_height
    if not 3.0 <= figure_height <= 36.0:
        raise QualityError("สัดส่วนพื้นที่ผิดปกติสำหรับแผนที่; ตรวจพิกัดก่อนสร้างภาพ")
    fig = plt.figure(figsize=(figure_width, figure_height), dpi=200)
    ax = fig.add_axes((side_margin / figure_width, footer_height / figure_height,
                       map_width / figure_width, map_height / figure_height))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("#F8FAFC")

    _draw_admin_layers(ax, boundary, roads, tambon_lines)
    if len(od_lines):
        od_lines.plot(ax=ax, color="#4C78A8", linewidth=0.75, alpha=0.68, zorder=3)

    styles = {
        "MATCHED_10KM": dict(marker="o", facecolors="white", edgecolors="#1F4E79", s=52, linewidths=1.1),
        "OVER_10KM": dict(marker="o", c="#A9AEB5", s=34, edgecolors="white", linewidths=0.5),
        "UNREACHABLE": dict(marker="X", c="#E69F00", s=55, edgecolors="white", linewidths=0.6),
        "EXCLUDED_IS_MAGNET": dict(marker="D", c="#7B3294", s=48, edgecolors="white", linewidths=0.6),
        "COORD_BOUNDARY_EXCEPTION": dict(marker="^", c="#7B3294", s=65, edgecolors="white", linewidths=0.7),
    }
    if "สถานะ" not in status_points.columns:
        raise QualityError("ชั้นสถานะไม่มี field สถานะ")
    for status, style in styles.items():
        subset = status_points[status_points["สถานะ"] == status]
        if len(subset):
            ax.scatter(subset.geometry.x, subset.geometry.y, zorder=5, **style)

    if len(magnets):
        ax.scatter(
            magnets.geometry.x,
            magnets.geometry.y,
            marker="*",
            s=225,
            c="#E69F00",
            edgecolors="#7A3E00",
            linewidths=0.8,
            zorder=7,
        )

    ax.set_xlim(xmin - pad_x, xmax + pad_x)
    ax.set_ylim(ymin - pad_y, ymax + pad_y)
    ax.set_aspect("equal", adjustable="box")
    ax.axis("off")
    fig.text(side_margin / figure_width, 1 - 0.17 / figure_height,
             f"แผนที่ผลการจับคู่โรงเรียนขนาดเล็ก — {area['name']}",
             fontsize=17, color="#17365D", weight="bold", va="top")
    fig.text(side_margin / figure_width, 1 - 0.53 / figure_height,
             "แผน 1.2 — ระยะทางตามโครงข่ายถนนไม่เกิน 10 กิโลเมตร",
             fontsize=11, color="#4B5563", va="top")

    ax.annotate(
        "N",
        xy=(0.955, 0.97),
        xytext=(0.955, 0.90),
        xycoords="axes fraction",
        ha="center",
        va="center",
        fontsize=11,
        weight="bold",
        arrowprops=dict(arrowstyle="-|>", color="#17365D", lw=1.5),
        zorder=10,
    )
    scale = _nice_scale_length(width)
    sx = xmin - pad_x + (width + 2 * pad_x) * 0.06
    sy = ymin - pad_y + (height + 2 * pad_y) * 0.055
    ax.plot([sx, sx + scale], [sy, sy], color="#111111", linewidth=2.2, zorder=10)
    ax.plot([sx, sx], [sy - height * 0.006, sy + height * 0.006], color="#111111", linewidth=1.3, zorder=10)
    ax.plot([sx + scale, sx + scale], [sy - height * 0.006, sy + height * 0.006], color="#111111", linewidth=1.3, zorder=10)
    ax.text(sx + scale / 2, sy + height * 0.012, f"{scale / 1000:g} กม.", ha="center", va="bottom", fontsize=10,
            bbox=dict(fc="white", ec="none", alpha=0.85, pad=2))

    # Place readable labels within the map and away from other magnet labels.
    # Offsets affect text only; school coordinates and pair geometry never move.
    if len(magnets) and "school_name" in magnets.columns:
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        map_box = ax.get_window_extent(renderer).padded(-4)
        occupied = []
        for point in magnets.geometry:
            px, py = ax.transData.transform((point.x, point.y))
            occupied.append(Bbox.from_extents(px - 12, py - 12, px + 12, py + 12))
        # จุดโรงเรียนต้นทางทุกสถานะ (วงกลม/กากบาท/ข้าวหลามตัด รัศมี ≤ 10 px ที่ 200 dpi): กรอบป้ายชื่อโรงเรียนคุณภาพชุมชน
        # ต้องไม่บังจุดเหล่านี้ (พบที่เชียงราย เขต 1 รุ่น 1.7.0 ป้ายบังวงกลมคู่ที่จับได้ 2 จุด — บันทึกการเปลี่ยนแปลง.md §14.11)
        if status_points is not None and len(status_points):
            for point in status_points.geometry:
                px, py = ax.transData.transform((point.x, point.y))
                occupied.append(Bbox.from_extents(px - 11, py - 11, px + 11, py + 11))
        offsets = [(10, 10), (-10, 10), (10, -10), (-10, -10),
                   (10, 28), (-10, 28), (10, -28), (-10, -28),
                   (10, 48), (-10, 48), (10, -48), (-10, -48)]
        for _, row in magnets.sort_values("school_name", kind="stable").iterrows():
            label = str(row.get("school_name", "") or "").strip()
            if not label or label.lower() in {"true", "false", "nan", "none"}:
                continue
            if len(label) > 28 and "(" in label:
                label = label.replace("(", "\n(", 1)
            candidates = []
            for dx, dy in offsets:
                text_item = ax.annotate(
                    label, (row.geometry.x, row.geometry.y), xytext=(dx, dy),
                    textcoords="offset points", ha="left" if dx > 0 else "right",
                    va="bottom" if dy > 0 else "top", fontsize=11.5,
                    color="#243447", linespacing=1.35,
                    bbox=dict(boxstyle="round,pad=0.24", fc="white", ec="#DEE5EC", alpha=0.94, lw=0.5),
                    zorder=8,
                )
                box = text_item.get_window_extent(renderer).padded(5)
                inside = Bbox.intersection(box, map_box)
                outside = box.width * box.height - (inside.width * inside.height if inside else 0)
                overlap = 0.0
                for other in occupied:
                    intersection = Bbox.intersection(box, other)
                    if intersection is not None:
                        overlap += intersection.width * intersection.height
                score = outside * 100 + overlap * 10 + abs(dy)
                candidates.append((score, dx, dy, box))
                text_item.remove()
            _, dx, dy, box = min(candidates, key=lambda item: item[0])
            ax.annotate(
                label, (row.geometry.x, row.geometry.y), xytext=(dx, dy),
                textcoords="offset points", ha="left" if dx > 0 else "right",
                va="bottom" if dy > 0 else "top", fontsize=11.5,
                color="#243447", linespacing=1.35,
                bbox=dict(boxstyle="round,pad=0.24", fc="white", ec="#DEE5EC", alpha=0.94, lw=0.5),
                arrowprops=dict(arrowstyle="-", color="#697B8B", lw=0.6, shrinkA=3, shrinkB=7),
                zorder=8,
            )
            occupied.append(box)

    legend_handles = [
        Line2D([0], [0], color="#4C78A8", lw=1.3, label="เส้นเชื่อมคู่ที่ผ่านเกณฑ์"),
        Line2D([0], [0], marker="*", color="none", markerfacecolor="#E69F00", markeredgecolor="#7A3E00", markersize=11, label="โรงเรียนคุณภาพชุมชน"),
        Line2D([0], [0], marker="o", color="none", markerfacecolor="white", markeredgecolor="#1F4E79", markersize=6, label="จับคู่ได้ภายใน 10 กม."),
        Line2D([0], [0], marker="o", color="none", markerfacecolor="#A9AEB5", markeredgecolor="white", markersize=6, label="ปลายทางใกล้ที่สุดเกิน 10 กม."),
    ]
    for status, marker, color, label in (
        ("UNREACHABLE", "X", "#E69F00", "ไม่พบเส้นทาง"),
        ("EXCLUDED_IS_MAGNET", "D", "#7B3294", "ต้นทางเป็นโรงเรียนคุณภาพชุมชนด้วย"),
        ("COORD_BOUNDARY_EXCEPTION", "^", "#7B3294", "พิกัด–ขอบเขตต้องตรวจสอบ"),
    ):
        if (status_points["สถานะ"] == status).any():
            legend_handles.append(Line2D([0], [0], marker=marker, color="none", markerfacecolor=color,
                                         markeredgecolor="white", markersize=7, label=label))
    if tambon_lines is not None and len(tambon_lines):
        legend_handles.append(_tambon_legend_handle(Line2D))
    fig.legend(
        handles=legend_handles,
        loc="upper center", bbox_to_anchor=(0.5, 1.25 / figure_height),
        frameon=False, fontsize=9.5, ncol=4, columnspacing=1.7, labelspacing=0.8,
    )
    fig.text(
        0.5,
        0.16 / figure_height,
        "เส้นเชื่อมแสดงความสัมพันธ์ตามผลการจับคู่ มิใช่แนวเส้นทางเดินรถจริง\n"
        + map_source_footer(tambon_lines),
        ha="center",
        va="bottom",
        fontsize=9.0, linespacing=1.6,
        color="#4B5563",
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    # Keep the declared canvas at 2,400 px wide; the map fills its aspect-matched
    # axes without stretching geography or cropping content at export time.
    fig.savefig(output_path, dpi=200, facecolor="white")
    plt.close(fig)
    try:
        from PIL import Image

        with Image.open(output_path) as image:
            if image.width < 2000:
                raise QualityError(f"PNG กว้างเพียง {image.width}px; ต้องอย่างน้อย 2000px")
    except ImportError as exc:
        raise PipelineError("ไม่พบ Pillow สำหรับตรวจ PNG") from exc


def render_standard_map(
    area: dict[str, Any],
    boundary_path: Path,
    roads_path: Path,
    od_path: Path,
    status_path: Path,
    magnets_path: Path,
    output_path: Path,
    tambon_lines: Any = None,
) -> None:
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas ใน QGIS Python") from exc
    render_map_frames(
        area,
        gpd.read_file(boundary_path),
        gpd.read_file(roads_path),
        gpd.read_file(od_path),
        gpd.read_file(status_path),
        gpd.read_file(magnets_path),
        output_path,
        tambon_lines=tambon_lines,
    )


def size_bin_label(students: int) -> str:
    for low, high, _ in SIZE_BINS:
        if low <= students <= high:
            return f"{low}-{high}"
    raise QualityError(
        f"จำนวนนักเรียน {students} อยู่นอกช่วงโรงเรียนขนาดเล็ก 0–{SMALL_SCHOOL_MAX_STUDENTS} คน"
    )


def size_map_bins() -> tuple[tuple[int, int, str], ...]:
    """ช่วงที่วาดบนแผนที่ขนาด (ไม่เกิน SIZE_MAP_MAX_STUDENTS); ช่วงที่เหลือนับใน QC และตารางแต่ไม่วาดจุด"""
    drawn = tuple(item for item in SIZE_BINS if item[1] <= SIZE_MAP_MAX_STUDENTS)
    missing = [f"{low}-{high}" for low, high, _ in drawn if f"{low}-{high}" not in SIZE_MAP_MARKER_AREA]
    if not drawn or missing or drawn[-1][1] != SIZE_MAP_MAX_STUDENTS:
        raise PipelineError("SIZE_BINS, SIZE_MAP_MAX_STUDENTS และ SIZE_MAP_MARKER_AREA ไม่สอดคล้องกัน")
    return drawn


def school_size_profile(
    small_rows: list[dict[str, Any]], expected_total: int | None = None
) -> dict[str, Any]:
    """สรุปจำนวนโรงเรียนต่อช่วงขนาดจากทะเบียน DMC (ตัวเลขรวมเท่านั้น ไม่มีรายโรงเรียน)"""
    values = sorted(int(row["students"]) for row in small_rows)
    if not values:
        raise QualityError("ไม่มีโรงเรียนขนาดเล็กสำหรับสรุปขนาด")
    if expected_total is not None and len(values) != int(expected_total):
        raise QualityError(
            f"จำนวนโรงเรียนในทะเบียน ({len(values)}) ไม่ตรงค่าตรวจรับ ({expected_total})"
        )
    counts = {f"{low}-{high}": 0 for low, high, _ in SIZE_BINS}
    for value in values:
        counts[size_bin_label(value)] += 1
    count = len(values)
    middle = count // 2
    median = float(values[middle]) if count % 2 else (values[middle - 1] + values[middle]) / 2
    return {
        "basis": "จำนวนนักเรียนรวมตามทะเบียน DMC682 ของโรงเรียนขนาดเล็ก (ไม่เกิน 120 คน)",
        "bins": [
            {"label": f"{low}-{high}", "min": low, "max": high, "count": counts[f"{low}-{high}"]}
            for low, high, _ in SIZE_BINS
        ],
        "total": count,
        "students_min": values[0],
        "students_max": values[-1],
        "students_median": median,
        "students_total": sum(values),
        "map_max_students": SIZE_MAP_MAX_STUDENTS,
    }


def public_size_analysis(profile: dict[str, Any]) -> str:
    total = int(profile["total"])
    counts = {item["label"]: int(item["count"]) for item in profile["bins"]}
    shown = sum(counts.get(f"{low}-{high}", 0) for low, high, _ in size_map_bins())
    hidden = total - shown
    tiny = counts.get("0-10", 0) + counts.get("11-20", 0)

    def percent(value: int) -> str:
        return f"{value / total * 100:.1f}" if total else "0.0"

    median = float(profile["students_median"])
    median_text = f"{int(median)}" if median.is_integer() else f"{median:.1f}"
    tiny_text = (
        f"ในจำนวนนี้มีนักเรียนไม่เกิน 20 คน จำนวน {tiny} แห่ง"
        if tiny else "ในจำนวนนี้ไม่มีโรงเรียนที่มีนักเรียนไม่เกิน 20 คน"
    )
    return (
        f"โรงเรียนขนาดเล็กตามทะเบียนในพื้นที่ {total} แห่ง มีนักเรียนตั้งแต่ "
        f"{int(profile['students_min'])} ถึง {int(profile['students_max'])} คนต่อโรงเรียน (มัธยฐาน {median_text} คน) "
        f"แผนที่แสดงเฉพาะโรงเรียนที่มีนักเรียนไม่เกิน {SIZE_MAP_MAX_STUDENTS} คน จำนวน {shown} แห่ง (ร้อยละ {percent(shown)}) "
        f"{tiny_text} ส่วนโรงเรียนที่มีนักเรียน "
        f"{SIZE_MAP_MAX_STUDENTS + 1}–{SMALL_SCHOOL_MAX_STUDENTS} คนอีก {hidden} แห่ง (ร้อยละ {percent(hidden)}) "
        "นับรวมในตารางแต่ไม่แสดงบนแผนที่ ข้อมูลนี้แสดงสภาพขนาดโรงเรียนตามทะเบียน มิใช่เกณฑ์ตัดสินการควบรวมหรือย้ายนักเรียน"
    )


def portal_size_profile(profile: dict[str, Any]) -> dict[str, Any]:
    """รูปแบบสำหรับ catalog: ค่าคงที่ของแผนที่มาจากโค้ด (ไม่ใช่ QC) เพื่อให้ QC รุ่นก่อนและรุ่นใหม่ให้ payload เดียวกัน"""
    bins = [
        {
            "label": str(item["label"]),
            "min": int(item["min"]),
            "max": int(item["max"]),
            "count": int(item["count"]),
        }
        for item in profile["bins"]
    ]
    return {
        "total": int(profile["total"]),
        "studentsMin": int(profile["students_min"]),
        "studentsMax": int(profile["students_max"]),
        "studentsMedian": profile["students_median"],
        "mapMaxStudents": SIZE_MAP_MAX_STUDENTS,
        "onMapTotal": sum(item["count"] for item in bins if item["max"] <= SIZE_MAP_MAX_STUDENTS),
        "bins": bins,
    }


def size7_code(students: int) -> str:
    """รหัสเกณฑ์ 7 ขนาดตาม DMC682 สำหรับจำนวนนักเรียนรวม"""
    if students < 0:
        raise QualityError(f"จำนวนนักเรียนติดลบ: {students}")
    for low, high, code in SIZE7_BOUNDS:
        if students >= low and (high is None or students <= high):
            return code
    raise QualityError(f"จำแนกเกณฑ์ 7 ขนาดไม่ได้: {students}")


def size7_label(code: str) -> str:
    """ป้ายอ่านง่ายจากรหัส เช่น '3_(201-300)' -> '3 (201–300)', '7_(>=2500)' -> '7 (2,500 ขึ้นไป)'"""
    for low, high, known in SIZE7_BOUNDS:
        if known == code:
            level = code.split("_", 1)[0]
            if high is None:
                return f"{level} ({low:,} ขึ้นไป)"
            if low == high:
                return f"{level} ({low:,})"
            return f"{level} ({low:,}–{high:,})"
    raise QualityError(f"รหัสเกณฑ์ 7 ขนาดไม่รู้จัก: {code!r}")


def load_master_records(area: dict[str, Any]) -> list[dict[str, Any]]:
    """โรงเรียน สพป. ทุกขนาดในเขตจากทะเบียนหลัก DMC682_สพป.csv (คอลัมน์ areacode, โรงเรียน, ชื่อโรงเรียน, รวมนักเรียน, เกณฑ์7ขนาด)"""
    code = area["areacode"]
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in csv_rows(MASTER_CSV):
        if str(row.get("areacode", "")).strip() != code:
            continue
        school_id = normalize_school_id(row.get("โรงเรียน"), "รหัสโรงเรียนในทะเบียนหลัก")
        if school_id in seen:
            raise QualityError(f"ทะเบียนหลักมีรหัสโรงเรียนซ้ำ: {school_id}")
        seen.add(school_id)
        students = int(float(row.get("รวมนักเรียน") or 0))
        size_code = str(row.get("เกณฑ์7ขนาด", "")).strip()
        if size7_code(students) != size_code:
            raise QualityError(f"เกณฑ์7ขนาดของ {school_id} ({size_code!r}) ไม่ตรงจำนวนนักเรียน {students}")
        rows.append({
            "school_id": school_id,
            "school_name": str(row.get("ชื่อโรงเรียน", "")).strip(),
            "students": students,
            "size_code": size_code,
        })
    if not rows:
        raise QualityError(f"ทะเบียนหลักไม่มีโรงเรียนของ {area['name']} ({code})")
    return rows


def read_od_pairs(path: Path) -> list[dict[str, Any]]:
    """คู่จับคู่ที่ผ่านเกณฑ์จากชั้น OD ที่ตรวจรับแล้ว: ระยะเป็นเมตรจาก total_cost หรือ distance_m (legacy) หรือ distance_km*1000"""
    ogr, _ = ogr_modules()
    dataset = ogr.Open(str(ensure_file(path, "OD ที่ตรวจรับแล้ว")), 0)
    if dataset is None or dataset.GetLayerCount() < 1:
        raise QualityError(f"เปิด OD ไม่ได้: {path}")
    layer = dataset.GetLayer(0)
    definition = layer.GetLayerDefn()
    fields = {definition.GetFieldDefn(index).GetName() for index in range(definition.GetFieldCount())}
    if not {"origin_id", "destination_id"}.issubset(fields):
        raise QualityError(f"OD ไม่มี field origin_id/destination_id: {path.name}")
    if "total_cost" in fields:
        field, scale = "total_cost", 1.0
    elif "distance_m" in fields:
        field, scale = "distance_m", 1.0
    elif "distance_km" in fields:
        field, scale = "distance_km", 1000.0
    else:
        raise QualityError(f"OD ไม่มี field ระยะทาง: {path.name}")
    pairs: dict[str, dict[str, Any]] = {}
    for feature in layer:
        origin = normalize_school_id(feature.GetField("origin_id"), "origin_id")
        destination = normalize_school_id(feature.GetField("destination_id"), "destination_id")
        try:
            metres = float(feature.GetField(field)) * scale
        except (TypeError, ValueError) as exc:
            raise QualityError(f"OD มีระยะไม่ใช่ตัวเลขสำหรับ {origin}") from exc
        if not math.isfinite(metres) or metres < 0 or metres > MAX_DISTANCE_METRES + 1e-6:
            raise QualityError(f"OD มีระยะผิดปกติสำหรับ {origin}: {metres}")
        if origin in pairs:
            raise QualityError(f"OD มี origin ซ้ำ: {origin}")
        pairs[origin] = {"origin": origin, "destination": destination, "metres": metres}
    dataset = None
    ordered = [pairs[key] for key in sorted(pairs)]
    for item in ordered:
        item["field"] = field
    return ordered


def accepted_od_path(area: dict[str, Any]) -> Path:
    """พาธชั้น OD ที่ตรวจรับแล้วของพื้นที่ (ไฟล์นำร่อง หรือชื่อมาตรฐานในโฟลเดอร์พื้นที่)"""
    if isinstance(area.get("pilot"), dict):
        return _pilot_file(area, "od")
    area_dir = ensure_within(AREA_ROOT / area["folder_name"], AREA_ROOT, "โฟลเดอร์รายพื้นที่")
    return ensure_file(area_dir / layout.area_relpath("od", output_names(area)["od"]), "OD ของพื้นที่")


def consolidation_profile(
    area: dict[str, Any],
    small_rows: list[dict[str, Any]],
    magnets: list[dict[str, Any]],
    overlap: set[str],
    od_pairs: list[dict[str, Any]],
    master_rows: list[dict[str, Any]],
    summary: dict[str, int],
    od_path: Path,
    od_name: str | None = None,
) -> dict[str, Any]:
    """ภาพจำลองหากโรงเรียนขนาดเล็กที่จับคู่ได้ทุกแห่งไปรวมกับโรงเรียนคุณภาพชุมชนปลายทาง (คำนวณซ้ำได้จาก DMC + OD ที่ตรวจรับ)

    ตัวเลขเป็นภาพจำลองประกอบการพิจารณา มิใช่แผนควบรวมหรือย้ายนักเรียน; ไม่มีชื่อโรงเรียนขนาดเล็กในผลลัพธ์
    """
    small_by_id = {row["school_id"]: row for row in small_rows}
    magnet_ids = {row["school_id"] for row in magnets}
    master_by_id = {row["school_id"]: row for row in master_rows}
    if len(od_pairs) != int(summary["matched"]):
        raise QualityError(f"คู่ OD {len(od_pairs)} ไม่เท่าจำนวนจับคู่ได้ {summary['matched']} ของ {area['name']}")
    for pair in od_pairs:
        if pair["origin"] not in small_by_id or pair["origin"] in overlap:
            raise QualityError(f"OD มีต้นทางที่ไม่ใช่โรงเรียนขนาดเล็กของพื้นที่: {pair['origin']}")
        if pair["destination"] not in magnet_ids:
            raise QualityError(f"OD มีปลายทางที่ไม่ใช่โรงเรียนคุณภาพชุมชนของพื้นที่: {pair['destination']}")
    missing = sorted((set(small_by_id) | magnet_ids) - set(master_by_id))
    if missing:
        raise QualityError(f"ทะเบียนหลักไม่มีโรงเรียนของพื้นที่: {missing[:5]}")
    for pair in od_pairs:
        if int(small_by_id[pair['origin']]['students']) != int(master_by_id[pair['origin']]['students']):
            raise QualityError(f"จำนวนนักเรียนของ {pair['origin']} ในทะเบียนโรงเรียนขนาดเล็กไม่ตรงทะเบียนหลัก")

    feeders: dict[str, list[dict[str, Any]]] = {}
    for pair in od_pairs:
        feeders.setdefault(pair["destination"], []).append(pair)
    schools = []
    for magnet in magnets:
        master = master_by_id[magnet["school_id"]]
        received_pairs = feeders.get(magnet["school_id"], [])
        received = sum(int(small_by_id[pair["origin"]]["students"]) for pair in received_pairs)
        after = int(master["students"]) + received
        schools.append({
            "school_id": magnet["school_id"],
            "school_name": magnet["school_name"],
            "students_before": int(master["students"]),
            "size_code_before": master["size_code"],
            "feeders": len(received_pairs),
            "students_received": received,
            "students_after": after,
            "size_code_after": size7_code(after),
            "farthest_km": round(max(pair["metres"] for pair in received_pairs) / 1000.0, 2) if received_pairs else None,
            "is_small_school": magnet["school_id"] in overlap,
        })
    schools.sort(key=lambda item: (-item["students_after"], item["school_id"]))
    students_moved = sum(item["students_received"] for item in schools)
    if sum(item["feeders"] for item in schools) != len(od_pairs) or students_moved != sum(
        int(small_by_id[pair["origin"]]["students"]) for pair in od_pairs
    ):
        raise QualityError("ผลรวมรายโรงเรียนคุณภาพชุมชนไม่ตรงคู่ OD")
    metres = [pair["metres"] for pair in od_pairs]
    before_total = len(master_rows)
    aggregate = {
        "schools_before": before_total,
        "schools_after": before_total - len(od_pairs),
        "small_registered": int(summary["registered"]),
        "small_matched": len(od_pairs),
        "small_remaining": int(summary["registered"]) - len(od_pairs),
        "small_remaining_by_status": {
            key: int(summary.get(key, 0)) for key in ("over10km", "unreachable", "excluded", "coordExceptions")
        },
        "students_small_total": sum(int(row["students"]) for row in small_rows),
        "students_moved": students_moved,
        "magnets_total": len(magnets),
        "magnets_receiving": sum(1 for item in schools if item["feeders"]),
        "magnets_students_before": sum(item["students_before"] for item in schools),
        "magnets_students_after": sum(item["students_after"] for item in schools),
        "distance_mean_km": round(sum(metres) / len(metres) / 1000.0, 2) if metres else None,
        "distance_max_km": round(max(metres) / 1000.0, 2) if metres else None,
    }
    return {
        "version": CONSOLIDATION_VERSION,
        "basis": (
            "ภาพจำลองผลคัดกรองตามแผน 1.2: สมมติว่าโรงเรียนขนาดเล็กทุกแห่งที่จับคู่ได้ภายใน 10 กิโลเมตรไปรวมกับ"
            "โรงเรียนคุณภาพชุมชนปลายทางตามผล นับนักเรียนตามทะเบียน DMC682 ทั้งหมด (นักเรียนรวมเชิงสมมติ) "
            "มิใช่แผนควบรวม ข้อเสนอย้ายนักเรียน หรือการคาดการณ์จำนวนนักเรียนจริง"
        ),
        "distance_field": od_pairs[0]["field"] if od_pairs else None,
        "aggregate": aggregate,
        "magnets": schools,
        "inputs": {
            "dmc_master_sha256": sha256_file(MASTER_CSV),
            "od": {**_plain_file_fingerprint(od_path), "filename": od_name or od_path.name},
        },
    }


def build_consolidation(area: dict[str, Any], summary: dict[str, int], od_path: Path) -> dict[str, Any]:
    """ชื่อไฟล์ OD ที่บันทึกใน inputs ใช้ชื่อสุดท้ายจาก output_names เสมอ (ตอน run ไฟล์ยังเป็น .partial)"""
    small_rows, magnets, overlap = load_area_records(area)
    return consolidation_profile(
        area, small_rows, magnets, overlap, read_od_pairs(od_path), load_master_records(area), summary, od_path,
        od_name=output_names(area)["od"],
    )


def public_consolidation_note(profile: dict[str, Any]) -> str:
    agg = profile["aggregate"]
    excluded = int(agg["small_remaining_by_status"].get("excluded", 0))
    excluded_clause = (
        f" (ในจำนวนนี้ {excluded} แห่งเป็นโรงเรียนขนาดเล็กตามทะเบียนที่มีบทบาทโรงเรียนคุณภาพชุมชนด้วย)" if excluded else ""
    )
    return (
        f"ภาพจำลองนี้คำนวณจากผลคัดกรองตามแผน 1.2 โดยสมมติว่าโรงเรียนขนาดเล็กทุกแห่งที่จับคู่ได้ภายใน 10 กิโลเมตร "
        f"จำนวน {agg['small_matched']:,} แห่ง ไปรวมกับโรงเรียนคุณภาพชุมชนปลายทางตามผล และนับนักเรียนตามทะเบียน DMC682 ทั้งหมด "
        f"โรงเรียน สพป. ในเขตจะลดจาก {agg['schools_before']:,} เหลือ {agg['schools_after']:,} แห่ง "
        f"โรงเรียนคุณภาพชุมชน {agg['magnets_receiving']} จาก {agg['magnets_total']} แห่งรับนักเรียนเพิ่มรวม {agg['students_moved']:,} คน "
        f"ส่วนโรงเรียนขนาดเล็กอีก {agg['small_remaining']:,} แห่งยังคงอยู่ตามผลคัดกรอง{excluded_clause} "
        "ตัวเลขทั้งหมดเป็นภาพจำลองเพื่อประกอบการพิจารณา มิใช่แผนควบรวม ข้อเสนอย้ายนักเรียน หรือการคาดการณ์จำนวนนักเรียนจริง"
    )


def portal_consolidation(profile: dict[str, Any]) -> dict[str, Any]:
    """รูปแบบสำหรับ catalog: ไม่มีรหัสโรงเรียนและไม่มีชื่อโรงเรียนขนาดเล็กต้นทาง; ชื่อโรงเรียนคุณภาพชุมชนทุกแห่ง
    (รวมแห่งที่เป็นโรงเรียนขนาดเล็กตามทะเบียนด้วย ซึ่งติดเครื่องหมาย †) แสดงได้เพราะปรากฏบนแผนที่อยู่แล้ว"""
    agg = profile["aggregate"]
    schools = [
        {
            "name": str(item["school_name"]),
            "studentsBefore": int(item["students_before"]),
            "feeders": int(item["feeders"]),
            "studentsReceived": int(item["students_received"]),
            "studentsAfter": int(item["students_after"]),
            "sizeBefore": size7_label(item["size_code_before"]),
            "sizeAfter": size7_label(item["size_code_after"]),
            "farthestKm": item["farthest_km"],
            "isSmallSchool": bool(item["is_small_school"]),
        }
        for item in profile["magnets"]
    ]
    for item in schools:
        if re.search(r"\d{8,}", item["name"]):
            raise QualityError("ชื่อโรงเรียนคุณภาพชุมชนมีรหัสตัวเลขปน")
    return {
        "schoolsBefore": int(agg["schools_before"]),
        "schoolsAfter": int(agg["schools_after"]),
        "smallRegistered": int(agg["small_registered"]),
        "smallMatched": int(agg["small_matched"]),
        "smallRemaining": int(agg["small_remaining"]),
        "studentsMoved": int(agg["students_moved"]),
        "qualityReceiving": int(agg["magnets_receiving"]),
        "qualityTotal": int(agg["magnets_total"]),
        "distanceMeanKm": agg["distance_mean_km"],
        "distanceMaxKm": agg["distance_max_km"],
        "note": public_consolidation_note(profile),
        "schools": schools,
    }


def verify_consolidation(area: dict[str, Any], qc: dict[str, Any]) -> None:
    """คำนวณภาพจำลองใหม่จาก DMC ปัจจุบันและ OD ที่ตรวจรับ แล้วเทียบกับ QC ทุกค่า (แบบเดียวกับ verify_school_size)"""
    recorded = qc.get("consolidation")
    if not isinstance(recorded, dict):
        raise QualityError(f"QC {area['name']} ไม่มี consolidation; ต้องรัน refresh-assets")
    od_path = accepted_od_path(area)
    expected = build_consolidation(area, _qc_summary(qc), od_path)
    keys = ("version", "distance_field", "aggregate", "magnets")
    if {key: recorded.get(key) for key in keys} != {key: expected[key] for key in keys}:
        raise QualityError(f"ภาพจำลองหลังจับคู่ใน QC ของ {area['name']} ไม่ตรงข้อมูลปัจจุบัน; ต้องรัน refresh-assets")
    inputs = recorded.get("inputs") or {}
    if inputs.get("dmc_master_sha256") != expected["inputs"]["dmc_master_sha256"]:
        raise QualityError(f"ทะเบียนหลัก DMC เปลี่ยนจากที่ QC {area['name']} ใช้; ต้องรัน refresh-assets")
    if inputs.get("od") != expected["inputs"]["od"]:
        raise QualityError(f"ชั้น OD ของ {area['name']} เปลี่ยนจากที่ QC ใช้; ต้องรัน refresh-assets")


def load_tambon_lines(machine: dict[str, Any], boundary: Any) -> tuple[Any, dict[str, Any]]:
    """อ่านแนวเขตตำบลจาก tha_admin3 ภายนอกด้วย exact path เฉพาะส่วนที่ทับขอบเขตพื้นที่ แล้ว clip ให้อยู่ในขอบเขต

    อ่านอย่างเดียว ไม่คัดลอกชั้นข้อมูลทั้งประเทศเข้า workspace และไม่เขียนไฟล์ข้างฐานภายนอก (Project Tor.md หัวข้อ 6)
    """
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas ใน QGIS Python") from exc
    if boundary is None or getattr(boundary, "crs", None) is None or not len(boundary):
        raise QualityError("ต้องมีขอบเขตพื้นที่ที่มี CRS ก่อนอ่านแนวเขตตำบล")
    boundary_utm = boundary if str(boundary.crs).upper() == "EPSG:32647" else boundary.to_crs("EPSG:32647")
    path = exact_external_paths(machine)["admin3"]
    tambon = gpd.read_file(path, mask=boundary_utm, columns=["adm3_pcode", "adm3_name1", "adm2_name1"])
    if tambon.empty:
        raise QualityError("ไม่พบตำบลที่ทับขอบเขตพื้นที่ในชั้นข้อมูล tha_admin3")
    tambon = tambon.to_crs("EPSG:32647")
    clipped = gpd.clip(tambon, boundary_utm, keep_geom_type=True)
    # ตัด sliver จากความคลาดเคลื่อนทศนิยมของจุดยอด (พื้นที่ต่ำกว่า 1 ตร.ม.) ก่อนนับตำบลและก่อนสร้างเส้น
    clipped = clipped[clipped.geometry.area > 1.0]
    if clipped.empty:
        raise QualityError("clip แนวเขตตำบลแล้วไม่เหลือรูปหลายเหลี่ยม")
    try:
        import shapely
    except ImportError as exc:
        raise PipelineError("ไม่พบ shapely ใน QGIS Python") from exc
    # รวมเส้นขอบที่ตำบลข้างเคียงใช้ร่วมกันให้เหลือเส้นเดียว มิฉะนั้นเส้นประสองเส้นซ้อนกันจะกลายเป็นเส้นทึบ
    merged = shapely.line_merge(shapely.union_all(list(clipped.geometry.boundary)))
    lines = gpd.GeoDataFrame(geometry=[merged], crs="EPSG:32647").explode(index_parts=False)
    lines = lines[lines.geometry.notna() & ~lines.geometry.is_empty].reset_index(drop=True)
    if lines.empty:
        raise QualityError("clip แนวเขตตำบลแล้วไม่เหลือเส้น")
    meta = {
        "source": "admin3",
        "layer": path.name,
        "fingerprint": shapefile_fingerprint(path),
        "tambon_count": int(clipped["adm3_pcode"].nunique()),
    }
    return lines, meta


def fresh_presentation(tambon_meta: dict[str, Any] | None) -> dict[str, Any]:
    """บล็อก presentation สำหรับ QC ที่สร้างใหม่ (run/pilot-assets) ให้บันทึกที่มาชั้นแสดงผลไว้ที่เดียวกับ refresh"""
    return {
        "renderer_version": VERSION,
        "mode": "fresh-render",
        "school_size_map": True,
        "size_map_max_students": SIZE_MAP_MAX_STUDENTS,
        "size_map_labels_version": SIZE_MAP_LABELS_VERSION,
        "size_summary_page_version": SIZE_SUMMARY_PAGE_VERSION,
        "consolidation_version": CONSOLIDATION_VERSION,
        "tambon_lines": tambon_meta,
    }


def _draw_admin_layers(ax: Any, boundary: Any, roads: Any, tambon_lines: Any) -> None:
    """ลำดับชั้นร่วมของทั้งสองแผนที่: พื้นขอบเขต → แนวเขตตำบล (เส้นประ) → ถนน → เส้นขอบเขต สพป. ทับบนสุดให้คมชัด"""
    boundary.plot(ax=ax, facecolor="#EEF4F9", edgecolor="none", zorder=1)
    if tambon_lines is not None and len(tambon_lines):
        if str(tambon_lines.crs).upper() != "EPSG:32647":
            tambon_lines = tambon_lines.to_crs("EPSG:32647")
        tambon_lines.plot(ax=ax, color=TAMBON_LINE_COLOR, linewidth=0.7, linestyle=(0, (3, 2)), zorder=1.5)
    if roads is not None and len(roads):
        roads.plot(ax=ax, color="#C8CDD3", linewidth=0.30, alpha=0.85, zorder=2)
    boundary.boundary.plot(ax=ax, color="#17365D", linewidth=1.5, zorder=2.6)


def _tambon_legend_handle(line_cls: Any) -> Any:
    return line_cls([0], [0], color=TAMBON_LINE_COLOR, lw=1.1, linestyle=(0, (3, 2)), label="แนวเขตตำบล")


def map_source_footer(tambon_lines: Any) -> str:
    parts = ["ที่มาถนน: © OpenStreetMap contributors"]
    if tambon_lines is not None and len(tambon_lines):
        parts.append("แนวเขตตำบล: ชั้นข้อมูลขอบเขตการปกครอง tha_admin3")
    parts += ["ชุดข้อมูล DMC682", "ฉบับนำร่องเพื่อการตรวจสอบ"]
    return "  |  ".join(parts)


def _segment_point_distance(x1: float, y1: float, x2: float, y2: float, px: float, py: float) -> float:
    """ระยะจากจุด (px, py) ถึงส่วนของเส้นตรง (x1, y1)–(x2, y2) ในพิกัดหน้าจอ"""
    dx, dy = x2 - x1, y2 - y1
    length_sq = dx * dx + dy * dy
    if length_sq <= 1e-9:
        return math.hypot(px - x1, py - y1)
    t = max(0.0, min(1.0, ((px - x1) * dx + (py - y1) * dy) / length_sq))
    return math.hypot(px - (x1 + t * dx), py - (y1 + t * dy))


def _box_point_distance(box: Any, px: float, py: float) -> float:
    """ระยะจากจุดถึงกรอบสี่เหลี่ยม (0 เมื่ออยู่ในกรอบ)"""
    cx = min(max(px, box.x0), box.x1)
    cy = min(max(py, box.y0), box.y1)
    return math.hypot(px - cx, py - cy)


def _place_size_map_labels(fig: Any, ax: Any, drawn: Any, bbox_cls: Any, decorations: Iterable[Any] = ()) -> int:
    """วางป้ายชื่อโรงเรียนตามทะเบียนให้ทุกโรงเรียนที่วาดบนแผนที่ขนาด (ตัดสินร่วมกับผู้ใช้ 8 ก.ย. 2569 รอบห้า)

    เลื่อนเฉพาะตำแหน่งข้อความ: ลองตำแหน่งรอบจุด 28 แบบจากใกล้ไปไกล ให้คะแนนโทษเมื่อ (1) ป้ายล้นกรอบแผนที่
    (2) ป้ายทับจุดอื่น ป้ายอื่น ลูกศรทิศเหนือ หรือมาตราส่วน (3) เส้นโยงจากจุดไปป้ายพาดผ่านจุดของโรงเรียนอื่น
    (4) ป้ายอยู่ใกล้จุดของโรงเรียนอื่นมากกว่าจุดของตนเอง แล้วบวกระยะเล็กน้อยเพื่อชอบตำแหน่งใกล้
    หยุดค้นเมื่อพบตำแหน่งที่ไม่มีโทษใดเลย ตำแหน่งที่เลือกคือคะแนนต่ำสุดในที่ลองมา พิกัดจุดไม่เลื่อน
    ทุกป้ายมีเส้นโยงไปยังจุดของตน ลำดับการวางคงที่ (นักเรียนน้อยก่อน แล้วชื่อ แล้วรหัส) เพื่อให้ภาพซ้ำได้
    คืนจำนวนป้ายที่วาง ผู้เรียกต้องตรวจว่าเท่าจำนวนโรงเรียนที่วาด
    """
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    map_box = ax.get_window_extent(renderer).padded(-4)
    px_per_pt = fig.dpi / 72.0
    # กรอบขาวของป้ายกว้างกว่าตัวอักษร pad 0.2 em (Text.get_window_extent ไม่รวมกรอบ) บวกช่องไฟเล็กน้อย
    box_pad = 0.2 * SIZE_MAP_LABEL_FONTSIZE * px_per_pt + 2.0
    marker_at: dict[Any, tuple[float, float, float]] = {}
    occupied = []
    for index, row in drawn.iterrows():
        px, py = ax.transData.transform((row.geometry.x, row.geometry.y))
        radius = math.sqrt(SIZE_MAP_MARKER_AREA[row["bin"]]) / 2 * px_per_pt
        marker_at[index] = (float(px), float(py), float(radius))
        occupied.append(bbox_cls.from_extents(px - radius - 1.5, py - radius - 1.5, px + radius + 1.5, py + radius + 1.5))
    for artist in decorations:  # ลูกศรทิศเหนือและมาตราส่วน: ป้ายต้องไม่ทับ
        occupied.append(artist.get_window_extent(renderer).padded(4))
    label_box = dict(boxstyle="round,pad=0.2", fc="white", ec="#DEE5EC", alpha=0.92, lw=0.4)
    placed = 0
    order = drawn.sort_values(["students", "name", "school_id"], kind="stable")
    for index, row in order.iterrows():
        label = " ".join(str(row["name"]).split())  # ยุบช่องว่างซ้ำจากทะเบียน ใช้เฉพาะแสดงผล ไม่ใช้ join
        if not label or label.lower() in {"nan", "none", "true", "false"}:
            raise QualityError(f"โรงเรียน {row['school_id']} ไม่มีชื่อในทะเบียนสำหรับป้ายบนแผนที่ขนาด")
        if len(label) > 24 and "(" in label:
            label = label.replace("(", "\n(", 1)
        ox, oy, radius_px = marker_at[index]
        others = [marker for key, marker in marker_at.items() if key != index]
        radius_pt = radius_px / px_per_pt
        near = radius_pt + 2.5
        offsets = [
            (near, near), (-near, near), (near, -near), (-near, -near),
            (near + 3, 0), (-(near + 3), 0), (0, near + 3), (0, -(near + 3)),
            (near, near + 14), (-near, near + 14), (near, -(near + 14)), (-near, -(near + 14)),
            (near + 14, near), (-(near + 14), near), (near + 14, -near), (-(near + 14), -near),
            (near, near + 30), (-near, near + 30), (near, -(near + 30)), (-near, -(near + 30)),
            (near + 48, near), (-(near + 48), near), (near + 48, -near), (-(near + 48), -near),
            (near, near + 48), (-near, near + 48), (near, -(near + 48)), (-near, -(near + 48)),
        ]
        candidates = []
        for dx, dy in offsets:
            ha = "center" if dx == 0 else ("left" if dx > 0 else "right")
            va = "center" if dy == 0 else ("bottom" if dy > 0 else "top")
            text_item = ax.annotate(
                label, (row.geometry.x, row.geometry.y), xytext=(dx, dy), textcoords="offset points",
                ha=ha, va=va, fontsize=SIZE_MAP_LABEL_FONTSIZE, color="#243447", linespacing=1.3,
                bbox=label_box, zorder=8,
            )
            box = text_item.get_window_extent(renderer).padded(box_pad)
            text_item.remove()
            inside = bbox_cls.intersection(box, map_box)
            outside = box.width * box.height - (inside.width * inside.height if inside else 0)
            overlap = 0.0
            for other in occupied:
                intersection = bbox_cls.intersection(box, other)
                if intersection is not None:
                    overlap += intersection.width * intersection.height
            # เส้นโยงวิ่งจากจุดของตนไปกลางป้าย (หยุดที่ขอบป้าย): ต้องไม่พาดผ่านจุดของโรงเรียนอื่น
            cx, cy = (box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2
            crossings = sum(
                1 for mx, my, mr in others if _segment_point_distance(ox, oy, cx, cy, mx, my) < mr + 1.0
            )
            # ป้ายไม่ควรอยู่ใกล้จุดของโรงเรียนอื่นมากกว่าจุดของตน (ผู้อ่านจะจับคู่ชื่อผิด)
            own_distance = _box_point_distance(box, ox, oy)
            nearer_foreign = sum(1 for mx, my, mr in others if _box_point_distance(box, mx, my) - mr < own_distance)
            score = outside * 100 + overlap * 10 + crossings * 400 + nearer_foreign * 60 + (abs(dx) + abs(dy)) * 0.8
            candidates.append((score, dx, dy, ha, va, box))
            if outside == 0 and overlap == 0 and crossings == 0 and nearer_foreign == 0:
                break  # หยุดค้นเมื่อพบตำแหน่งที่ไม่มีโทษ (offsets เรียงจากใกล้ไปไกล); ตำแหน่งที่ใช้คือคะแนนต่ำสุดในที่ลองมา
        _, dx, dy, ha, va, box = min(candidates, key=lambda item: item[0])
        ax.annotate(
            label, (row.geometry.x, row.geometry.y), xytext=(dx, dy), textcoords="offset points",
            ha=ha, va=va, fontsize=SIZE_MAP_LABEL_FONTSIZE, color="#243447", linespacing=1.3,
            bbox=label_box,
            arrowprops=dict(arrowstyle="-", color="#5B6B7C", lw=0.6, shrinkA=2, shrinkB=radius_pt + 1),
            zorder=8,
        )
        occupied.append(box)
        placed += 1
    return placed


def render_size_map_frames(
    area: dict[str, Any],
    boundary: Any,
    roads: Any,
    small_rows: list[dict[str, Any]],
    output_path: Path,
    profile: dict[str, Any],
    tambon_lines: Any = None,
) -> None:
    """แผนที่โรงเรียนขนาดเล็กจำแนกตามจำนวนนักเรียน รูปแบบเดียวกับแผนที่ผลการจับคู่

    วาดเฉพาะโรงเรียนที่มีนักเรียนไม่เกิน SIZE_MAP_MAX_STUDENTS คน ไล่ทั้งสีและขนาดจุดตามช่วง และติดป้ายชื่อตามทะเบียน
    ให้ทุกโรงเรียนที่วาด (SIZE_MAP_LABELS_VERSION; แผนที่ผลการจับคู่ยังมีป้ายเฉพาะโรงเรียนคุณภาพชุมชน) ไม่แสดงโรงเรียนคุณภาพชุมชน
    ส่วน profile (QC school_size) ยังนับครบทุกช่วงถึง 120 คน (Project Tor.md หัวข้อ 8 และ 18.4; บันทึกการเปลี่ยนแปลง.md §14.6–14.7, §14.10)
    """
    ensure_within(output_path, WORKSPACE, "PNG แผนที่ขนาดโรงเรียน")
    try:
        import matplotlib

        matplotlib.use("Agg", force=True)
        import matplotlib.pyplot as plt
        from matplotlib.lines import Line2D
        from matplotlib.transforms import Bbox
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ matplotlib/GeoPandas สำหรับสร้างแผนที่") from exc

    font_name = _plot_font_name()
    plt.rcParams.update({"font.family": font_name, "axes.unicode_minus": False})
    if boundary is None or getattr(boundary, "crs", None) is None:
        raise QualityError("ชั้นขอบเขตไม่มี CRS")
    if str(boundary.crs).upper() != "EPSG:32647":
        boundary = boundary.to_crs("EPSG:32647")
    if roads is not None and len(roads):
        if getattr(roads, "crs", None) is None:
            raise QualityError("ชั้นถนนไม่มี CRS")
        if str(roads.crs).upper() != "EPSG:32647":
            roads = roads.to_crs("EPSG:32647")
    if not small_rows:
        raise QualityError("ไม่มีโรงเรียนขนาดเล็กสำหรับวาดแผนที่ขนาด")
    points = gpd.GeoDataFrame(
        {
            "students": [int(row["students"]) for row in small_rows],
            "school_id": [str(row["school_id"]) for row in small_rows],
            "name": [str(row.get("school_name", "") or "").strip() for row in small_rows],
        },
        geometry=gpd.points_from_xy(
            [row["lon"] for row in small_rows], [row["lat"] for row in small_rows]
        ),
        crs="EPSG:4326",
    ).to_crs("EPSG:32647")
    points["bin"] = [size_bin_label(int(value)) for value in points["students"]]
    if int(profile["total"]) != len(points):
        raise QualityError("จำนวนโรงเรียนในสรุปขนาดไม่ตรงจุดที่วาด")
    counts = {item["label"]: int(item["count"]) for item in profile["bins"]}
    drawn_bins = size_map_bins()
    drawn = points[points["students"] <= SIZE_MAP_MAX_STUDENTS]
    shown = int(len(drawn))
    if shown != sum(counts.get(f"{low}-{high}", 0) for low, high, _ in drawn_bins):
        raise QualityError("จำนวนโรงเรียนที่วาดไม่ตรงผลรวมช่วงขนาดใน QC")

    extents = [frame.total_bounds for frame in (boundary, drawn) if len(frame)]
    xmin = min(float(bounds[0]) for bounds in extents)
    ymin = min(float(bounds[1]) for bounds in extents)
    xmax = max(float(bounds[2]) for bounds in extents)
    ymax = max(float(bounds[3]) for bounds in extents)
    width = max(xmax - xmin, 1000.0)
    height = max(ymax - ymin, 1000.0)
    pad_x = width * 0.03
    pad_y = height * 0.03
    figure_width, side_margin = 12.0, 0.24
    header_height, footer_height = 0.86, 1.75  # ท้ายภาพ 1.75 นิ้ว ให้คำอธิบายสัญลักษณ์ 2 แถวไม่ชิดข้อความท้ายภาพ 3 บรรทัด
    map_width = figure_width - 2 * side_margin
    map_height = map_width * (height + 2 * pad_y) / (width + 2 * pad_x)
    figure_height = map_height + header_height + footer_height
    if not 3.0 <= figure_height <= 36.0:
        raise QualityError("สัดส่วนพื้นที่ผิดปกติสำหรับแผนที่ขนาดโรงเรียน; ตรวจพิกัดก่อนสร้างภาพ")
    fig = plt.figure(figsize=(figure_width, figure_height), dpi=200)
    ax = fig.add_axes((side_margin / figure_width, footer_height / figure_height,
                       map_width / figure_width, map_height / figure_height))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("#F8FAFC")

    _draw_admin_layers(ax, boundary, roads, tambon_lines)
    # วาดกลุ่มใหญ่ก่อน ให้จุดโรงเรียนเล็กสุดอยู่บนสุดและไม่ถูกบัง
    for low, high, colour in reversed(drawn_bins):
        label = f"{low}-{high}"
        subset = drawn[drawn["bin"] == label]
        if len(subset):
            ax.scatter(
                subset.geometry.x, subset.geometry.y, marker="o", s=SIZE_MAP_MARKER_AREA[label], c=colour,
                edgecolors="#1F2933", linewidths=0.55, zorder=5,
            )

    ax.set_xlim(xmin - pad_x, xmax + pad_x)
    ax.set_ylim(ymin - pad_y, ymax + pad_y)
    ax.set_aspect("equal", adjustable="box")
    ax.axis("off")
    median = float(profile["students_median"])
    median_text = f"{int(median)}" if median.is_integer() else f"{median:.1f}"
    total = int(profile["total"])
    fig.text(side_margin / figure_width, 1 - 0.17 / figure_height,
             f"แผนที่โรงเรียนขนาดเล็กจำแนกตามจำนวนนักเรียน (ไม่เกิน {SIZE_MAP_MAX_STUDENTS} คน) — {area['name']}",
             fontsize=17, color="#17365D", weight="bold", va="top")
    fig.text(side_margin / figure_width, 1 - 0.53 / figure_height,
             f"แสดงเฉพาะโรงเรียนที่มีนักเรียนไม่เกิน {SIZE_MAP_MAX_STUDENTS} คน ({shown} จาก {total} แห่งตามทะเบียน) · "
             f"โรงเรียนขนาดเล็กทั้งหมดมีนักเรียน {int(profile['students_min'])}–{int(profile['students_max'])} คน "
             f"มัธยฐาน {median_text} คน · ชุดข้อมูล DMC682",
             fontsize=11, color="#4B5563", va="top")

    north_arrow = ax.annotate(
        "N", xy=(0.955, 0.97), xytext=(0.955, 0.90), xycoords="axes fraction",
        ha="center", va="center", fontsize=11, weight="bold",
        arrowprops=dict(arrowstyle="-|>", color="#17365D", lw=1.5), zorder=10,
    )
    scale = _nice_scale_length(width)
    sx = xmin - pad_x + (width + 2 * pad_x) * 0.06
    sy = ymin - pad_y + (height + 2 * pad_y) * 0.055
    scale_bar = [
        *ax.plot([sx, sx + scale], [sy, sy], color="#111111", linewidth=2.2, zorder=10),
        *ax.plot([sx, sx], [sy - height * 0.006, sy + height * 0.006], color="#111111", linewidth=1.3, zorder=10),
        *ax.plot([sx + scale, sx + scale], [sy - height * 0.006, sy + height * 0.006], color="#111111", linewidth=1.3, zorder=10),
        ax.text(sx + scale / 2, sy + height * 0.012, f"{scale / 1000:g} กม.", ha="center", va="bottom", fontsize=10,
                bbox=dict(fc="white", ec="none", alpha=0.85, pad=2)),
    ]

    # ป้ายชื่อทุกโรงเรียนที่วาด: วางหลังกำหนดกรอบแผนที่แล้ว เพราะการเลือกตำแหน่งใช้พิกัดหน้าจอจริง
    placed = _place_size_map_labels(fig, ax, drawn, Bbox, decorations=[north_arrow, *scale_bar])
    if placed != shown:
        raise QualityError("จำนวนป้ายชื่อบนแผนที่ขนาดไม่เท่าจำนวนโรงเรียนที่วาด")

    legend_handles = [
        Line2D([0], [0], marker="o", color="none", markerfacecolor=colour, markeredgecolor="#1F2933",
               markeredgewidth=0.55, markersize=math.sqrt(SIZE_MAP_MARKER_AREA[f"{low}-{high}"]),
               label=f"{low}–{high} คน  ({counts.get(f'{low}-{high}', 0)} แห่ง)")
        for low, high, colour in drawn_bins
    ]
    if tambon_lines is not None and len(tambon_lines):
        legend_handles.append(_tambon_legend_handle(Line2D))
    legend = fig.legend(
        handles=legend_handles, loc="upper center", bbox_to_anchor=(0.5, 1.65 / figure_height),
        frameon=False, fontsize=9.5, ncol=4, columnspacing=1.6, labelspacing=0.9, handletextpad=0.8,
        title="จำนวนนักเรียนรวมต่อโรงเรียน — สีและขนาดจุดไล่ตามช่วง (จำนวนโรงเรียนในวงเล็บ)", title_fontsize=9.5,
    )
    legend.get_title().set_color("#243447")
    footer_text = fig.text(
        0.5, 0.16 / figure_height,
        f"จำแนกขนาดจากจำนวนนักเรียนรวมในทะเบียน โรงเรียนที่มีนักเรียน {SIZE_MAP_MAX_STUDENTS + 1}–{SMALL_SCHOOL_MAX_STUDENTS} คน "
        "ไม่แสดงบนแผนที่ ข้อมูลนี้แสดงสภาพขนาดโรงเรียน มิใช่เกณฑ์ตัดสินการควบรวม\n"
        "ป้ายชื่อคือชื่อโรงเรียนตามทะเบียน DMC682 ของทุกโรงเรียนที่วาด ตำแหน่งข้อความเลื่อนได้เพื่อให้อ่านง่าย จุดโรงเรียนไม่เลื่อน\n"
        + map_source_footer(tambon_lines),
        ha="center", va="bottom", fontsize=9.0, linespacing=1.6, color="#4B5563",
    )
    # ข้อความท้ายภาพต้องอยู่ในความกว้างภาพ (เคยล้นทั้งสองข้างในภาพทดสอบ 8 ก.ย. 2569 โดยไม่มีอะไรเตือน)
    fig.canvas.draw()
    footer_extent = footer_text.get_window_extent(fig.canvas.get_renderer())
    if footer_extent.x0 < 0 or footer_extent.x1 > fig.bbox.width:
        raise QualityError("ข้อความท้ายแผนที่ขนาดโรงเรียนล้นความกว้างภาพ; ย่อข้อความก่อนเผยแพร่")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=200, facecolor="white")
    plt.close(fig)
    try:
        from PIL import Image

        with Image.open(output_path) as image:
            if image.width < 2000:
                raise QualityError(f"PNG ขนาดโรงเรียนกว้างเพียง {image.width}px; ต้องอย่างน้อย 2000px")
    except ImportError as exc:
        raise PipelineError("ไม่พบ Pillow สำหรับตรวจ PNG") from exc


def render_standard_size_map(
    area: dict[str, Any],
    boundary_path: Path,
    roads_path: Path,
    small_rows: list[dict[str, Any]],
    output_path: Path,
    profile: dict[str, Any],
    tambon_lines: Any = None,
) -> None:
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas ใน QGIS Python") from exc
    render_size_map_frames(
        area, gpd.read_file(boundary_path), gpd.read_file(roads_path), small_rows, output_path, profile,
        tambon_lines=tambon_lines,
    )


def read_vector_frame(path: Path) -> Any:
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas ใน QGIS Python") from exc
    return gpd.read_file(path)


def public_analysis(area: dict[str, Any], summary: dict[str, Any]) -> str:
    if area["areacode"] == "43010000":
        return (
            "จากโรงเรียนขนาดเล็กตามทะเบียน 99 แห่ง มีผลชั้นเส้นทางเดิม 96 แห่ง "
            "พบ 46 แห่งจับคู่ได้ภายในเกณฑ์ 10 กิโลเมตร 50 แห่งอยู่นอกเกณฑ์ตามผลที่ตรวจรับ "
            "และมีกรณีพิกัด–ขอบเขตไม่สอดคล้อง 3 แห่ง ผลพื้นที่นี้นำเข้าจาก Excel ที่ตรวจรับแล้ว "
            "และมี NetworkX fallback บางรายการ จึงไม่ควรอธิบายว่าทั้ง 99 แห่งมาจาก QNEAT3 ชุดเดียวกัน"
        )
    text = (
        f"จากโรงเรียนต้นทางที่วิเคราะห์ {summary['origins']} แห่ง พบ "
        f"{summary['matched']} แห่งจับคู่ได้ภายในเกณฑ์ 10 กิโลเมตร และ "
        f"{summary['over10km']} แห่งอยู่นอกเกณฑ์"
    )
    review = int(summary.get("unreachable", 0)) + int(summary.get("coordExceptions", 0))
    if review:
        text += f" โดยมีกรณีที่ต้องตรวจข้อมูลหรือการเชื่อมต่อโครงข่าย {review} แห่ง"
    if int(summary.get("excluded", 0)):
        text += f" พร้อมแยกกรณีเป็นโรงเรียนคุณภาพชุมชนด้วย {summary['excluded']} แห่ง"
    return text + " ผลนี้เป็นเครื่องมือคัดกรองเบื้องต้น มิใช่ข้อเสนอควบรวมหรือย้ายนักเรียน"


def _reportlab_fonts() -> tuple[str, str]:
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError as exc:
        raise PipelineError("ไม่พบ ReportLab สำหรับสร้าง PDF") from exc
    font_root = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    regular = font_root / "tahoma.ttf"
    bold = font_root / "tahomabd.ttf"
    if regular.is_file():
        pdfmetrics.registerFont(TTFont("SMS-Thai", str(regular)))
        if bold.is_file():
            pdfmetrics.registerFont(TTFont("SMS-Thai-Bold", str(bold)))
            pdfmetrics.registerFontFamily(
                "SMS-Thai", normal="SMS-Thai", bold="SMS-Thai-Bold", italic="SMS-Thai", boldItalic="SMS-Thai-Bold"
            )
            return "SMS-Thai", "SMS-Thai-Bold"
        return "SMS-Thai", "SMS-Thai"
    return "Helvetica", "Helvetica-Bold"


def render_summary_pdf(
    area: dict[str, Any],
    summary: dict[str, Any],
    image_path: Path,
    output_path: Path,
    size_image_path: Path | None = None,
    consolidation: dict[str, Any] | None = None,
    size_profile: dict[str, Any] | None = None,
) -> None:
    ensure_within(output_path, WORKSPACE, "PDF สรุป")
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen import canvas
        from reportlab.platypus import Paragraph, Table, TableStyle
        from reportlab.lib.enums import TA_RIGHT
        from xml.sax.saxutils import escape as xml_escape
        from PIL import Image
    except ImportError as exc:
        raise PipelineError("ไม่พบ ReportLab/Pillow สำหรับสร้าง PDF") from exc

    regular, bold = _reportlab_fonts()
    page_width, page_height = A4
    output_path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(output_path), pagesize=A4)
    pdf.setTitle(f"สรุปผลการจับคู่ {area['name']}")
    pdf.setAuthor("โครงการวิเคราะห์การจับคู่โรงเรียนขนาดเล็ก")
    # Give the map its own page; squeezing it between metrics and prose made
    # school labels unreadable. Page two uses the same summary as the portal.
    with Image.open(image_path) as source_image:
        iw, ih = source_image.size
    map_page_width, map_page_height = landscape(A4) if iw > ih else A4
    pdf.setPageSize((map_page_width, map_page_height))
    map_margin = 16
    map_scale = min((map_page_width - 2 * map_margin) / iw,
                    (map_page_height - 2 * map_margin) / ih)
    draw_w, draw_h = iw * map_scale, ih * map_scale
    pdf.drawImage(ImageReader(str(image_path)), (map_page_width - draw_w) / 2,
                  (map_page_height - draw_h) / 2, width=draw_w, height=draw_h,
                  preserveAspectRatio=True, mask="auto")
    pdf.showPage()
    pdf.setPageSize(A4)
    margin = 38
    pdf.setFillColor(colors.HexColor("#17365D"))
    pdf.setFont(bold, 16)
    pdf.drawString(margin, page_height - 48, f"สรุปผลการจับคู่ — {area['name']}")
    pdf.setFont(regular, 8.5)
    pdf.setFillColor(colors.HexColor("#4B5563"))
    pdf.drawString(margin, page_height - 65, "ฉบับนำร่องเพื่อการตรวจสอบ | ชุดข้อมูล DMC682")
    pdf.drawRightString(page_width - margin, page_height - 65, f"รหัสพื้นที่ {area['areacode']}")

    metrics = [
        ("ต้นทางวิเคราะห์", int(summary.get("origins", 0))),
        ("โรงเรียนคุณภาพชุมชน", int(summary.get("magnets", 0))),
        ("จับคู่ได้ ≤10 กม.", int(summary.get("matched", 0))),
        ("นอกเกณฑ์", int(summary.get("over10km", 0))),
        (
            "ต้องตรวจ/แยก",
            int(summary.get("unreachable", 0))
            + int(summary.get("excluded", 0))
            + int(summary.get("coordExceptions", 0)),
        ),
    ]
    card_gap = 6
    card_width = (page_width - 2 * margin - card_gap * 4) / 5
    card_y = page_height - 112
    for index, (label, value) in enumerate(metrics):
        x = margin + index * (card_width + card_gap)
        pdf.setFillColor(colors.HexColor("#EEF4F9"))
        pdf.roundRect(x, card_y, card_width, 35, 4, fill=1, stroke=0)
        pdf.setFillColor(colors.HexColor("#17365D"))
        pdf.setFont(bold, 13)
        pdf.drawCentredString(x + card_width / 2, card_y + 18, f"{value:,}")
        pdf.setFont(regular, 6.7)
        pdf.drawCentredString(x + card_width / 2, card_y + 6, label)

    max_w = page_width - 2 * margin

    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "ThaiBody",
        parent=styles["BodyText"],
        fontName=regular,
        fontSize=11,
        leading=18,
        textColor=colors.HexColor("#263238"),
        splitLongWords=True,
    )
    note_y = card_y - 30
    analysis = Paragraph(public_analysis(area, summary), body_style)
    _, analysis_h = analysis.wrap(max_w, 80)
    analysis.drawOn(pdf, margin, note_y - analysis_h)
    note_y -= analysis_h + 22
    limitation = Paragraph(
        "<b>วิธีอ่านและข้อจำกัด:</b> เส้นเชื่อมบนแผนที่แสดงความสัมพันธ์ตามคู่ที่ผ่านเกณฑ์ "
        "มิใช่แนวเส้นทางเดินรถจริง ผลนี้ใช้คัดกรองเบื้องต้นและต้องพิจารณาความพร้อมของพื้นที่ "
        "ความจุโรงเรียน การเดินทางจริง ภูมิประเทศ และความเห็นของชุมชนร่วมด้วย "
        "ผลการจับคู่เป็นเครื่องมือคัดกรอง มิใช่ข้อเสนอควบรวมหรือย้ายนักเรียน",
        body_style,
    )
    _, limitation_h = limitation.wrap(max_w, 70)
    limitation.drawOn(pdf, margin, note_y - limitation_h)
    if note_y - limitation_h < 45:
        raise QualityError("ข้อความสรุปยาวเกินหน้ากระดาษ PDF; หยุดก่อนเผยแพร่")

    def draw_footer() -> None:
        pdf.setStrokeColor(colors.HexColor("#D5DCE3"))
        pdf.line(margin, 28, page_width - margin, 28)
        pdf.setFont(regular, 6.8)
        pdf.setFillColor(colors.HexColor("#5F6B76"))
        pdf.drawString(margin, 17, "วิธี: แผน 1.2 — ระยะทางตามโครงข่ายถนนไม่เกิน 10 กิโลเมตร")
        pdf.drawRightString(page_width - margin, 17, "ที่มาถนน: © OpenStreetMap contributors")

    if consolidation is not None:
        # ส่วนภาพจำลองหลังจับคู่: ข้อความสรุป (ส่วน A) และตารางรายโรงเรียนคุณภาพชุมชน (ส่วน B) ตัวเลขชุดเดียวกับเว็บ
        agg = consolidation["aggregate"]

        def fmt_km(value: Any) -> str:
            return f"{float(value):.2f}" if value is not None else "—"

        y = note_y - limitation_h - 20
        pdf.setFillColor(colors.HexColor("#17365D"))
        pdf.setFont(bold, 11.5)
        pdf.drawString(margin, y - 11, "ภาพจำลองหากจับคู่ตามผลคัดกรอง (ภาพจำลองประกอบการพิจารณา ไม่ใช่แผนควบรวม)")
        y -= 22
        small_style = ParagraphStyle("ThaiSmall", parent=body_style, fontSize=9.5, leading=14)
        block_a = Paragraph(
            f"โรงเรียน สพป. ในเขต {agg['schools_before']:,} → {agg['schools_after']:,} แห่ง · "
            f"โรงเรียนขนาดเล็กที่จับคู่ได้ {agg['small_matched']:,} แห่ง ยังคงอยู่ {agg['small_remaining']:,} แห่ง · "
            f"นักเรียนที่ย้ายตามภาพจำลอง {agg['students_moved']:,} คน<br/>"
            f"โรงเรียนคุณภาพชุมชนที่รับนักเรียน {agg['magnets_receiving']} จาก {agg['magnets_total']} แห่ง · "
            f"ระยะทางคู่ที่จับได้เฉลี่ย {fmt_km(agg['distance_mean_km'])} กม. ไกลสุด {fmt_km(agg['distance_max_km'])} กม. ตามโครงข่ายถนน<br/>"
            + xml_escape(public_consolidation_note(consolidation)),
            small_style,
        )
        _, block_h = block_a.wrap(max_w, 160)
        block_a.drawOn(pdf, margin, y - block_h)
        y -= block_h + 10
        cell = ParagraphStyle("ThaiCell", parent=body_style, fontSize=8.5, leading=11)
        cell_right = ParagraphStyle("ThaiCellRight", parent=cell, alignment=TA_RIGHT)
        head = ParagraphStyle("ThaiHead", parent=cell, fontName=bold, fontSize=8, leading=10,
                              textColor=colors.HexColor("#17365D"))
        head_right = ParagraphStyle("ThaiHeadRight", parent=head, alignment=TA_RIGHT)
        headers = ("โรงเรียนคุณภาพชุมชน", "นักเรียน<br/>ก่อน", "รร.เล็ก<br/>ที่รับ", "นักเรียน<br/>ที่รับเพิ่ม",
                   "นักเรียน<br/>หลังรวม", "ขนาด 7 ระดับ<br/>ก่อน → หลัง", "ไกลสุด<br/>(กม.)")
        right_cols = {1, 2, 3, 4, 6}
        rows = [[Paragraph(text, head_right if index in right_cols else head) for index, text in enumerate(headers)]]
        for school in consolidation["magnets"]:
            name = xml_escape(str(school["school_name"])) + (" †" if school.get("is_small_school") else "")
            rows.append([
                Paragraph(name, cell),
                Paragraph(f"{school['students_before']:,}", cell_right),
                Paragraph(f"{school['feeders']:,}", cell_right),
                Paragraph(f"{school['students_received']:,}", cell_right),
                Paragraph(f"{school['students_after']:,}", cell_right),
                Paragraph(f"{size7_label(school['size_code_before'])} → {size7_label(school['size_code_after'])}", cell),
                Paragraph(fmt_km(school['farthest_km']) if school["feeders"] else "—", cell_right),
            ])
        rows.append([
            Paragraph("รวม / ไกลสุด", head), Paragraph(f"{agg['magnets_students_before']:,}", head_right),
            Paragraph(f"{agg['small_matched']:,}", head_right), Paragraph(f"{agg['students_moved']:,}", head_right),
            Paragraph(f"{agg['magnets_students_after']:,}", head_right), Paragraph("", head),
            Paragraph(fmt_km(agg['distance_max_km']), head_right),
        ])
        col_widths = [126, 52, 50, 62, 58, 118, max_w - 466]
        table = Table(rows, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF4F9")),
            ("LINEBELOW", (0, 0), (-1, 0), 0.7, colors.HexColor("#9FB3C8")),
            ("LINEBELOW", (0, 1), (-1, -2), 0.3, colors.HexColor("#D5DCE3")),
            ("LINEABOVE", (0, -1), (-1, -1), 0.7, colors.HexColor("#9FB3C8")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        footnote = (
            "† โรงเรียนขนาดเล็กตามทะเบียนที่มีบทบาทโรงเรียนคุณภาพชุมชนด้วย  |  ขนาด 7 ระดับตามเกณฑ์ สพฐ. ในทะเบียน DMC682  |  "
            f"ตารางนี้ไม่แสดงชื่อโรงเรียนขนาดเล็กต้นทาง (ชื่อโรงเรียนขนาดเล็กไม่เกิน {SIZE_MAP_MAX_STUDENTS} คนอยู่บนแผนที่ขนาดโรงเรียนหน้าถัดไป)"
        )
        footnote_style = ParagraphStyle("ThaiFoot", parent=body_style, fontSize=7.5, leading=10,
                                        textColor=colors.HexColor("#5F6B76"))
        note_par = Paragraph(footnote, footnote_style)
        _, footnote_h = note_par.wrap(max_w, 40)
        available = y - 45 - footnote_h - 6
        _, table_h = table.wrap(max_w, available)
        if len(consolidation["magnets"]) > 10 or table_h > available:
            draw_footer()
            pdf.showPage()
            pdf.setPageSize(A4)
            pdf.setFillColor(colors.HexColor("#17365D"))
            pdf.setFont(bold, 12)
            pdf.drawString(margin, page_height - 48, f"ภาพจำลองหากจับคู่ตามผลคัดกรอง (ต่อ) — {area['name']}")
            y = page_height - 70
            available = y - 45 - footnote_h - 6
            _, table_h = table.wrap(max_w, available)
            if table_h > available:
                raise QualityError("ตารางภาพจำลองยาวเกินหนึ่งหน้า; หยุดก่อนเผยแพร่")
        table.drawOn(pdf, margin, y - table_h)
        note_par.drawOn(pdf, margin, y - table_h - 6 - footnote_h)
    draw_footer()
    pdf.showPage()
    if size_image_path is not None:
        # หน้าที่ 3: แผนที่ขนาดโรงเรียนเต็มหน้า ใช้กติกาเดียวกับหน้าแผนที่จับคู่
        with Image.open(size_image_path) as size_image:
            sw, sh = size_image.size
        size_page_width, size_page_height = landscape(A4) if sw > sh else A4
        pdf.setPageSize((size_page_width, size_page_height))
        size_scale = min((size_page_width - 2 * map_margin) / sw,
                         (size_page_height - 2 * map_margin) / sh)
        size_w, size_h = sw * size_scale, sh * size_scale
        pdf.drawImage(ImageReader(str(size_image_path)), (size_page_width - size_w) / 2,
                      (size_page_height - size_h) / 2, width=size_w, height=size_h,
                      preserveAspectRatio=True, mask="auto")
        pdf.showPage()
    if size_profile is not None:
        # หน้าที่ 4: คำอธิบายขนาดโรงเรียนขนาดเล็ก ตัวเลขชุดเดียวกับการ์ด "สภาพโรงเรียนขนาดเล็กในพื้นที่" บนเว็บ
        # (ตัดสินร่วมกับผู้ใช้ 14 ก.ย. 2569 — บันทึกการเปลี่ยนแปลง.md §14.15)
        pdf.setPageSize(A4)
        pdf.setFillColor(colors.HexColor("#17365D"))
        pdf.setFont(bold, 16)
        pdf.drawString(margin, page_height - 48, f"สภาพโรงเรียนขนาดเล็กในพื้นที่ — {area['name']}")
        pdf.setFont(regular, 8.5)
        pdf.setFillColor(colors.HexColor("#4B5563"))
        pdf.drawString(margin, page_height - 65, "ฉบับนำร่องเพื่อการตรวจสอบ | ชุดข้อมูล DMC682 | คำอธิบายประกอบแผนที่ขนาดโรงเรียนหน้าก่อนหน้า")
        pdf.drawRightString(page_width - margin, page_height - 65, f"รหัสพื้นที่ {area['areacode']}")
        size_total = int(size_profile["total"])
        size_counts = {str(item["label"]): int(item["count"]) for item in size_profile["bins"]}
        drawn_bins = size_map_bins()
        size_shown = sum(size_counts.get(f"{low}-{high}", 0) for low, high, _ in drawn_bins)
        size_hidden = size_total - size_shown
        size_median = float(size_profile["students_median"])
        size_median_text = f"{int(size_median)}" if size_median.is_integer() else f"{size_median:.1f}"
        size_metrics = [
            ("โรงเรียนขนาดเล็กตามทะเบียน", f"{size_total:,}"),
            ("นักเรียนรวม (คน)", f"{int(size_profile['students_total']):,}"),
            ("น้อยสุด–มากสุด (คน)", f"{int(size_profile['students_min'])}–{int(size_profile['students_max'])}"),
            ("มัธยฐาน (คน)", size_median_text),
            (f"แสดงบนแผนที่ (ไม่เกิน {SIZE_MAP_MAX_STUDENTS} คน)", f"{size_shown:,}"),
        ]
        for index, (label, value) in enumerate(size_metrics):
            x = margin + index * (card_width + card_gap)
            pdf.setFillColor(colors.HexColor("#EEF4F9"))
            pdf.roundRect(x, card_y, card_width, 35, 4, fill=1, stroke=0)
            pdf.setFillColor(colors.HexColor("#17365D"))
            pdf.setFont(bold, 13)
            pdf.drawCentredString(x + card_width / 2, card_y + 18, value)
            pdf.setFont(regular, 6.7)
            pdf.drawCentredString(x + card_width / 2, card_y + 6, label)
        y = card_y - 30
        size_par = Paragraph(xml_escape(public_size_analysis(size_profile)), body_style)
        _, size_par_h = size_par.wrap(max_w, 140)
        size_par.drawOn(pdf, margin, y - size_par_h)
        y -= size_par_h + 18
        pdf.setFillColor(colors.HexColor("#17365D"))
        pdf.setFont(bold, 11.5)
        pdf.drawString(margin, y - 11, "จำนวนโรงเรียนขนาดเล็กต่อช่วงจำนวนนักเรียนรวม (ตัวเลขรวมระดับพื้นที่ ไม่มีรายโรงเรียน)")
        y -= 22

        def size_percent(value: int) -> str:
            return f"{value / size_total * 100:.1f}" if size_total else "0.0"

        size_cell = ParagraphStyle("ThaiSizeCell", parent=body_style, fontSize=9, leading=12)
        size_cell_right = ParagraphStyle("ThaiSizeCellRight", parent=size_cell, alignment=TA_RIGHT)
        size_head = ParagraphStyle("ThaiSizeHead", parent=size_cell, fontName=bold, fontSize=8.5, leading=11,
                                   textColor=colors.HexColor("#17365D"))
        size_head_right = ParagraphStyle("ThaiSizeHeadRight", parent=size_head, alignment=TA_RIGHT)
        size_headers = ("สี", "ช่วงนักเรียน (คน)", "จำนวนโรงเรียน", "ร้อยละ", "บนแผนที่ขนาดโรงเรียน")
        size_rows = [[Paragraph(h, size_head_right if i in (2, 3) else size_head) for i, h in enumerate(size_headers)]]
        swatch_rows = []
        for low, high, colour in drawn_bins:
            count = size_counts.get(f"{low}-{high}", 0)
            swatch_rows.append((len(size_rows), colour))
            size_rows.append([
                Paragraph("", size_cell), Paragraph(f"{low}–{high}", size_cell),
                Paragraph(f"{count:,}", size_cell_right), Paragraph(size_percent(count), size_cell_right),
                Paragraph("แสดง (สีและขนาดจุดตามช่วง)", size_cell),
            ])
        grey_row = len(size_rows)
        size_rows.append([
            Paragraph("", size_cell), Paragraph(f"{SIZE_MAP_MAX_STUDENTS + 1}–{SMALL_SCHOOL_MAX_STUDENTS}", size_cell),
            Paragraph(f"{size_hidden:,}", size_cell_right), Paragraph(size_percent(size_hidden), size_cell_right),
            Paragraph("ไม่แสดงบนแผนที่ (นับรวมในตาราง)", size_cell),
        ])
        size_rows.append([
            Paragraph("", size_head), Paragraph("รวม", size_head), Paragraph(f"{size_total:,}", size_head_right),
            Paragraph("100.0" if size_total else "0.0", size_head_right), Paragraph(f"แสดง {size_shown:,} แห่ง", size_head),
        ])
        size_style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF4F9")),
            ("LINEBELOW", (0, 0), (-1, 0), 0.7, colors.HexColor("#9FB3C8")),
            ("LINEBELOW", (0, 1), (-1, -2), 0.3, colors.HexColor("#D5DCE3")),
            ("LINEABOVE", (0, -1), (-1, -1), 0.7, colors.HexColor("#9FB3C8")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("BACKGROUND", (0, grey_row), (0, grey_row), colors.HexColor("#C9CED6")),
            ("TEXTCOLOR", (1, grey_row), (-1, grey_row), colors.HexColor("#5F6B76")),
        ]
        for row_index, colour in swatch_rows:
            size_style.append(("BACKGROUND", (0, row_index), (0, row_index), colors.HexColor(colour)))
        size_table = Table(size_rows, colWidths=[28, 112, 92, 70, max_w - 302], repeatRows=1)
        size_table.setStyle(TableStyle(size_style))
        _, size_table_h = size_table.wrap(max_w, y - 60)
        size_table.drawOn(pdf, margin, y - size_table_h)
        y -= size_table_h + 18
        pdf.setFillColor(colors.HexColor("#17365D"))
        pdf.setFont(bold, 11.5)
        pdf.drawString(margin, y - 11, "วิธีอ่านแผนที่ขนาดโรงเรียน (หน้าก่อนหน้า)")
        y -= 22
        size_guide_style = ParagraphStyle("ThaiSizeGuide", parent=body_style, fontSize=9.5, leading=14.5)
        size_guide = Paragraph(
            "• สีและขนาดจุดไล่ตามช่วงจำนวนนักเรียน นักเรียนมาก = จุดใหญ่ ตามคำอธิบายสัญลักษณ์ใต้แผนที่<br/>"
            "• ป้ายชื่อคือชื่อโรงเรียนตามทะเบียน DMC682 ของทุกโรงเรียนที่วาด ตำแหน่งข้อความเลื่อนได้เพื่อให้อ่านง่าย จุดโรงเรียนไม่เลื่อน<br/>"
            f"• โรงเรียนที่มีนักเรียน {SIZE_MAP_MAX_STUDENTS + 1}–{SMALL_SCHOOL_MAX_STUDENTS} คน ไม่วาดบนแผนที่แต่นับรวมในตารางข้างบน "
            "และไม่แสดงโรงเรียนคุณภาพชุมชน (ดูแผนที่ผลการจับคู่หน้าที่ 1)<br/>"
            "• เส้นประคือแนวเขตตำบล (tha_admin3) เส้นทึบสีน้ำเงินเข้มคือขอบเขต สพป.<br/>"
            f"• เกณฑ์ 7 ขนาดของ สพฐ. จัดโรงเรียนที่มีนักเรียนไม่เกิน {SMALL_SCHOOL_MAX_STUDENTS} คนเป็นขนาดที่ 1 (แบ่งย่อย 1.1–1.6 ช่วงละ 20 คน) "
            "ช่วง 6 ช่วงบนแผนที่เป็นการแบ่งของโครงการที่ละเอียดกว่าช่วง 1.1–1.3<br/>"
            "• ข้อมูลนี้แสดงสภาพขนาดโรงเรียนตามทะเบียน มิใช่เกณฑ์ตัดสินการควบรวมหรือย้ายนักเรียน",
            size_guide_style,
        )
        _, size_guide_h = size_guide.wrap(max_w, 200)
        if y - size_guide_h < 45:
            raise QualityError("หน้าคำอธิบายขนาดโรงเรียนใน PDF ยาวเกินหน้ากระดาษ; หยุดก่อนเผยแพร่")
        size_guide.drawOn(pdf, margin, y - size_guide_h)
        draw_footer()
        pdf.showPage()
    pdf.save()
    with output_path.open("rb") as handle:
        if handle.read(5) != b"%PDF-":
            raise QualityError(f"PDF ไม่สมบูรณ์: {output_path}")


def _plain_file_fingerprint(path: Path) -> dict[str, Any]:
    exact = ensure_file(path, "ไฟล์ต้นทางพื้นที่นำร่อง")
    return {
        "filename": exact.name,
        "bytes": exact.stat().st_size,
        "sha256": sha256_file(exact),
    }


def _frame_epsg(frame: Any) -> int | None:
    crs = getattr(frame, "crs", None)
    if crs is None:
        return None
    try:
        return crs.to_epsg()
    except AttributeError:
        text_value = str(crs).upper()
        return 32647 if text_value == "EPSG:32647" else None


def _require_pilot_crs(frame: Any, label: str) -> None:
    if _frame_epsg(frame) != 32647:
        raise QualityError(f"{label} ต้องเป็น EPSG:32647 แต่พบ {getattr(frame, 'crs', None)}")


def _value_as_school_id(value: Any) -> str:
    if value is None:
        return ""
    text_value = str(value).strip()
    if not text_value or text_value.lower() == "nan":
        return ""
    try:
        return str(int(float(text_value)))
    except (TypeError, ValueError, OverflowError):
        return text_value


def _frame_id_set(frame: Any, field: str, label: str) -> set[str]:
    if field not in frame.columns:
        raise QualityError(f"{label} ไม่มี field {field}")
    values = [_value_as_school_id(value) for value in frame[field].tolist()]
    if any(not value for value in values):
        raise QualityError(f"{label} มีรหัสโรงเรียนว่าง")
    if len(values) != len(set(values)):
        raise QualityError(f"{label} มีรหัสโรงเรียนซ้ำ")
    return set(values)


def _pilot_file(area: dict[str, Any], key: str) -> Path:
    files = area.get("pilot", {}).get("files", {})
    name = files.get(key)
    if not isinstance(name, str) or Path(name).name != name:
        raise QualityError(f"areas.json ไม่มีชื่อไฟล์นำร่องที่ปลอดภัยสำหรับ {area['name']} ({key})")
    area_dir = ensure_within(AREA_ROOT / area["folder_name"], AREA_ROOT, "โฟลเดอร์นำร่อง")
    return ensure_file(area_dir / layout.area_relpath(key, name), f"ไฟล์นำร่อง {key}")


def _magnet_points_from_catalog(area: dict[str, Any]) -> Any:
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas ใน QGIS Python") from exc
    _, magnet_rows, _ = load_area_records(area)
    frame = gpd.GeoDataFrame(
        {
            "school_code": [row["school_id"] for row in magnet_rows],
            "school_name": [row["school_name"] for row in magnet_rows],
        },
        geometry=gpd.points_from_xy(
            [row["lon"] for row in magnet_rows],
            [row["lat"] for row in magnet_rows],
        ),
        crs="EPSG:4326",
    ).to_crs("EPSG:32647")
    return frame


def load_qneat3_pilot_frames(
    area: dict[str, Any],
) -> tuple[dict[str, Any], dict[str, int], dict[str, Any]]:
    try:
        import geopandas as gpd
        import pandas as pd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas/Pandas ใน QGIS Python") from exc

    paths = {
        key: _pilot_file(area, key)
        for key in ("boundary", "small", "magnets", "roads", "od", "status")
    }
    frames = {key: gpd.read_file(path) for key, path in paths.items()}
    for key, frame in frames.items():
        _require_pilot_crs(frame, f"ชั้นนำร่อง {key}")

    expected = area["pilot"]["expected"]
    if len(frames["boundary"]) != 1:
        raise QualityError(f"ขอบเขตนำร่อง {area['name']} ต้องมี 1 feature")
    if len(frames["small"]) != int(expected["origins"]):
        raise QualityError("จำนวนชั้นโรงเรียนต้นทางนำร่องไม่ตรงค่าตรวจรับ")
    if len(frames["status"]) != int(expected["origins"]):
        raise QualityError("จำนวนชั้นสถานะนำร่องไม่ตรงค่าตรวจรับ")
    if len(frames["magnets"]) != int(expected["magnets"]):
        raise QualityError("จำนวนชั้นโรงเรียนคุณภาพชุมชนนำร่องไม่ตรงค่าตรวจรับ")

    status_ids = _frame_id_set(frames["status"], "โรงเรียน", "ชั้นสถานะนำร่อง")
    od_ids = _frame_id_set(frames["od"], "origin_id", "OD ผ่านเกณฑ์นำร่อง")
    if not od_ids.issubset(status_ids):
        raise QualityError("OD นำร่องมี origin ที่ไม่อยู่ในชั้นสถานะ")
    if len(od_ids) != int(expected["matched"]):
        raise QualityError("จำนวน origin ใน OD นำร่องไม่ตรงค่าตรวจรับ")
    if "total_cost" in frames["od"].columns:
        costs = pd.to_numeric(frames["od"]["total_cost"], errors="coerce")
        if costs.isna().any() or (costs > 10000.0).any() or (costs < 0).any():
            raise QualityError("OD ที่ระบุว่าผ่านเกณฑ์มี total_cost ว่างหรือนอกช่วง 0–10,000 เมตร")

    status_frame = frames["status"].copy()
    status_frame["สถานะ"] = [
        "MATCHED_10KM" if school_id in od_ids else "OVER_10KM"
        for school_id in status_frame["โรงเรียน"].map(_value_as_school_id)
    ]
    summary = {
        "registered": int(expected["registered"]),
        "origins": len(status_frame),
        "magnets": len(frames["magnets"]),
        "matched": len(od_ids),
        "over10km": len(status_frame) - len(od_ids),
        "unreachable": 0,
        "excluded": int(expected["excluded"]),
        "coordExceptions": 0,
    }
    validate_expected_summary(area, summary)

    # Use the official central coordinates/names for labels.  This also avoids
    # the corrupted Boolean school_name type in one legacy magnet GPKG.
    display_magnets = _magnet_points_from_catalog(area)
    frames_for_map = {
        "boundary": frames["boundary"],
        "roads": frames["roads"],
        "od": frames["od"],
        "status": status_frame,
        "magnets": display_magnets,
    }
    inputs = {key: _plain_file_fingerprint(path) for key, path in paths.items()}
    inputs["road_feature_count"] = len(frames["roads"])
    return frames_for_map, summary, inputs


def _read_nong_khai_status_excel(path: Path) -> dict[str, str]:
    try:
        import openpyxl
    except ImportError as exc:
        raise PipelineError("ไม่พบ openpyxl สำหรับอ่านผลหนองคายที่ตรวจรับแล้ว") from exc
    workbook = openpyxl.load_workbook(
        ensure_file(path, "Excel หนองคายที่ตรวจรับแล้ว"),
        read_only=True,
        data_only=True,
    )
    try:
        if "สถานะครบ 99 แห่ง" not in workbook.sheetnames:
            raise QualityError("Excel หนองคายไม่มี worksheet สถานะครบ 99 แห่ง")
        sheet = workbook["สถานะครบ 99 แห่ง"]
        header = {
            str(cell.value).strip(): index
            for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)), start=1)
            if cell.value is not None
        }
        id_column = header.get("รหัส รร.เล็ก")
        status_column = header.get("สถานะตามแผน 1.2")
        if not id_column or not status_column:
            raise QualityError("Excel หนองคายไม่มีคอลัมน์รหัส/สถานะที่กำหนด")
        mapping: dict[str, str] = {}
        labels = {
            "จับคู่ได้ภายในเกณฑ์": "MATCHED_10KM",
            "อยู่นอกเกณฑ์—มีผลเส้นทาง": "OVER_10KM",
            "อยู่นอกเกณฑ์—กรณีพิกัด–ขอบเขต": "COORD_BOUNDARY_EXCEPTION",
        }
        for row in sheet.iter_rows(min_row=2, values_only=True):
            school_id = _value_as_school_id(row[id_column - 1])
            raw_status = str(row[status_column - 1] or "").strip()
            if not school_id:
                continue
            if school_id in mapping:
                raise QualityError(f"Excel หนองคายมีรหัสซ้ำ: {school_id}")
            if raw_status not in labels:
                raise QualityError(f"Excel หนองคายมีสถานะที่ไม่รองรับ: {raw_status!r}")
            mapping[school_id] = labels[raw_status]
        return mapping
    finally:
        workbook.close()


def load_legacy_nong_khai_frames(
    machine: dict[str, Any], area: dict[str, Any]
) -> tuple[dict[str, Any], dict[str, int], dict[str, Any]]:
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas ใน QGIS Python") from exc

    small_rows, _, _ = load_area_records(area)
    status_mapping = _read_nong_khai_status_excel(_pilot_file(area, "excel"))
    small_ids = {row["school_id"] for row in small_rows}
    if set(status_mapping) != small_ids:
        missing = sorted(small_ids - set(status_mapping))
        extra = sorted(set(status_mapping) - small_ids)
        raise QualityError(
            f"Excel หนองคายไม่ตรงประชากร DMC: missing={missing[:5]}, extra={extra[:5]}"
        )
    status_frame = gpd.GeoDataFrame(
        {
            "โรงเรียน": [row["school_id"] for row in small_rows],
            "ชื่อโรงเรียน": [row["school_name"] for row in small_rows],
            "สถานะ": [status_mapping[row["school_id"]] for row in small_rows],
        },
        geometry=gpd.points_from_xy(
            [row["lon"] for row in small_rows],
            [row["lat"] for row in small_rows],
        ),
        crs="EPSG:4326",
    ).to_crs("EPSG:32647")

    roads_path = _pilot_file(area, "roads")
    od_path = _pilot_file(area, "od")
    roads = gpd.read_file(roads_path)
    od_lines = gpd.read_file(od_path)
    _require_pilot_crs(roads, "ถนน legacy หนองคาย")
    _require_pilot_crs(od_lines, "OD legacy หนองคาย")
    od_ids = _frame_id_set(od_lines, "origin_id", "OD legacy หนองคาย")
    matched_ids = {
        school_id for school_id, status in status_mapping.items() if status == "MATCHED_10KM"
    }
    if od_ids != matched_ids:
        raise QualityError("OD 46 คู่ของหนองคายไม่ตรงกลุ่มผ่านเกณฑ์ใน Excel")

    boundary_path = exact_external_paths(machine)["education_boundaries"]
    escaped_name = area["boundary_name"].replace("'", "''")
    try:
        boundary = gpd.read_file(
            boundary_path,
            where=f'"{BOUNDARY_NAME_FIELD}" = \'{escaped_name}\'',
        )
    except (TypeError, ValueError, OSError):
        # Safe fallback: this reads only the one exact boundary shapefile, not
        # the external folder and never the nationwide roads layer.
        all_boundaries = gpd.read_file(boundary_path)
        if BOUNDARY_NAME_FIELD not in all_boundaries.columns:
            raise QualityError(f"ขอบเขต สพป. ไม่มี field {BOUNDARY_NAME_FIELD}")
        boundary = all_boundaries[
            all_boundaries[BOUNDARY_NAME_FIELD].astype(str).str.strip()
            == area["boundary_name"]
        ].copy()
    if len(boundary) != 1:
        raise QualityError(f"ขอบเขตหนองคายต้องพบ 1 feature แต่พบ {len(boundary)}")
    boundary = boundary.to_crs("EPSG:32647")

    magnets = _magnet_points_from_catalog(area)
    expected = area["pilot"]["expected"]
    summary = {
        "registered": len(status_frame),
        "origins": int(expected["origins"]),
        "magnets": len(magnets),
        "matched": int((status_frame["สถานะ"] == "MATCHED_10KM").sum()),
        "over10km": int((status_frame["สถานะ"] == "OVER_10KM").sum()),
        "unreachable": 0,
        "excluded": 0,
        "coordExceptions": int(
            (status_frame["สถานะ"] == "COORD_BOUNDARY_EXCEPTION").sum()
        ),
    }
    validate_expected_summary(area, summary)
    input_paths = {
        "roads": roads_path,
        "od": od_path,
        "small_legacy": _pilot_file(area, "small"),
        "excel": _pilot_file(area, "excel"),
    }
    inputs = {key: _plain_file_fingerprint(path) for key, path in input_paths.items()}
    inputs["education_boundaries"] = shapefile_fingerprint(boundary_path)
    inputs["lineage_note"] = (
        "หมวดหมู่มาจาก Excel ที่ตรวจรับแล้ว; ไม่รัน QNEAT3 ซ้ำและมี NetworkX fallback ในผลเดิม"
    )
    return {
        "boundary": boundary,
        "roads": roads,
        "od": od_lines,
        "status": status_frame,
        "magnets": magnets,
    }, summary, inputs


def load_pilot_frames(
    machine: dict[str, Any], area: dict[str, Any]
) -> tuple[dict[str, Any], dict[str, int], dict[str, Any]]:
    adapter = area.get("pilot", {}).get("adapter")
    if adapter == "qneat3_pilot":
        return load_qneat3_pilot_frames(area)
    if adapter == "legacy_nong_khai_1":
        return load_legacy_nong_khai_frames(machine, area)
    raise QualityError(f"ไม่รองรับ pilot adapter สำหรับ {area['name']}: {adapter!r}")


def pilot_asset_paths(
    area: dict[str, Any], token: str
) -> tuple[Path, dict[str, Path], dict[str, Path]]:
    area_dir = ensure_within(AREA_ROOT / area["folder_name"], AREA_ROOT, "โฟลเดอร์นำร่อง")
    if not area_dir.is_dir():
        raise SafetyError(f"ไม่พบโฟลเดอร์ผลนำร่อง: {area_dir}")
    names = output_names(area)
    final = {
        key: ensure_within(
            area_dir / layout.area_relpath(key, names[key]), AREA_ROOT, f"ผลนำร่อง {key}"
        )
        for key in ("image", "pdf", "size_image", "qc")
    }
    partial = {key: make_partial(path, token) for key, path in final.items()}
    return area_dir, final, partial


def pilot_qc_payload(
    area: dict[str, Any],
    summary: dict[str, int],
    input_meta: dict[str, Any],
    partial: dict[str, Path],
    size_profile: dict[str, Any] | None = None,
    presentation: dict[str, Any] | None = None,
    consolidation: dict[str, Any] | None = None,
) -> dict[str, Any]:
    validate_expected_summary(area, summary)
    outputs: dict[str, Any] = {}
    for key in ("image", "pdf", "size_image"):
        path = ensure_file(partial[key], f"ผลนำร่อง {key}")
        outputs[key] = {
            "filename": output_names(area)[key],
            "bytes": path.stat().st_size,
            "sha256": sha256_file(path),
        }
    return {
        "schema_version": 1,
        "status": "PASSED",
        "publication_label": "ฉบับนำร่องเพื่อการตรวจสอบ",
        "dataset": "ชุดข้อมูล DMC682",
        "method": "แผน 1.2 — ระยะทางตามโครงข่ายถนนไม่เกิน 10 กิโลเมตร",
        "area": {
            "id": area["id"],
            "areacode": area["areacode"],
            "name": area["name"],
            "province": area["province"],
        },
        "generated_at": local_now(),
        "pipeline_version": VERSION,
        "analysis": {
            "adapter": area["pilot"]["adapter"],
            "summary": summary,
            "analysis_note": area["pilot"].get("analysis"),
        },
        "source_fingerprints": input_meta,
        "outputs": outputs,
        "school_size": size_profile,
        "presentation": presentation,
        "consolidation": consolidation,
    }


def cleanup_partials_under(
    partial: dict[str, Path], parent: Path, token: str
) -> None:
    for path in partial.values():
        try:
            ensure_within(path, parent, "ไฟล์ .partial สำหรับ cleanup")
            if f".partial.{token}" not in path.name:
                raise SafetyError(f"ชื่อ partial ไม่มี token ที่คาดหมาย: {path}")
            if path.is_file():
                path.unlink()
        except (OSError, PipelineError) as exc:
            print(f"[คำเตือน] เก็บ partial ไม่สำเร็จ: {path} ({exc})")


def check_area_layout(area_dir: Path, qc_name: str) -> None:
    """บังคับโครงสร้างโฟลเดอร์ย่อยรายพื้นที่ตาม area_layout (Project Tor.md หัวข้อ 4)

    root ของพื้นที่มีได้เฉพาะ QC JSON ชื่อตรงและไฟล์ .qgz; โฟลเดอร์ย่อยต้องครบทั้งสามชุด
    และไฟล์ในโฟลเดอร์ย่อยต้องมีนามสกุลตรงกับโฟลเดอร์นั้น ไฟล์ .partial ถูกข้ามเพราะเป็นงานที่ยังไม่ผ่าน QC
    """
    for sub in layout.SUBDIRS:
        if not (area_dir / sub).is_dir():
            raise QualityError(f"ไม่พบโฟลเดอร์ย่อย {sub!r} ใน {area_dir.name}")
    for entry in sorted(area_dir.iterdir()):
        if ".partial." in entry.name:
            continue
        if entry.is_dir():
            if entry.name not in layout.SUBDIRS:
                raise QualityError(f"โฟลเดอร์ย่อยไม่รองรับใน {area_dir.name}: {entry.name}")
            allowed = layout.SUBDIR_EXTENSIONS[entry.name]
            for child in sorted(entry.iterdir()):
                if ".partial." in child.name:
                    continue
                if child.is_dir() or child.suffix.lower() not in allowed:
                    raise QualityError(
                        f"ไฟล์อยู่ผิดโฟลเดอร์ย่อยใน {area_dir.name}: "
                        f"{child.relative_to(area_dir).as_posix()}"
                    )
        elif not (entry.name == qc_name or entry.suffix.lower() == ".qgz"):
            raise QualityError(
                f"ไฟล์นอกเหนือ QC/QGZ อยู่ที่ root ของ {area_dir.name}: {entry.name}"
            )


def verify_school_size(area: dict[str, Any], qc: dict[str, Any]) -> None:
    """คำนวณช่วงขนาดใหม่จาก DMC ปัจจุบันแล้วเทียบกับที่ QC บันทึกไว้ทุกค่า"""
    recorded = qc.get("school_size")
    if not isinstance(recorded, dict):
        raise QualityError(f"QC {area['name']} ไม่มี school_size")
    small_rows, _, _ = load_area_records(area)
    expected = school_size_profile(small_rows, expected_total=_qc_summary(qc)["registered"])
    keys = ("bins", "total", "students_min", "students_max", "students_median", "students_total")
    if {key: recorded.get(key) for key in keys} != {key: expected[key] for key in keys}:
        raise QualityError(f"ข้อมูลขนาดโรงเรียนใน QC ของ {area['name']} ไม่ตรง DMC ปัจจุบัน")


def verify_area_assets(area: dict[str, Any], require_size_map: bool = True) -> dict[str, Any]:
    token = "verification"
    _, final, _ = pilot_asset_paths(area, token)
    qc = load_json(final["qc"], f"QC {area['name']}")
    if qc.get("status") != "PASSED":
        raise QualityError(f"QC {area['name']} ไม่ได้สถานะ PASSED")
    if qc.get("area", {}).get("areacode") != area["areacode"]:
        raise QualityError(f"QC ชี้พื้นที่ผิดสำหรับ {area['name']}")
    summary = qc.get("analysis", {}).get("summary")
    if not isinstance(summary, dict):
        raise QualityError(f"QC {area['name']} ไม่มี analysis.summary")
    validate_expected_summary(area, summary)
    if area.get("boundary_override") is not None:
        boundary_path = resolve_catalog_data_path(
            area["boundary_override"], f"boundary_override ของ {area['areacode']}"
        )
        current_boundary = vector_source_fingerprint(boundary_path)
        recorded_boundary = qc.get("source_fingerprints", {}).get("education_boundaries")
        if portable_fingerprint(recorded_boundary) != portable_fingerprint(current_boundary):
            raise QualityError(
                f"QC {area['name']} ไม่ตรง boundary override รุ่นปัจจุบัน; ต้องคำนวณใหม่"
            )
        if qc.get("source_fingerprints", {}).get("boundary_source_kind") != "workspace_override":
            raise QualityError(f"QC {area['name']} ไม่ได้บันทึกการใช้ boundary override")
        recorded_mapping = str(
            qc.get("source_fingerprints", {}).get("boundary_mapping_sha256", "")
        ).upper()
        if recorded_mapping != sha256_file(BOUNDARY_AUDIT_CSV):
            raise QualityError(f"QC {area['name']} ไม่ตรง boundary mapping รุ่นปัจจุบัน")
    output_metadata = qc.get("outputs")
    if not isinstance(output_metadata, dict):
        raise QualityError(f"QC {area['name']} ไม่มี outputs")
    names = output_names(area)
    area_dir = ensure_within(
        AREA_ROOT / area["folder_name"], AREA_ROOT, f"โฟลเดอร์ {area['name']}"
    )
    for key, metadata in output_metadata.items():
        if key not in names or key == "qc" or not isinstance(metadata, dict):
            raise QualityError(f"QC {area['name']} มี output key ไม่รองรับ: {key!r}")
        path = ensure_file(
            area_dir / layout.area_relpath(key, names[key]), f"ผล {key} ของ {area['name']}"
        )
        if metadata.get("filename") != path.name:
            raise QualityError(f"QC filename {key} ไม่ตรงสำหรับ {area['name']}")
        if int(metadata.get("bytes", -1)) != path.stat().st_size:
            raise QualityError(f"QC size {key} ไม่ตรงสำหรับ {area['name']}")
        if str(metadata.get("sha256", "")).upper() != sha256_file(path):
            raise QualityError(f"QC hash {key} ไม่ตรงสำหรับ {area['name']}")
    required_outputs = ("image", "pdf", "size_image") if require_size_map else ("image", "pdf")
    for key in required_outputs:
        if key not in output_metadata:
            raise QualityError(f"QC {area['name']} ไม่มี output {key}")
    try:
        from PIL import Image

        with Image.open(final["image"]) as image:
            if image.width < 2000:
                raise QualityError(f"PNG {area['name']} กว้างต่ำกว่า 2,000 พิกเซล")
            image.verify()
    except ImportError as exc:
        raise PipelineError("ไม่พบ Pillow สำหรับตรวจรับ PNG") from exc
    with final["pdf"].open("rb") as handle:
        if handle.read(5) != b"%PDF-":
            raise QualityError(f"PDF {area['name']} ไม่สมบูรณ์")
    if "size_image" in output_metadata:
        from PIL import Image as SizeImage

        with SizeImage.open(final["size_image"]) as size_image:
            if size_image.width < 2000:
                raise QualityError(f"PNG ขนาดโรงเรียน {area['name']} กว้างต่ำกว่า 2,000 พิกเซล")
            size_image.verify()
        verify_school_size(area, qc)
        if require_size_map:
            presentation = qc.get("presentation") or {}
            if (presentation.get("size_map_max_students") != SIZE_MAP_MAX_STUDENTS
                    or presentation.get("size_map_labels_version") != SIZE_MAP_LABELS_VERSION
                    or presentation.get("size_summary_page_version") != SIZE_SUMMARY_PAGE_VERSION
                    or not isinstance(presentation.get("tambon_lines"), dict)
                    or presentation.get("consolidation_version") != CONSOLIDATION_VERSION):
                raise QualityError(
                    f"QC {area['name']} สร้างด้วยกติกาต่างจากรุ่นปัจจุบัน "
                    f"(แผนที่ขนาดไม่เกิน {SIZE_MAP_MAX_STUDENTS} คน + ป้ายชื่อทุกโรงเรียนที่วาด รุ่น {SIZE_MAP_LABELS_VERSION} "
                    f"+ หน้าคำอธิบายขนาดใน PDF รุ่น {SIZE_SUMMARY_PAGE_VERSION} "
                    f"+ แนวเขตตำบล + ภาพจำลองหลังจับคู่ รุ่น {CONSOLIDATION_VERSION}); "
                    "ต้องรัน refresh-assets"
                )
            verify_consolidation(area, qc)
    check_area_layout(area_dir, names["qc"])
    return qc


def _presentation_input_snapshot(
    machine: dict[str, Any], area: dict[str, Any], qc: dict[str, Any]
) -> dict[str, Any]:
    """Fingerprint only exact render inputs; never load nationwide roads."""
    snapshot: dict[str, Any] = {}
    sources = qc.get("source_fingerprints", {})
    adapter = area.get("pilot", {}).get("adapter")
    if adapter:
        keys = (
            ("small", "roads", "od", "excel")
            if adapter == "legacy_nong_khai_1"
            else ("boundary", "small", "magnets", "roads", "od", "status")
        )
        for key in keys:
            path = _pilot_file(area, key)
            fingerprint = _plain_file_fingerprint(path)
            source_key = "small_legacy" if adapter == "legacy_nong_khai_1" and key == "small" else key
            if sources.get(source_key) != fingerprint:
                raise QualityError(f"ต้นทาง {key} ของ {area['name']} ไม่ตรง QC ที่ตรวจรับ")
            snapshot[str(path)] = fingerprint
        if adapter == "legacy_nong_khai_1":
            boundary_path = exact_external_paths(machine)["education_boundaries"]
            fingerprint = shapefile_fingerprint(boundary_path)
            if portable_fingerprint(sources.get("education_boundaries")) != portable_fingerprint(fingerprint):
                raise QualityError("ขอบเขตต้นทางหนองคายเปลี่ยนจากรุ่นที่ตรวจรับ")
            snapshot[str(boundary_path)] = fingerprint
    else:
        area_dir = ensure_within(AREA_ROOT / area["folder_name"], AREA_ROOT, "พื้นที่สำหรับวาดใหม่")
        names = output_names(area)
        for key in ("boundary", "small", "magnets", "roads", "od", "status"):
            path = ensure_file(area_dir / layout.area_relpath(key, names[key]), f"ผล {key}")
            fingerprint = _plain_file_fingerprint(path)
            if qc.get("outputs", {}).get(key) != fingerprint:
                raise QualityError(f"ผล GIS {key} ของ {area['name']} ไม่ตรง QC ที่ตรวจรับ")
            snapshot[str(path)] = fingerprint
    admin3_path = exact_external_paths(machine)["admin3"]
    snapshot[str(admin3_path)] = shapefile_fingerprint(admin3_path)
    for path in (MASTER_CSV, SMALL_CSV, MAGNET_CSV, BOUNDARY_AUDIT_CSV, AREAS_PATH, MACHINE_PATH):
        snapshot[str(path)] = _plain_file_fingerprint(path)
    return snapshot


def _load_accepted_render_frames(
    machine: dict[str, Any], area: dict[str, Any], qc: dict[str, Any]
) -> dict[str, Any]:
    """Use accepted status/OD geometry without computing or changing a pair."""
    summary = _qc_summary(qc)
    if isinstance(area.get("pilot"), dict):
        frames, actual_summary, input_meta = load_pilot_frames(machine, area)
        if actual_summary != summary or portable_fingerprint(input_meta) != portable_fingerprint(qc.get("source_fingerprints")):
            raise QualityError(f"ผลนำร่องที่อ่านเพื่อวาด {area['name']} ไม่ตรง QC เดิม")
        return frames
    try:
        import geopandas as gpd
        import pandas as pd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas/Pandas ใน QGIS Python") from exc
    names = output_names(area)
    area_dir = ensure_within(AREA_ROOT / area["folder_name"], AREA_ROOT, "พื้นที่สำหรับวาดใหม่")
    frames = {
        key: gpd.read_file(area_dir / layout.area_relpath(key, names[key]))
        for key in ("boundary", "roads", "od", "status", "magnets")
    }
    for key, frame in frames.items():
        _require_pilot_crs(frame, f"ชั้น {key} สำหรับวาดใหม่")
    status = frames["status"]
    if "สถานะ" not in status.columns:
        raise QualityError(f"ชั้นสถานะของ {area['name']} ไม่มี field สถานะ")
    status_keys = {
        "matched": "MATCHED_10KM", "over10km": "OVER_10KM",
        "unreachable": "UNREACHABLE", "excluded": "EXCLUDED_IS_MAGNET",
        "coordExceptions": "COORD_BOUNDARY_EXCEPTION",
    }
    actual = {key: int((status["สถานะ"] == value).sum()) for key, value in status_keys.items()}
    actual.update({"registered": len(status), "magnets": len(frames["magnets"])})
    if any(actual[key] != summary[key] for key in actual):
        raise QualityError(f"จำนวนจากชั้น GIS ของ {area['name']} ไม่ตรงตัวเลข QC")
    if sum(actual[key] for key in status_keys) != len(status) or len(frames["boundary"]) != 1:
        raise QualityError(f"หมวดสถานะหรือขอบเขตของ {area['name']} ไม่ครบ")
    status_ids = _frame_id_set(status, ORIGIN_ID_FIELD, "ชั้นสถานะที่ตรวจรับ")
    od_ids = _frame_id_set(frames["od"], "origin_id", "OD ผ่านเกณฑ์ที่ตรวจรับ")
    matched_ids = _frame_id_set(
        status[status["สถานะ"] == "MATCHED_10KM"], ORIGIN_ID_FIELD, "สถานะผ่านเกณฑ์"
    )
    if od_ids != matched_ids or not od_ids.issubset(status_ids):
        raise QualityError(f"เส้น OD และโรงเรียนผ่านเกณฑ์ของ {area['name']} ไม่ตรงกัน")
    costs = pd.to_numeric(frames["od"]["total_cost"], errors="coerce")
    if costs.isna().any() or (costs < 0).any() or (costs > 10000).any():
        raise QualityError(f"OD ของ {area['name']} มีค่าระยะนอกเกณฑ์")
    return frames


def refresh_assets(
    machine: dict[str, Any], catalog: dict[str, Any], confirmation: str,
    area_codes: list[str] | None = None,
) -> list[dict[str, Any]]:
    """Refresh accepted PNG/PDF together, preserving all analysis and GIS bytes."""
    import copy
    import gc

    if confirmation.strip().upper() != "REFRESH":
        raise SafetyError("ต้องพิมพ์ REFRESH เพื่อยืนยันการแทนที่เฉพาะภาพ/PDF ที่ตรวจรับแล้ว")
    ready_areas = [area for area in catalog["areas"] if area["status"] == "ready"]
    if area_codes:
        requested = {code.strip() for code in area_codes}
        unknown = requested - {area["areacode"] for area in ready_areas}
        if unknown:
            raise SafetyError("refresh-assets รับเฉพาะพื้นที่ ready: " + ", ".join(sorted(unknown)))
        ready_areas = [area for area in ready_areas if area["areacode"] in requested]
    if not ready_areas:
        raise SafetyError("ไม่มีพื้นที่ที่ตรวจรับแล้วสำหรับปรับภาพ")
    require_ready_status(machine, catalog)
    token = uuid.uuid4().hex[:12]
    all_partial: dict[str, Path] = {}
    all_final: dict[str, Path] = {}
    staged: list[dict[str, Any]] = []
    with pipeline_lock("refresh-assets"):
        try:
            processes = qgis_processes()
            if processes["gui"] or processes["workers"]:
                raise SafetyError("พบ QGIS/qgis_process หลัง preflight; ยังไม่เริ่มปรับภาพ")
            # Read and validate all accepted inputs before rendering any area.
            for area in ready_areas:
                _, final, partial = pilot_asset_paths(area, token)
                preflight_outputs(final, overwrite=True)
                # ตรวจผลเดิมแบบไม่บังคับแผนที่ขนาด เพราะ QC รุ่นก่อน 1.4.0 ยังไม่มี
                original = verify_area_assets(area, require_size_map=False)
                validate_expected_result(area, original.get("analysis", {}))
                staged.append({
                    "area": area, "final": final, "partial": partial, "original": original,
                    "inputs": _presentation_input_snapshot(machine, area, original),
                    "qc_hash": sha256_file(final["qc"]),
                })
                # Insertion order commits all QC files after all image/PDF files.
                for key in ("image", "pdf", "size_image"):
                    composite = f"{area['areacode']}:{key}"
                    all_partial[composite], all_final[composite] = partial[key], final[key]
            for item in staged:
                area, partial, original = item["area"], item["partial"], item["original"]
                composite = f"{area['areacode']}:qc"
                all_partial[composite], all_final[composite] = partial["qc"], item["final"]["qc"]
                if available_memory_bytes() < float(machine["limits"]["minimum_available_ram_gb"]) * GIB:
                    raise SafetyError("RAM ว่างต่ำกว่าเกณฑ์ก่อนปรับภาพพื้นที่ถัดไป")
                deadline = time.monotonic() + 90 * 60
                print(f"[ปรับภาพ] {area['name']} — ใช้ผลที่ตรวจรับแล้ว", flush=True)
                frames = _load_accepted_render_frames(machine, area, original)
                require_deadline(deadline, "อ่านแนวเขตตำบล")
                tambon_lines, tambon_meta = load_tambon_lines(machine, frames["boundary"])
                render_map_frames(
                    area, frames["boundary"], frames["roads"], frames["od"],
                    frames["status"], frames["magnets"], partial["image"],
                    tambon_lines=tambon_lines,
                )
                small_rows, _, _ = load_area_records(area)
                size_profile = school_size_profile(
                    small_rows, expected_total=_qc_summary(original)["registered"]
                )
                require_deadline(deadline, "ภาพจำลองหลังจับคู่")
                consolidation = build_consolidation(area, _qc_summary(original), accepted_od_path(area))
                require_deadline(deadline, "วาดแผนที่ขนาดโรงเรียน")
                render_size_map_frames(
                    area, frames["boundary"], frames["roads"], small_rows, partial["size_image"], size_profile,
                    tambon_lines=tambon_lines,
                )
                del frames, small_rows, tambon_lines
                gc.collect()
                require_deadline(deadline, "สร้าง PDF จากผลเดิม")
                render_summary_pdf(
                    area, _qc_summary(original), partial["image"], partial["pdf"],
                    size_image_path=partial["size_image"], consolidation=consolidation, size_profile=size_profile,
                )
                require_deadline(deadline, "ตรวจภาพก่อนเผยแพร่")
                from PIL import Image

                with Image.open(partial["image"]) as image:
                    width, height = image.size
                    if width < 2000:
                        raise QualityError("PNG ใหม่กว้างต่ำกว่า 2,000 พิกเซล")
                    image.verify()
                with Image.open(partial["size_image"]) as size_image:
                    size_width, size_height = size_image.size
                    if size_width < 2000:
                        raise QualityError("PNG แผนที่ขนาดโรงเรียนกว้างต่ำกว่า 2,000 พิกเซล")
                    size_image.verify()
                with partial["pdf"].open("rb") as handle:
                    if handle.read(5) != b"%PDF-":
                        raise QualityError("PDF ใหม่ไม่สมบูรณ์")
                refreshed = copy.deepcopy(original)
                for key in ("image", "pdf", "size_image"):
                    metadata = _plain_file_fingerprint(partial[key])
                    metadata["filename"] = item["final"][key].name
                    refreshed["outputs"][key] = metadata
                refreshed["presentation"] = {
                    "refreshed_at": local_now(), "renderer_version": VERSION,
                    "mode": "accepted-results-render-only", "image_width": width, "image_height": height,
                    "size_image_width": size_width, "size_image_height": size_height,
                    "school_size_map": True, "size_map_max_students": SIZE_MAP_MAX_STUDENTS,
                    "size_map_labels_version": SIZE_MAP_LABELS_VERSION,
                    "size_summary_page_version": SIZE_SUMMARY_PAGE_VERSION,
                    "tambon_lines": tambon_meta, "consolidation_version": CONSOLIDATION_VERSION,
                    "analysis_unchanged": True,
                }
                refreshed["school_size"] = size_profile
                refreshed["consolidation"] = consolidation
                if (refreshed["analysis"] != original["analysis"]
                        or refreshed["source_fingerprints"] != original["source_fingerprints"]
                        or refreshed["generated_at"] != original["generated_at"]
                        or refreshed["pipeline_version"] != original["pipeline_version"]):
                    raise QualityError("การปรับภาพเปลี่ยนข้อมูลวิเคราะห์หรือประวัติผลเดิม")
                write_json(partial["qc"], refreshed)
                item["refreshed"] = refreshed
            for item in staged:
                area, original = item["area"], item["original"]
                if (sha256_file(item["final"]["qc"]) != item["qc_hash"]
                        or _presentation_input_snapshot(machine, area, original) != item["inputs"]):
                    raise QualityError(f"ข้อมูล {area['name']} เปลี่ยนระหว่างปรับภาพ; ยกเลิกทั้งชุด")
                verify_area_assets(area, require_size_map=False)
            processes = qgis_processes()
            if processes["gui"] or processes["workers"]:
                raise SafetyError("พบ QGIS/qgis_process ก่อนแทนที่ภาพ; ยกเลิกทั้งชุด")
            transactional_promote(all_partial, all_final, overwrite=True, token=token)
            print(f"[สำเร็จ] ปรับภาพ/PDF {len(staged)} พื้นที่; ตัวเลขและผล GIS คงเดิม", flush=True)
            return [item["refreshed"] for item in staged]
        except BaseException:
            cleanup_partials_under(all_partial, AREA_ROOT, token)
            raise


def build_pilot_assets(
    machine: dict[str, Any],
    catalog: dict[str, Any],
    confirmation: str,
    overwrite: bool = False,
    area_codes: list[str] | None = None,
) -> list[dict[str, Any]]:
    if confirmation.strip().upper() != "PILOT":
        raise SafetyError("ต้องพิมพ์ PILOT เพื่อยืนยันการสร้างภาพ/PDF นำร่อง")
    require_ready_status(machine, catalog)
    ready_areas = [area for area in catalog["areas"] if area["status"] == "ready"]
    if area_codes:
        requested = {code.strip() for code in area_codes}
        unknown = requested - {area["areacode"] for area in ready_areas}
        if unknown:
            raise SafetyError(
                "pilot-assets รับเฉพาะรหัสพื้นที่สถานะ ready: " + ", ".join(sorted(unknown))
            )
        ready_areas = [area for area in ready_areas if area["areacode"] in requested]
    results: list[dict[str, Any]] = []
    with pipeline_lock("build-pilot-assets"):
        processes = qgis_processes()
        if processes["gui"] or processes["workers"]:
            raise SafetyError("พบ QGIS/qgis_process หลัง preflight; ยังไม่เริ่มเขียนผล")
        for area in ready_areas:
            token = uuid.uuid4().hex[:12]
            _, final, partial = pilot_asset_paths(area, token)
            existing = [path for path in final.values() if path.exists()]
            if len(existing) == len(final) and not overwrite:
                qc = verify_area_assets(area)
                print(f"[คงเดิม] ภาพ/PDF/QC ผ่านการตรวจแล้ว: {area['name']}")
                results.append(qc)
                continue
            if not isinstance(area.get("pilot"), dict):
                raise QualityError(
                    f"ผลสถานะ ready ของ {area['name']} ไม่ครบ; "
                    "ให้รันคำนวณพื้นที่ใหม่แทนการสร้าง pilot-assets"
                )
            preflight_outputs(final, overwrite)
            try:
                frames, summary, input_meta = load_pilot_frames(machine, area)
                tambon_lines, tambon_meta = load_tambon_lines(machine, frames["boundary"])
                render_map_frames(
                    area,
                    frames["boundary"],
                    frames["roads"],
                    frames["od"],
                    frames["status"],
                    frames["magnets"],
                    partial["image"],
                    tambon_lines=tambon_lines,
                )
                small_rows, _, _ = load_area_records(area)
                size_profile = school_size_profile(small_rows, expected_total=summary["registered"])
                consolidation = build_consolidation(area, summary, _pilot_file(area, "od"))
                render_size_map_frames(
                    area, frames["boundary"], frames["roads"], small_rows, partial["size_image"], size_profile,
                    tambon_lines=tambon_lines,
                )
                render_summary_pdf(
                    area, summary, partial["image"], partial["pdf"], size_image_path=partial["size_image"],
                    consolidation=consolidation, size_profile=size_profile,
                )
                qc = pilot_qc_payload(
                    area, summary, input_meta, partial, size_profile=size_profile,
                    presentation=fresh_presentation(tambon_meta), consolidation=consolidation,
                )
                write_json(partial["qc"], qc)
                # QC is inserted last and therefore acts as the commit marker.
                transactional_promote(partial, final, overwrite, token)
                results.append(verify_area_assets(area))
                print(f"[สำเร็จ] สร้างภาพ/PDF/QC: {area['name']}")
            except BaseException:
                cleanup_partials_under(partial, AREA_ROOT, token)
                raise
    return results


def _qc_summary(qc: dict[str, Any]) -> dict[str, int]:
    summary = qc.get("analysis", {}).get("summary")
    if not isinstance(summary, dict):
        raise QualityError("QC ไม่มี summary")
    return {key: int(value) for key, value in summary.items()}


def build_portal_catalog_payload(
    catalog: dict[str, Any], qcs: dict[str, dict[str, Any]]
) -> dict[str, Any]:
    status_counts: dict[str, int] = {
        "ready": 0,
        "pending": 0,
        "blocked": 0,
        "empty": 0,
    }
    for area in catalog["areas"]:
        status_counts[area["status"]] += 1
    province_by_code: dict[str, str] = {}
    for area in catalog["areas"]:
        province_code = "TH" + area["areacode"][:2]
        existing = province_by_code.setdefault(province_code, area["province"])
        if existing != area["province"]:
            raise QualityError(f"รหัสจังหวัด {province_code} map ไปหลายชื่อ")
    provinces = [
        {"code": code, "name": name, "svgId": f"province-{code}"}
        for code, name in sorted(province_by_code.items())
    ]
    entries: list[dict[str, Any]] = []
    ready_update_dates = [
        str(qc.get("generated_at", ""))[:10]
        for qc in qcs.values()
        if str(qc.get("generated_at", ""))[:10]
    ]
    catalog_updated = max(ready_update_dates) if ready_update_dates else None
    null_summary = {
        "registered": None,
        "origins": None,
        "magnets": None,
        "matched": None,
        "over10km": None,
        "unreachable": None,
        "excluded": None,
        "coordExceptions": None,
    }
    for area in catalog["areas"]:
        entry = {
            "id": area["id"],
            "areacode": area["areacode"],
            "province": area["province"],
            "name": area["name"],
            "status": area["status"],
            "summary": dict(null_summary),
            "image": None,
            "pdf": None,
            "sizeImage": None,
            "sizeProfile": None,
            "sizeAnalysis": None,
            "consolidation": None,
            "analysis": area.get("status_note"),
            "updated": None,
        }
        if area["status"] == "ready":
            qc = qcs.get(area["areacode"])
            if qc is None:
                raise QualityError(f"ไม่มี QC สำหรับพื้นที่พร้อมดู {area['name']}")
            summary = _qc_summary(qc)
            validate_expected_summary(area, summary)
            entry["summary"] = {key: summary[key] for key in null_summary}
            entry["image"] = f"files/{area['id']}-map.png"
            entry["pdf"] = f"files/{area['id']}-summary.pdf"
            profile = qc.get("school_size")
            if not isinstance(profile, dict):
                raise QualityError(f"QC {area['name']} ไม่มี school_size สำหรับพอร์ทัล")
            entry["sizeImage"] = f"files/{area['id']}-sizes.png"
            entry["sizeProfile"] = portal_size_profile(profile)
            entry["sizeAnalysis"] = public_size_analysis(profile)
            scenario = qc.get("consolidation")
            if not isinstance(scenario, dict):
                raise QualityError(f"QC {area['name']} ไม่มี consolidation สำหรับพอร์ทัล; ต้องรัน refresh-assets")
            entry["consolidation"] = portal_consolidation(scenario)
            entry["analysis"] = public_analysis(area, summary)
            entry["updated"] = str(qc.get("generated_at", ""))[:10] or None
        entries.append(entry)
    payload = {
        "meta": {
            "schemaVersion": 1,
            "title": "พอร์ทัลผลการจับคู่โรงเรียนขนาดเล็ก",
            "releaseLabel": "ฉบับนำร่องเพื่อการตรวจสอบ",
            "dataset": "DMC682",
            "method": "แผน 1.2 — ระยะทางตามโครงข่ายถนนไม่เกิน 10 กิโลเมตร",
            "thresholdMeters": 10000,
            # Derive this from accepted QC so verification remains stable on
            # future dates.  Using today's date made an unchanged valid portal
            # fail acceptance after midnight.
            "updated": catalog_updated,
            "provinceCount": len(provinces),
            "areaCount": len(entries),
            "statusCounts": status_counts,
        },
        "provinces": provinces,
        "areas": entries,
    }
    if len(provinces) != 77 or len(entries) != 183:
        raise QualityError("catalog พอร์ทัลต้องมี 77 จังหวัดและ 183 พื้นที่")
    expected_status_counts = catalog.get("expected_status_counts")
    if status_counts != expected_status_counts:
        raise QualityError(
            f"จำนวนสถานะพอร์ทัลเปลี่ยน: ได้ {status_counts}, "
            f"คาด {expected_status_counts}"
        )
    return payload


def portal_catalog_text(payload: dict[str, Any]) -> str:
    body = json.dumps(payload, ensure_ascii=False, indent=2)
    return (
        '"use strict";\n\n'
        "// บัญชีระดับจังหวัด/พื้นที่และตัวเลขรวมที่อนุญาตให้เผยแพร่เท่านั้น\n"
        "// ไม่มีรายชื่อโรงเรียนขนาดเล็กต้นทาง รหัสโรงเรียน หรือพิกัดในไฟล์นี้; ชื่อและจำนวนนักเรียน (DMC682) ของโรงเรียนคุณภาพชุมชนแสดงเฉพาะในภาพจำลองหลังจับคู่\n"
        f"window.SMS_CATALOG = {body};\n"
    )


def parse_portal_catalog(path: Path) -> dict[str, Any]:
    text_value = ensure_file(path, "catalog-data.js").read_text(encoding="utf-8-sig")
    marker = "window.SMS_CATALOG = "
    start = text_value.find(marker)
    if start < 0:
        raise QualityError("catalog-data.js ไม่มี window.SMS_CATALOG")
    json_text = text_value[start + len(marker) :].strip()
    if not json_text.endswith(";"):
        raise QualityError("catalog-data.js ไม่มีเครื่องหมายปิด")
    try:
        value = json.loads(json_text[:-1])
    except json.JSONDecodeError as exc:
        raise QualityError(f"catalog-data.js มี JSON ไม่สมบูรณ์: {exc}") from exc
    if not isinstance(value, dict):
        raise QualityError("catalog-data.js ต้องเป็น object")
    return value


def publish_portal(
    catalog: dict[str, Any], confirmation: str, overwrite: bool = False
) -> dict[str, Any]:
    if confirmation.strip().upper() != "PUBLISH":
        raise SafetyError("ต้องพิมพ์ PUBLISH เพื่อยืนยันการเผยแพร่พอร์ทัล")
    validate_catalog(catalog, verify_hashes=True)
    ensure_within(PORTAL_ROOT, WORKSPACE, "โฟลเดอร์พอร์ทัล")
    files_dir = ensure_within(PORTAL_ROOT / "files", PORTAL_ROOT, "โฟลเดอร์ไฟล์พอร์ทัล")
    files_dir.mkdir(parents=True, exist_ok=True)
    qcs: dict[str, dict[str, Any]] = {}
    area_assets: dict[str, dict[str, Path]] = {}
    for area in catalog["areas"]:
        if area["status"] != "ready":
            continue
        qc = verify_area_assets(area)
        qcs[area["areacode"]] = qc
        _, final, _ = pilot_asset_paths(area, "portal-source")
        area_assets[area["areacode"]] = final

    payload = build_portal_catalog_payload(catalog, qcs)
    token = uuid.uuid4().hex[:12]
    final: dict[str, Path] = {}
    for area in catalog["areas"]:
        if area["status"] != "ready":
            continue
        final[f"{area['id']}-image"] = ensure_within(
            files_dir / f"{area['id']}-map.png", PORTAL_ROOT, "PNG พอร์ทัล"
        )
        final[f"{area['id']}-pdf"] = ensure_within(
            files_dir / f"{area['id']}-summary.pdf", PORTAL_ROOT, "PDF พอร์ทัล"
        )
        final[f"{area['id']}-sizes"] = ensure_within(
            files_dir / f"{area['id']}-sizes.png", PORTAL_ROOT, "PNG ขนาดโรงเรียนพอร์ทัล"
        )
    final["catalog"] = ensure_within(PORTAL_ROOT / "catalog-data.js", PORTAL_ROOT, "catalog พอร์ทัล")
    partial = {key: make_partial(path, token) for key, path in final.items()}
    preflight_outputs(final, overwrite)
    with pipeline_lock("publish-portal"):
        try:
            for area in catalog["areas"]:
                if area["status"] != "ready":
                    continue
                sources = area_assets[area["areacode"]]
                shutil.copy2(sources["image"], partial[f"{area['id']}-image"])
                shutil.copy2(sources["pdf"], partial[f"{area['id']}-pdf"])
                shutil.copy2(sources["size_image"], partial[f"{area['id']}-sizes"])
            partial["catalog"].write_text(portal_catalog_text(payload), encoding="utf-8", newline="\n")
            # catalog is promoted last and is the visible commit marker for assets.
            transactional_promote(partial, final, overwrite, token)
        except BaseException:
            cleanup_partials_under(partial, PORTAL_ROOT, token)
            raise
    print(f"เผยแพร่พอร์ทัลออฟไลน์แล้ว: {PORTAL_ROOT / 'index.html'}")
    return payload


def verify_delivery(catalog: dict[str, Any]) -> dict[str, Any]:
    catalog_check = validate_catalog(catalog, verify_hashes=True)
    qcs: dict[str, dict[str, Any]] = {}
    ready_areas = [area for area in catalog["areas"] if area["status"] == "ready"]
    for area in ready_areas:
        qcs[area["areacode"]] = verify_area_assets(area)

    for filename in ("index.html", "portal.css", "portal.js", "catalog-data.js"):
        ensure_file(PORTAL_ROOT / filename, f"ไฟล์พอร์ทัล {filename}")
    portal_payload = parse_portal_catalog(PORTAL_ROOT / "catalog-data.js")
    expected_payload = build_portal_catalog_payload(catalog, qcs)
    if portal_payload != expected_payload:
        raise QualityError("catalog-data.js ไม่ตรง QC/areas.json ปัจจุบัน; ต้องสร้างพอร์ทัลใหม่")
    catalog_text = (PORTAL_ROOT / "catalog-data.js").read_text(encoding="utf-8-sig")
    if re.search(r"\d{9,}", catalog_text):
        raise QualityError("catalog-data.js มีตัวเลขยาวคล้ายรหัสโรงเรียน")
    for area in ready_areas:
        entry = next(item for item in portal_payload["areas"] if item["areacode"] == area["areacode"])
        small, magnets, overlap = load_area_records(area)
        names = {str(row["name"]) for row in (entry.get("consolidation") or {}).get("schools", [])}
        magnet_names = {row["school_name"] for row in magnets}
        origin_names = {row["school_name"] for row in small if row["school_id"] not in overlap}
        if not names <= magnet_names or len(names) != len(magnets) or names & origin_names:
            raise QualityError(f"catalog มีชื่อโรงเรียนนอกกลุ่มโรงเรียนคุณภาพชุมชนของ {area['name']}")
    for area in ready_areas:
        _, area_final, _ = pilot_asset_paths(area, "verify-portal")
        portal_image = ensure_file(
            PORTAL_ROOT / "files" / f"{area['id']}-map.png", "PNG ในพอร์ทัล"
        )
        portal_pdf = ensure_file(
            PORTAL_ROOT / "files" / f"{area['id']}-summary.pdf", "PDF ในพอร์ทัล"
        )
        if sha256_file(portal_image) != sha256_file(area_final["image"]):
            raise QualityError(f"PNG พอร์ทัลไม่ตรงต้นฉบับพื้นที่: {area['name']}")
        if sha256_file(portal_pdf) != sha256_file(area_final["pdf"]):
            raise QualityError(f"PDF พอร์ทัลไม่ตรงต้นฉบับพื้นที่: {area['name']}")
        portal_sizes = ensure_file(
            PORTAL_ROOT / "files" / f"{area['id']}-sizes.png", "PNG ขนาดโรงเรียนในพอร์ทัล"
        )
        if sha256_file(portal_sizes) != sha256_file(area_final["size_image"]):
            raise QualityError(f"PNG ขนาดโรงเรียนในพอร์ทัลไม่ตรงต้นฉบับพื้นที่: {area['name']}")

    portal_sources = "\n".join(
        ensure_file(PORTAL_ROOT / name, name).read_text(encoding="utf-8-sig")
        for name in ("index.html", "portal.css", "portal.js", "catalog-data.js")
    ).lower()
    for forbidden in ("http://", "https://", "fetch(", "xmlhttprequest", "websocket"):
        if forbidden in portal_sources:
            raise QualityError(f"พอร์ทัลมีการอ้างเครือข่ายที่ห้ามใช้: {forbidden}")
    report = {
        "status": "PASSED",
        "checked_at": local_now(),
        "areas": catalog_check["area_count"],
        "provinces": catalog_check["province_count"],
        "ready_assets": len(ready_areas),
        "status_counts": catalog_check["status_counts"],
        "offline_network_tokens": 0,
    }
    print("\nผลตรวจรับระบบส่งมอบ: PASSED")
    print(f"พื้นที่ทั้งหมด {report['areas']} | จังหวัด {report['provinces']} | พร้อมดู {report['ready_assets']}")
    return report


def read_od_pair_costs(path: Path) -> dict[str, tuple[str, float]]:
    ogr, _ = ogr_modules()
    dataset = ogr.Open(str(ensure_file(path, "OD สำหรับเปรียบเทียบ")), 0)
    if dataset is None or dataset.GetLayerCount() < 1:
        raise QualityError(f"เปิด OD สำหรับเปรียบเทียบไม่ได้: {path}")
    layer = dataset.GetLayer(0)
    fields = {
        layer.GetLayerDefn().GetFieldDefn(index).GetName()
        for index in range(layer.GetLayerDefn().GetFieldCount())
    }
    required = {"origin_id", "destination_id", "total_cost"}
    if not required.issubset(fields):
        raise QualityError(f"OD สำหรับเปรียบเทียบไม่มี field {sorted(required - fields)}")
    pairs: dict[str, tuple[str, float]] = {}
    for feature in layer:
        origin = normalize_school_id(feature.GetField("origin_id"), "origin_id")
        destination = normalize_school_id(
            feature.GetField("destination_id"), "destination_id"
        )
        try:
            cost = float(feature.GetField("total_cost"))
        except (TypeError, ValueError) as exc:
            raise QualityError(f"OD มี total_cost ไม่ใช่ตัวเลขสำหรับ {origin}") from exc
        if not math.isfinite(cost) or cost < 0:
            raise QualityError(f"OD มี total_cost ไม่ถูกต้องสำหรับ {origin}")
        if origin in pairs:
            raise QualityError(f"OD ที่เปรียบเทียบมี origin ซ้ำ: {origin}")
        pairs[origin] = (destination, cost)
    dataset = None
    return pairs


def compare_od_pairs(
    accepted: dict[str, tuple[str, float]],
    candidate: dict[str, tuple[str, float]],
    tolerance_metres: float = 0.001,
) -> dict[str, Any]:
    if set(accepted) != set(candidate):
        missing = sorted(set(accepted) - set(candidate))
        extra = sorted(set(candidate) - set(accepted))
        raise QualityError(
            f"ชุด origin ของผลคำนวณซ้ำไม่ตรง: missing={missing[:5]}, extra={extra[:5]}"
        )
    maximum_difference = 0.0
    pair_tokens: list[str] = []
    millimetre_tokens: list[str] = []
    for origin in sorted(accepted):
        accepted_destination, accepted_cost = accepted[origin]
        candidate_destination, candidate_cost = candidate[origin]
        if accepted_destination != candidate_destination:
            raise QualityError(
                f"ผลคำนวณซ้ำเลือกปลายทางต่างกันสำหรับ {origin}: "
                f"{candidate_destination} แทน {accepted_destination}"
            )
        difference = abs(candidate_cost - accepted_cost)
        maximum_difference = max(maximum_difference, difference)
        if difference > tolerance_metres:
            raise QualityError(
                f"ระยะผลคำนวณซ้ำต่างเกิน {tolerance_metres} เมตรสำหรับ {origin}: "
                f"ต่าง {difference:.6f} เมตร"
            )
        pair_tokens.append(f"{origin}:{candidate_destination}")
        millimetre_tokens.append(
            f"{origin}:{candidate_destination}:{candidate_cost:.3f}"
        )
    return {
        "pair_count": len(candidate),
        "maximum_distance_difference_metres": maximum_difference,
        "pair_digest": hashlib.sha256(
            "\n".join(pair_tokens).encode("utf-8")
        ).hexdigest().upper(),
        "distance_digest_millimetre": hashlib.sha256(
            "\n".join(millimetre_tokens).encode("utf-8")
        ).hexdigest().upper(),
    }


def recompute_check(
    machine: dict[str, Any],
    catalog: dict[str, Any],
    area: dict[str, Any],
    confirmation: str,
    runs: int = 2,
) -> dict[str, Any]:
    code = area["areacode"]
    if confirmation.strip() != code:
        raise SafetyError("รหัสยืนยันไม่ตรง; ยังไม่เริ่ม recompute-check")
    if area.get("pilot", {}).get("adapter") != "qneat3_pilot":
        raise SafetyError("recompute-check ใช้ได้เฉพาะพื้นที่นำร่อง QNEAT3")
    if runs not in {1, 2}:
        raise SafetyError("recompute-check อนุญาต 1 หรือ 2 รอบเท่านั้น")
    preflight = require_ready_status(machine, catalog)
    accepted = read_od_pair_costs(_pilot_file(area, "od"))
    expected = area["pilot"]["expected"]
    run_reports: list[dict[str, Any]] = []
    candidate_runs: list[dict[str, tuple[str, float]]] = []
    with pipeline_lock("recompute-check", code):
        processes = qgis_processes()
        if processes["gui"] or processes["workers"]:
            raise SafetyError("พบ QGIS/qgis_process หลัง preflight; ยังไม่เริ่ม recompute-check")
        for run_number in range(1, runs + 1):
            token = uuid.uuid4().hex[:12]
            run_root = local_state_root() / "runs" / code
            run_root.mkdir(parents=True, exist_ok=True)
            run_dir = run_root / (
                f"recompute-{run_number}-{dt.datetime.now():%Y%m%d-%H%M%S}-{token}"
            )
            run_dir.mkdir(parents=False, exist_ok=False)
            temporary = {
                "boundary": run_dir / "boundary.gpkg",
                "small": run_dir / "small.gpkg",
                "magnets": run_dir / "magnets.gpkg",
                "roads": run_dir / "roads.gpkg",
                "od": run_dir / "od_matched.gpkg",
                "status": run_dir / "status.gpkg",
            }
            deadline = time.monotonic() + float(machine["limits"]["area_timeout_seconds"])
            completed = False
            try:
                print(f"\nrecompute-check รอบ {run_number}/{runs}: {area['name']}")
                small_all, magnets, overlap = load_area_records(area)
                analysis_origins = [
                    row for row in small_all if row["school_id"] not in overlap
                ]
                create_school_points(
                    temporary["small"], "small_school_origins", analysis_origins, magnet=False
                )
                create_school_points(
                    temporary["magnets"], "magnet_schools", magnets, magnet=True
                )
                preparation = prepare_boundary_and_roads(
                    machine, area, temporary, run_dir, deadline
                )
                require_deadline(deadline, "QNEAT3 recompute-check")
                full_od = run_qneat3(machine, temporary, run_dir, deadline)
                require_deadline(deadline, "วิเคราะห์ recompute-check")
                result = analyse_od_and_write_outputs(
                    machine,
                    full_od,
                    temporary["boundary"],
                    small_all,
                    analysis_origins,
                    magnets,
                    overlap,
                    temporary["od"],
                    temporary["status"],
                )
                validate_expected_summary(area, result["summary"])
                if result["full_od_rows"] != result["expected_full_od_rows"]:
                    raise QualityError(
                        "พื้นที่นำร่องต้องได้ Full OD ครบ: "
                        f"ได้ {result['full_od_rows']}, คาด {result['expected_full_od_rows']}"
                    )
                candidate_pairs = read_od_pair_costs(temporary["od"])
                comparison = compare_od_pairs(
                    accepted,
                    candidate_pairs,
                    tolerance_metres=float(machine["analysis"]["tie_tolerance_metres"]),
                )
                candidate_runs.append(candidate_pairs)
                report = {
                    "run": run_number,
                    "summary": result["summary"],
                    "full_od_rows": result["full_od_rows"],
                    "expected_full_od_rows": result["expected_full_od_rows"],
                    "road_feature_count": preparation["roads"]["feature_count"],
                    "oneway_counts": preparation["oneway_counts"],
                    **comparison,
                }
                run_reports.append(report)
                completed = True
                print(
                    f"[ผ่าน] รอบ {run_number}: Full OD {report['full_od_rows']}, "
                    f"คู่ผ่าน {report['pair_count']}, ต่างสูงสุด "
                    f"{report['maximum_distance_difference_metres']:.9f} เมตร"
                )
                print(f"  pair_digest: {report['pair_digest']}")
            finally:
                if completed:
                    try:
                        cleanup_successful_run_dir(run_dir)
                    except (OSError, PipelineError) as exc:
                        print(f"[คำเตือน] ลบผล recompute-check หลังผ่านไม่ได้: {exc}")
                else:
                    print(f"เก็บ log recompute-check ที่ไม่สำเร็จไว้: {run_dir}")

    if len(run_reports) == 2:
        first, second = run_reports
        if first["pair_digest"] != second["pair_digest"]:
            raise QualityError("recompute-check สองรอบได้ pair digest ไม่เหมือนกัน")
        between_runs = compare_od_pairs(
            candidate_runs[0],
            candidate_runs[1],
            tolerance_metres=float(machine["analysis"]["tie_tolerance_metres"]),
        )
        if between_runs["distance_digest_millimetre"] != first["distance_digest_millimetre"]:
            # This is informational only when both direct comparisons are
            # inside the explicit 1 mm tolerance.  Exact decimal hashes are
            # overly sensitive to harmless floating-point boundary rounding.
            print("[คำเตือน] distance digest ระดับมิลลิเมตรต่าง แต่ค่าดิบอยู่ใน tolerance")
        if first["summary"] != second["summary"]:
            raise QualityError("recompute-check สองรอบได้ summary ไม่เหมือนกัน")
    final_report = {
        "status": "PASSED",
        "area": {"areacode": code, "name": area["name"]},
        "runs": run_reports,
        "preflight_checked_at": preflight["checked_at"],
        "workspace_outputs_modified": False,
    }
    print(f"recompute-check ผ่าน {runs} รอบ โดยไม่แก้ผลใน workspace")
    return final_report


def area_by_code(catalog: dict[str, Any], code: str) -> dict[str, Any]:
    normalized = code.strip()
    for area in catalog["areas"]:
        if area["areacode"] == normalized or area["id"] == normalized:
            return area
    raise PipelineError(f"ไม่พบรหัสพื้นที่: {code}")


def search_areas(catalog: dict[str, Any], query: str) -> list[dict[str, Any]]:
    needle = query.strip().casefold()
    if not needle:
        return []
    return [
        area
        for area in catalog["areas"]
        if needle in area["name"].casefold()
        or needle in area["province"].casefold()
        or needle in area["areacode"]
    ]


def interactive_area(catalog: dict[str, Any]) -> dict[str, Any] | None:
    query = input("ค้นด้วยชื่อจังหวัด/สพป. หรือรหัส (Enter เพื่อยกเลิก): ").strip()
    if not query:
        return None
    matches = search_areas(catalog, query)
    if not matches:
        print("ไม่พบพื้นที่")
        return None
    for area in matches[:30]:
        print(f"  {area['areacode']}  {area['name']}  [{area['status']}]")
    if len(matches) > 30:
        print(f"  ... อีก {len(matches) - 30} พื้นที่ โปรดค้นให้แคบลง")
    code = input("พิมพ์รหัสพื้นที่ 8 หลักที่ต้องการ: ").strip()
    return area_by_code(catalog, code)


def menu(machine: dict[str, Any], catalog: dict[str, Any]) -> int:
    while True:
        print("\n" + "=" * 72)
        print("ระบบคำนวณและพอร์ทัลผลการจับคู่สำหรับ สพฐ.")
        print("1  ตรวจสถานะข้อมูลและเครื่อง")
        print("2  เลือกพื้นที่เพื่อคำนวณ")
        print("3  สร้าง/อัปเดตพอร์ทัล")
        print("4  ตรวจรับผล")
        print("0  ออกจากระบบ")
        choice = input("เลือกเมนู: ").strip()
        try:
            if choice == "1":
                print_status(status_report(machine, catalog, deep_fingerprints=True))
            elif choice == "2":
                area = interactive_area(catalog)
                if area is None:
                    continue
                print(f"เลือก {area['name']} ({area['areacode']})")
                confirmation = input("พิมพ์รหัสพื้นที่ซ้ำเพื่อยืนยัน: ").strip()
                _, final, _ = output_paths(area, "menu-check")
                overwrite = any(path.exists() for path in final.values())
                if overwrite:
                    overwrite = input("พบผลเดิม พิมพ์ OVERWRITE เพื่อสร้างรุ่นใหม่: ").strip().upper() == "OVERWRITE"
                    if not overwrite:
                        print("ยกเลิกโดยไม่แก้ผลเดิม")
                        continue
                run_area(machine, catalog, area, confirmation, overwrite=overwrite)
            elif choice == "3":
                pilot_confirmation = input("พิมพ์ PILOT เพื่อสร้าง/ตรวจภาพนำร่อง: ")
                build_pilot_assets(machine, catalog, pilot_confirmation, overwrite=False)
                publish_confirmation = input("พิมพ์ PUBLISH เพื่อเผยแพร่พอร์ทัล: ")
                portal_exists = (PORTAL_ROOT / "catalog-data.js").exists()
                publish_portal(catalog, publish_confirmation, overwrite=portal_exists)
                verify_delivery(catalog)
            elif choice == "4":
                verify_delivery(catalog)
            elif choice == "0":
                return 0
            else:
                print("กรุณาเลือก 0–4")
        except PipelineError as exc:
            print(f"[ไม่สำเร็จ] {exc}")

# ---------------------------------------------------------------------------
# relayout — ย้ายไฟล์ของพื้นที่เข้าโฟลเดอร์ย่อยตาม area_layout โดยไม่เปลี่ยนชื่อและไม่คำนวณใหม่
# ---------------------------------------------------------------------------

RELAYOUT_STALE_SUFFIXES = (".tmp", ".bak", ".lock", ".gpkg-wal", ".gpkg-shm", ".gpkg-journal")


def plan_relayout(area: dict[str, Any], area_dir: Path) -> dict[str, Any]:
    """วางแผนการย้ายแบบอ่านอย่างเดียว: คืนรายการย้าย ไฟล์ที่อยู่ถูกที่แล้ว และปัญหาที่ต้องหยุด"""
    qc_name = f"ผลตรวจ_{area['areacode']}.json"
    moves: list[tuple[Path, Path]] = []
    in_place: list[str] = []
    keep_root: list[str] = []
    projects: list[Path] = []
    problems: list[str] = []
    for entry in sorted(area_dir.iterdir()):
        name = entry.name
        if ".partial." in name or ".backup." in name or name.lower().endswith(RELAYOUT_STALE_SUFFIXES):
            problems.append(f"ไฟล์ค้างจากงานอื่น: {name}")
            continue
        if entry.is_dir():
            if name not in layout.SUBDIRS:
                problems.append(f"โฟลเดอร์ย่อยไม่รองรับ: {name}")
                continue
            allowed = layout.SUBDIR_EXTENSIONS[name]
            for child in sorted(entry.iterdir()):
                relative = child.relative_to(area_dir).as_posix()
                if child.is_dir():
                    problems.append(f"โฟลเดอร์ซ้อนไม่รองรับ: {relative}")
                elif ".partial." in child.name:
                    problems.append(f"ไฟล์ค้างจากงานอื่น: {relative}")
                elif child.suffix.lower() not in allowed:
                    problems.append(f"ไฟล์อยู่ผิดโฟลเดอร์ย่อย: {relative}")
                else:
                    in_place.append(relative)
            continue
        if name == qc_name:
            keep_root.append(name)
            continue
        if entry.suffix.lower() == ".qgz":
            keep_root.append(name)
            projects.append(entry)
            continue
        subdir = layout.subdir_for_extension(entry.suffix)
        if subdir is None:
            problems.append(f"ไม่มีกฎ layout สำหรับ: {name}")
            continue
        moves.append((entry, area_dir / subdir / name))
    return {
        "qc_name": qc_name,
        "moves": moves,
        "in_place": in_place,
        "keep_root": keep_root,
        "projects": projects,
        "problems": problems,
    }


def _qgs_entry_name(archive: zipfile.ZipFile, label: str) -> str:
    names = [name for name in archive.namelist() if name.lower().endswith(".qgs")]
    if len(names) != 1:
        raise QualityError(f"QGZ ต้องมีไฟล์ .qgs หนึ่งไฟล์: {label}")
    return names[0]


def _relative_gpkg_pattern(name: str) -> re.Pattern[str]:
    # lookbehind กัน ../x.gpkg และพาธที่ถูกแทนแล้ว (./01 ข้อมูล GIS/x.gpkg) ไม่ให้ถูกแทนซ้ำ
    return re.compile(r"(?<![\w./])\./" + re.escape(name))


def validate_qgz_layout(qgz: Path, area_dir: Path, gpkg_names: list[str]) -> dict[str, int]:
    """ตรวจว่าทุก datasource แบบ ./ ในโครงการชี้ไฟล์ที่มีอยู่จริง และไม่มีพาธเดิมเหลือ"""
    with zipfile.ZipFile(qgz) as archive:
        data = archive.read(_qgs_entry_name(archive, qgz.name))
    text = data.decode("utf-8")
    for name in gpkg_names:
        if _relative_gpkg_pattern(name).search(text):
            raise QualityError(f"QGZ ยังอ้างพาธเดิม ./{name}: {qgz.name}")
    try:
        root = ET.fromstring(data)
    except ET.ParseError as exc:
        raise QualityError(f"อ่าน XML ของ QGZ ไม่ผ่าน: {qgz.name} ({exc})") from exc
    absolute = root.find("./properties/Paths/Absolute")
    if absolute is None or (absolute.text or "").strip().lower() != "false":
        raise QualityError(f"QGZ ไม่ได้ตั้งค่า relative paths: {qgz.name}")
    sources = [element.text or "" for element in root.iter("datasource")]
    sources += [element.get("source", "") for element in root.iter("layer-tree-layer")]
    local = other = 0
    for source in sources:
        if source.startswith("./"):
            relative = source.split("|", 1)[0][2:]
            target = ensure_within(area_dir / relative, AREA_ROOT, "datasource ใน QGZ")
            if not target.is_file():
                raise QualityError(f"QGZ อ้างไฟล์ที่ไม่มี: {relative} ({qgz.name})")
            local += 1
        elif source:
            other += 1
    return {"local_sources": local, "other_sources": other}


def rewrite_qgz_for_layout(
    qgz: Path, area_dir: Path, gpkg_names: list[str], token: str
) -> dict[str, Any]:
    """แทน ./<gpkg> ด้วย ./01 ข้อมูล GIS/<gpkg> ในไฟล์ .qgs ภายใน .qgz โดยคัดลอก entry อื่นไบต์ต่อไบต์"""
    with zipfile.ZipFile(qgz) as archive:
        infos = archive.infolist()
        qgs_name = _qgs_entry_name(archive, qgz.name)
        raw = {info.filename: archive.read(info) for info in infos}
    # decode จาก bytes ตรง ๆ เพื่อคง CRLF/LF เดิมของแต่ละโครงการ
    text = raw[qgs_name].decode("utf-8")
    area_posix = area_dir.resolve().as_posix()
    replacements: list[dict[str, Any]] = []
    for name in gpkg_names:
        text, relative_count = _relative_gpkg_pattern(name).subn(
            f"./{layout.GIS_DIR}/{name}", text
        )
        absolute_pattern = re.compile(re.escape(f"{area_posix}/{name}"), re.IGNORECASE)
        text, absolute_count = absolute_pattern.subn(
            f"{area_posix}/{layout.GIS_DIR}/{name}", text
        )
        if relative_count or absolute_count:
            replacements.append(
                {"name": name, "relative": relative_count, "absolute": absolute_count}
            )
    if not replacements:
        return {"changed": False, "replacements": []}
    partial = make_partial(qgz, token)
    with zipfile.ZipFile(partial, "w") as archive:
        for info in infos:
            payload = text.encode("utf-8") if info.filename == qgs_name else raw[info.filename]
            entry = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            entry.compress_type = info.compress_type
            entry.external_attr = info.external_attr
            archive.writestr(entry, payload)
    try:
        validate_qgz_layout(partial, area_dir, gpkg_names)
    except PipelineError:
        partial.unlink(missing_ok=True)
        raise
    os.replace(partial, qgz)
    return {"changed": True, "replacements": replacements}


def qgis_check_project(machine: dict[str, Any], qgz: Path, area_dir: Path) -> list[dict[str, Any]]:
    """โหลดโครงการด้วย PyQGIS แบบ headless แล้วยืนยันว่าทุกชั้นที่อยู่ในโฟลเดอร์พื้นที่ valid"""
    global _QGIS_APP
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    try:
        from qgis.core import QgsApplication, QgsProject
    except ImportError as exc:
        raise PipelineError("ไม่พบ PyQGIS; ต้องรันด้วย python-qgis-ltr.bat") from exc
    if _QGIS_APP is None:
        QgsApplication.setPrefixPath(str(machine["qgis"]["prefix_path"]), True)
        _QGIS_APP = QgsApplication([], False)
        _QGIS_APP.initQgis()
    project = QgsProject.instance()
    project.clear()
    if not project.read(str(qgz)):
        raise QualityError(f"QGIS อ่านโครงการไม่สำเร็จ: {qgz.name}")
    area_posix = area_dir.resolve().as_posix().lower()
    report: list[dict[str, Any]] = []
    for layer in project.mapLayers().values():
        source = str(layer.source())
        head = source.split("|", 1)[0].replace("\\", "/").lower()
        local = head.startswith(area_posix)
        report.append({"layer": layer.name(), "valid": bool(layer.isValid()), "local": local})
    project.clear()
    invalid = [item["layer"] for item in report if item["local"] and not item["valid"]]
    if invalid:
        raise QualityError(
            f"ชั้นในโฟลเดอร์พื้นที่ไม่ valid หลังย้าย ({qgz.name}): {', '.join(invalid)}"
        )
    return report


def relayout_area(
    machine: dict[str, Any],
    catalog: dict[str, Any],
    area: dict[str, Any],
    confirmation: str | None,
    qgis_check: bool = False,
) -> dict[str, Any]:
    code = area["areacode"]
    area_dir = ensure_within(AREA_ROOT / area["folder_name"], AREA_ROOT, "โฟลเดอร์รายพื้นที่")
    if not area_dir.is_dir():
        raise SafetyError(f"ไม่พบโฟลเดอร์ผลของพื้นที่: {area_dir}")
    plan = plan_relayout(area, area_dir)
    print(f"relayout {area['name']} ({code}) -> {area_dir}")
    print(f"  คงที่ root: {len(plan['keep_root'])} ไฟล์ {plan['keep_root']}")
    print(f"  อยู่ถูกที่แล้ว: {len(plan['in_place'])} ไฟล์")
    print(f"  จะย้าย: {len(plan['moves'])} ไฟล์")
    for source, destination in plan["moves"]:
        print(f"    {source.name}  ->  {destination.relative_to(area_dir).as_posix()}")
    if plan["problems"]:
        for problem in plan["problems"]:
            print(f"  [ปัญหา] {problem}")
        raise SafetyError("พบไฟล์ที่จัด layout ไม่ได้; ยังไม่ย้ายไฟล์ใด")
    if confirmation is None:
        print("  dry-run: ยังไม่ย้าย ใส่ --confirm <areacode> เพื่อย้ายจริง")
        return {"dry_run": True, "moves": len(plan["moves"]), "in_place": len(plan["in_place"])}
    if confirmation.strip() != code:
        raise SafetyError("รหัสยืนยันไม่ตรง; ยังไม่ย้ายไฟล์ใด")

    with pipeline_lock("relayout", code):
        processes = qgis_processes()
        if processes["gui"] or processes["workers"]:
            raise SafetyError("พบ QGIS/qgis_process เปิดอยู่; ปิดก่อนย้ายไฟล์")
        for source, _ in plan["moves"]:
            if file_is_locked(source):
                raise SafetyError(f"ไฟล์ถูกเปิดใช้งานอยู่: {source.name}")
        for project in plan["projects"]:
            if file_is_locked(project):
                raise SafetyError(f"ไฟล์โครงการถูกเปิดใช้งานอยู่: {project.name}")
        token = uuid.uuid4().hex[:12]
        run_root = local_state_root() / "runs" / code
        run_root.mkdir(parents=True, exist_ok=True)
        run_dir = run_root / f"relayout-{dt.datetime.now():%Y%m%d-%H%M%S}-{token}"
        run_dir.mkdir(parents=False, exist_ok=False)
        records = [
            {
                "rel_from": source.relative_to(area_dir).as_posix(),
                "rel_to": destination.relative_to(area_dir).as_posix(),
                "bytes": source.stat().st_size,
                "sha256": sha256_file(source),
            }
            for source, destination in plan["moves"]
        ]
        project_records: list[dict[str, Any]] = []
        for project in plan["projects"]:
            backup = run_dir / f"{project.name}.before"
            shutil.copy2(project, backup)
            project_records.append(
                {"name": project.name, "before_sha256": sha256_file(project), "backup": str(backup)}
            )
        audit: dict[str, Any] = {
            "schema_version": 1,
            "status": "IN_PROGRESS",
            "action": "relayout",
            "pipeline_version": VERSION,
            "layout_version": layout.LAYOUT_VERSION,
            "subdirs": list(layout.SUBDIRS),
            "area": {"areacode": code, "name": area["name"], "folder_name": area["folder_name"]},
            "started_at": local_now(),
            "keep_root": plan["keep_root"],
            "already_in_place": plan["in_place"],
            "moves": records,
            "projects": project_records,
        }
        audit_partial = run_dir / "relayout_audit.partial.json"
        write_json(audit_partial, audit)

        ensure_area_subdirs(area_dir)
        for source, destination in plan["moves"]:
            ensure_within(destination, area_dir, "ปลายทางการย้าย")
            if destination.exists():
                raise SafetyError(f"ปลายทางมีไฟล์อยู่แล้ว: {destination}")
            os.replace(source, destination)
        for record in records:
            destination = area_dir / record["rel_to"]
            if (
                destination.stat().st_size != record["bytes"]
                or sha256_file(destination) != record["sha256"]
            ):
                raise QualityError(f"ไฟล์หลังย้ายไม่ตรงต้นฉบับ: {record['rel_to']}")

        gpkg_names = sorted(
            path.name
            for path in (area_dir / layout.GIS_DIR).iterdir()
            if path.is_file() and path.suffix.lower() == ".gpkg" and ".partial." not in path.name
        )
        for project, record in zip(plan["projects"], project_records):
            outcome = rewrite_qgz_for_layout(project, area_dir, gpkg_names, token)
            record.update(outcome)
            record["after_sha256"] = sha256_file(project)
            record["validation"] = validate_qgz_layout(project, area_dir, gpkg_names)
            if qgis_check:
                record["qgis_layers"] = qgis_check_project(machine, project, area_dir)

        if area.get("status") == "ready":
            verify_area_assets(area)
        else:
            check_area_layout(area_dir, plan["qc_name"])

        audit["status"] = "DONE"
        audit["finished_at"] = local_now()
        write_json(audit_partial, audit)
        os.replace(audit_partial, run_dir / "relayout_audit.json")
        print(f"  ย้ายสำเร็จ {len(records)} ไฟล์; audit: {run_dir / 'relayout_audit.json'}")
        return audit


def _boundary_from_external(machine: dict[str, Any], area: dict[str, Any]) -> Any:
    """อ่านขอบเขต สพป. หนึ่ง feature จาก shapefile ภายนอกด้วย exact path (ไม่สแกนโฟลเดอร์)"""
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas ใน QGIS Python") from exc
    boundary_path = exact_external_paths(machine)["education_boundaries"]
    boundary_name = str(area.get("boundary_name", "")).strip()
    if not boundary_name:
        raise QualityError(f"areas.json ไม่มี boundary_name สำหรับ {area['name']}")
    escaped_name = boundary_name.replace("'", "''")
    try:
        boundary = gpd.read_file(boundary_path, where=f'"{BOUNDARY_NAME_FIELD}" = \'{escaped_name}\'')
    except (TypeError, ValueError, OSError):
        all_boundaries = gpd.read_file(boundary_path)
        if BOUNDARY_NAME_FIELD not in all_boundaries.columns:
            raise QualityError(f"ขอบเขต สพป. ไม่มี field {BOUNDARY_NAME_FIELD}")
        boundary = all_boundaries[
            all_boundaries[BOUNDARY_NAME_FIELD].astype(str).str.strip() == boundary_name
        ].copy()
    if len(boundary) != 1:
        raise QualityError(f"ขอบเขต {area['name']} ต้องพบ 1 feature แต่พบ {len(boundary)}")
    return boundary.to_crs("EPSG:32647")


def _discover_area_frame(
    machine: dict[str, Any], area: dict[str, Any], area_dir: Path, kind: str, prefixes: tuple[str, ...]
) -> Any:
    """หาชั้น GIS ของพื้นที่จากชื่อมาตรฐานก่อน แล้วจึงชื่อ legacy ใน 01 ข้อมูล GIS; ขอบเขตใช้ฐานภายนอกเป็นทางสุดท้าย"""
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise PipelineError("ไม่พบ GeoPandas ใน QGIS Python") from exc
    names = output_names(area)
    candidate = area_dir / layout.area_relpath(kind, names[kind])
    if candidate.is_file():
        return gpd.read_file(candidate)
    gis_dir = area_dir / layout.GIS_DIR
    matches = (
        sorted(path for path in gis_dir.glob("*.gpkg") if path.name.startswith(prefixes) and ".partial." not in path.name)
        if gis_dir.is_dir() else []
    )
    if len(matches) == 1:
        print(f"[ข้อมูล] ใช้ชั้น {kind} จากไฟล์ legacy: {matches[0].name}")
        return gpd.read_file(matches[0])
    if kind == "boundary":
        print(f"[ข้อมูล] ไม่มีขอบเขตในโฟลเดอร์ อ่านจากฐานภายนอกด้วยชื่อ {area.get('boundary_name')}")
        return _boundary_from_external(machine, area)
    print(f"[คำเตือน] ไม่พบชั้น {kind} ของ {area['name']} วาดแผนที่โดยไม่มีชั้นนี้")
    return None


def render_size_map_pending(
    machine: dict[str, Any],
    catalog: dict[str, Any],
    area: dict[str, Any],
    confirmation: str,
    overwrite: bool = False,
) -> dict[str, Any]:
    """วาดเฉพาะแผนที่ขนาดโรงเรียนให้พื้นที่ pending ที่มีโฟลเดอร์แล้ว ไม่แตะ QC และไม่ขึ้นพอร์ทัล"""
    code = area["areacode"]
    if confirmation.strip() != code:
        raise SafetyError("รหัสยืนยันไม่ตรง; ยังไม่วาดแผนที่")
    if area.get("status") == "ready":
        raise SafetyError("พื้นที่ ready ให้ใช้ refresh-assets เพื่อให้ QC และพอร์ทัลตรงกัน")
    if area.get("status") in {"blocked", "empty"}:
        raise SafetyError(f"พื้นที่ {area['name']} มีสถานะ {area['status']} ยังวาดแผนที่ขนาดไม่ได้")
    area_dir = ensure_within(AREA_ROOT / area["folder_name"], AREA_ROOT, "โฟลเดอร์รายพื้นที่")
    if not area_dir.is_dir():
        raise SafetyError("ไม่มีโฟลเดอร์ผลของพื้นที่นี้; แผนที่ขนาดโรงเรียนสร้างเฉพาะพื้นที่ที่มีโฟลเดอร์แล้ว")
    require_ready_status(machine, catalog)
    names = output_names(area)
    final = ensure_within(
        area_dir / layout.area_relpath("size_image", names["size_image"]), AREA_ROOT, "PNG ขนาดโรงเรียน"
    )
    token = uuid.uuid4().hex[:12]
    partial = make_partial(final, token)
    preflight_outputs({"size_image": final}, overwrite)
    with pipeline_lock("size-map", code):
        processes = qgis_processes()
        if processes["gui"] or processes["workers"]:
            raise SafetyError("พบ QGIS/qgis_process เปิดอยู่; ยังไม่วาดแผนที่")
        try:
            small_rows, _, _ = load_area_records(area)
            profile = school_size_profile(small_rows)
            boundary = _discover_area_frame(machine, area, area_dir, "boundary", ("ขอบเขต_",))
            roads = _discover_area_frame(machine, area, area_dir, "roads", ("ถนน_", "roads_"))
            tambon_lines, tambon_meta = load_tambon_lines(machine, boundary)
            print(f"[ข้อมูล] แนวเขตตำบลจาก {tambon_meta['layer']}: {tambon_meta['tambon_count']} ตำบล")
            ensure_area_subdirs(area_dir)
            render_size_map_frames(area, boundary, roads, small_rows, partial, profile, tambon_lines=tambon_lines)
            transactional_promote({"size_image": partial}, {"size_image": final}, overwrite, token)
            check_area_layout(area_dir, names["qc"])
            print(f"[สำเร็จ] แผนที่ขนาดโรงเรียน {area['name']}: {final}")
            print("  ช่วงขนาด: " + " · ".join(f"{item['label']}={item['count']}" for item in profile["bins"]))
            shown = sum(int(item["count"]) for item in profile["bins"] if int(item["max"]) <= SIZE_MAP_MAX_STUDENTS)
            print(f"  วาดบนแผนที่ (ไม่เกิน {SIZE_MAP_MAX_STUDENTS} คน): {shown} จาก {profile['total']} แห่ง")
            return profile
        except BaseException:
            cleanup_partials_under({"size_image": partial}, AREA_ROOT, token)
            raise


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="ระบบคำนวณและพอร์ทัลผลการจับคู่โรงเรียนขนาดเล็ก"
    )
    parser.add_argument("--version", action="version", version=f"%(prog)s {VERSION}")
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("menu", help="เปิดเมนูแบบโต้ตอบ")
    status_parser = subparsers.add_parser("status", help="ตรวจสถานะข้อมูลและเครื่อง")
    status_parser.add_argument("--quick", action="store_true", help="ข้าม full SHA/fingerprint")
    run_parser = subparsers.add_parser("run", help="คำนวณหนึ่งพื้นที่")
    run_parser.add_argument("--area", required=True, help="areacode 8 หลัก")
    run_parser.add_argument("--confirm", required=True, help="พิมพ์ areacode ซ้ำ")
    run_parser.add_argument("--overwrite", action="store_true")
    pilot_parser = subparsers.add_parser("pilot-assets", help="สร้าง PNG/PDF/QC ของพื้นที่นำร่อง")
    pilot_parser.add_argument("--confirm", required=True, help="ต้องเป็น PILOT")
    pilot_parser.add_argument("--area", action="append", help="areacode ready; ระบุซ้ำได้")
    pilot_parser.add_argument("--overwrite", action="store_true")
    refresh_parser = subparsers.add_parser(
        "refresh-assets", help="ปรับ PNG/PDF ทุกพื้นที่ ready จากผลเดิมโดยไม่คำนวณ GIS ซ้ำ"
    )
    refresh_parser.add_argument("--confirm", required=True, help="ต้องเป็น REFRESH")
    refresh_parser.add_argument("--area", action="append", help="areacode ready; ระบุซ้ำได้")
    portal_parser = subparsers.add_parser("portal", help="สร้างพอร์ทัลจากผลที่ผ่าน QC")
    portal_parser.add_argument("--confirm", required=True, help="ต้องเป็น PUBLISH")
    portal_parser.add_argument("--overwrite", action="store_true")
    recompute_parser = subparsers.add_parser(
        "recompute-check", help="ทดสอบ QNEAT3 นำร่องใน LOCALAPPDATA โดยไม่ publish"
    )
    recompute_parser.add_argument("--area", required=True, help="areacode นำร่อง")
    recompute_parser.add_argument("--confirm", required=True, help="พิมพ์ areacode ซ้ำ")
    recompute_parser.add_argument("--runs", type=int, choices=(1, 2), default=2)
    relayout_parser = subparsers.add_parser(
        "relayout",
        help="ย้ายไฟล์ของพื้นที่เข้าโฟลเดอร์ย่อยตาม area_layout (dry-run เมื่อไม่ใส่ --confirm)",
    )
    relayout_parser.add_argument("--area", required=True, help="areacode 8 หลัก")
    relayout_parser.add_argument("--confirm", help="พิมพ์ areacode ซ้ำเพื่อย้ายจริง")
    relayout_parser.add_argument(
        "--qgis-check",
        action="store_true",
        help="โหลด .qgz ด้วย PyQGIS หลังย้ายเพื่อยืนยันว่าทุกชั้นในโฟลเดอร์พื้นที่ valid",
    )
    sizemap_parser = subparsers.add_parser(
        "size-map",
        help="วาดแผนที่ขนาดโรงเรียนให้พื้นที่ pending ที่มีโฟลเดอร์แล้ว (พื้นที่ ready ใช้ refresh-assets)",
    )
    sizemap_parser.add_argument("--area", required=True, help="areacode 8 หลัก")
    sizemap_parser.add_argument("--confirm", required=True, help="พิมพ์ areacode ซ้ำ")
    sizemap_parser.add_argument("--overwrite", action="store_true")
    subparsers.add_parser("verify", help="ตรวจรับผลและพอร์ทัล")
    return parser


def main(argv: list[str] | None = None) -> int:
    configure_console()
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        machine, catalog = load_settings()
        if args.command == "menu":
            return menu(machine, catalog)
        if args.command == "status":
            report = status_report(machine, catalog, deep_fingerprints=not args.quick)
            print_status(report)
            return 0 if report["ready"] else 2
        if args.command == "run":
            area = area_by_code(catalog, args.area)
            run_area(machine, catalog, area, args.confirm, overwrite=args.overwrite)
            return 0
        if args.command == "pilot-assets":
            build_pilot_assets(
                machine,
                catalog,
                args.confirm,
                overwrite=args.overwrite,
                area_codes=args.area,
            )
            return 0
        if args.command == "portal":
            publish_portal(catalog, args.confirm, overwrite=args.overwrite)
            return 0
        if args.command == "refresh-assets":
            refresh_assets(machine, catalog, args.confirm, area_codes=args.area)
            return 0
        if args.command == "recompute-check":
            area = area_by_code(catalog, args.area)
            recompute_check(machine, catalog, area, args.confirm, runs=args.runs)
            return 0
        if args.command == "size-map":
            area = area_by_code(catalog, args.area)
            render_size_map_pending(machine, catalog, area, args.confirm, overwrite=args.overwrite)
            return 0
        if args.command == "relayout":
            area = area_by_code(catalog, args.area)
            relayout_area(machine, catalog, area, args.confirm, qgis_check=args.qgis_check)
            return 0
        if args.command == "verify":
            verify_delivery(catalog)
            return 0
        parser.error(f"คำสั่งไม่รองรับ: {args.command}")
    except KeyboardInterrupt:
        if _ACTIVE_SUBPROCESS is not None:
            stop_process(_ACTIVE_SUBPROCESS)
        print("\nยกเลิกโดยผู้ใช้; ผลเดิมไม่ถูกเขียนทับ", file=sys.stderr)
        return 130
    except PipelineError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 2
    except Exception as exc:
        print(f"[ERROR] ข้อผิดพลาดที่ไม่ได้คาดหมาย: {exc}", file=sys.stderr)
        traceback.print_exc()
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
